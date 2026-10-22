"""Presentación Excel y cálculos económicos editables, independientes del motor."""
from copy import copy
import math
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from app_icociv.interfaz.widgets.tabla_numerica import numero_original, variacion_indice


def estilo_escenario(celda, editable=False):
    celda.fill = PatternFill('solid', fgColor='FFF2CC' if editable else 'E7E6E6')
    celda.font = Font(color='1F4E78' if editable else '202020')
    celda.protection = Protection(locked=not editable)


def agregar_desglose_importes(ws, registros, referencias, clave_valor, contexto):
    """Conserva principales registrados y enlaza los cuatro importes por ID."""
    from app_icociv.utilidades.importes_presupuestales import MODALIDADES, desglose_importe, desglose_presupuestal, _finito
    from app_icociv.interfaz.widgets.empalme_iccp_icociv import _escribir_tabla_excel
    if not registros:return
    assert len(registros)==len(referencias)
    inicio=ws.max_row+3;filas=[];desgloses=[]
    for r,ref in zip(registros,referencias):
        valor=r.get(clave_valor);d=desglose_presupuestal(r,valor);desgloses.append(d)
        b=desglose_importe(r,r.get('base_ajustable'),base=True);v=desglose_importe(r,valor)
        estado=' · '.join(dict.fromkeys(x['estado'] for x in (b,v) if x['estado'])) or ('Importes vinculados; contraparte por Cantidad' if ref else 'Importes fijos; sin vínculo vivo demostrable')
        fila=[r.get('calculo_id') or '',MODALIDADES[v['modalidad']],r.get('cantidad'),r.get('precio_base'),r.get('anticipo_amortizado'),r.get('base_ajustable'),valor,*d.values(),estado]
        filas.append([x if i in (0,1,11) or _finito(x) else None for i,x in enumerate(fila)])
    _escribir_tabla_excel(ws,inicio,'Desglose presupuestal complementario · '+contexto+' · Sin clasificar: principales conservados; pares vacíos',
        ['ID cálculo','Interpretación P/A','Cantidad presupuestal','P original','A original','Base principal original','Valor principal / escenario','Base ajustable unitaria','Base ajustable total','Valor ajustado unitario','Valor ajustado total','Disponibilidad'],filas)
    for f,(ref,d) in enumerate(zip(referencias,desgloses),inicio+2):
        if ref:
            ws.cell(f,3,'='+ref['cantidad']);ws.cell(f,7,'='+ref['valor'])
            for col,k in enumerate(d,8):
                if d[k] is not None:ws.cell(f,col,'='+ref[k])
        for col in range(3,12):
            valor=filas[f-inicio-2][col-1]
            ws.cell(f,col).number_format='0.0000E+00' if _finito(valor) and valor!=0 and (abs(valor)<0.00005 or abs(valor)>=1e12) else '#,##0.0000'


def agregar_calculos_editables(wb, registros, empalme=False):
    """Reproduce solamente las relaciones económicas de los servicios existentes."""
    from app_icociv.interfaz.widgets.empalme_iccp_icociv import _estilo_encabezado_tabla, _estilo_celda_tabla
    ws = wb.create_sheet('Cálculos editables')
    ws['A1'] = 'Cálculos económicos editables'
    ws['A2'] = 'Amarillo: entradas. Gris: fórmulas. Columna C: resultados registrados por SAVIP, independientes de la edición.'
    ws['A3'] = 'Cantidad es informativa. Los índices estadísticos se toman del análisis guardado; esta hoja no ejecuta ni reproduce el motor.'
    ws.freeze_panes = 'B5'
    ws.sheet_view.showGridLines = False
    fila = 5
    for i, registro in enumerate(registros, 1):
        ws.cell(fila, 1, f"{i}. {registro.get('item') or registro.get('insumo') or ''}").data_type = 's'
        ws.cell(fila, 4, f"{registro.get('fecha_inicial') or registro.get('periodo_base') or ''} → {registro.get('fecha_final') or registro.get('periodo_proyectado') or ''}; ID {registro.get('calculo_id') or ''}; {registro.get('caso') or 'ICOCIV'}").data_type = 's'
        fila += 1
        for col, title in enumerate(('Campo', 'Editable / calculado', 'Registrado SAVIP', 'Uso'), 1):
            _estilo_encabezado_tabla(ws.cell(fila, col, title))
        refs = {}
        def campo(nombre, clave=None, formula=None, editable=False, precision=2, nota=''):
            nonlocal fila
            fila += 1
            original = numero_original(registro.get(clave)) if clave else None
            for col, val in ((1, nombre), (2, formula if formula is not None else original), (3, original), (4, nota)):
                cell = ws.cell(fila, col, val)
                _estilo_celda_tabla(cell)
                if col != 2 and isinstance(val, str):
                    cell.data_type = 's'
            for col in (2, 3):
                ws.cell(fila, col).number_format = '#,##0.' + '0' * precision if precision else '0'
            cell = ws.cell(fila, 2)
            cell.fill = PatternFill('solid', fgColor='FFF2CC' if editable else 'E7E6E6')
            cell.font = Font(color='1F4E78' if editable else '202020')
            cell.protection = Protection(locked=not editable)
            refs[nombre] = f'B{fila}'
            return refs[nombre]
        acero = empalme and str(registro.get('tipo_calculo', '')).lower().endswith('acero')
        caso = registro.get('caso') if empalme else 'solo_icociv'
        if caso not in ('solo_iccp', 'solo_icociv', 'empalme_completo'):
            fila += 1
            ws.cell(fila, 1, 'Fórmula no incorporada: falta un escenario guardado demostrable. Consulte Resultados registrados.')
            fila += 2
            continue
        p = campo('P0 (COP)' if acero else 'P (COP)', 'p0' if acero else 'precio_base', editable=True)
        a = campo('A (COP)', 'anticipo_amortizado', editable=not acero, nota='El acero usa P0 sin restar anticipo.' if acero else '')
        if acero:
            ws[a] = '=0'
        campo('Cantidad', 'cantidad', editable=True, precision=4, nota='No multiplica automáticamente el ajuste.')
        entradas = [p] if acero else [p, a]
        validacion = [f'{p}<=0'] if acero else [f'{p}<=0', f'{a}<0', f'{a}>={p}']
        if caso in ('solo_iccp', 'empalme_completo'):
            i0c = campo('I0 ICCP', 'i0_iccp', editable=True, precision=4)
            ic = campo('I ICCP', 'i_iccp', editable=True, precision=4)
            entradas += [i0c, ic];validacion.append(f'{i0c}=0')
        if caso in ('solo_icociv', 'empalme_completo'):
            i0 = campo('I0 ICOCIV', 'i0_icociv', editable=True, precision=4)
            indice = campo('I ICOCIV', 'i_icociv', editable=True, precision=4)
            entradas += [i0, indice];validacion.append(f'{i0}<=0' if not empalme else f'{i0}=0')
        validacion = [f'NOT(ISNUMBER({c}))' for c in entradas] + validacion
        estado = campo('Estado de entradas', formula=f'=IF(OR({",".join(validacion)}),"Entradas inválidas","Cálculo disponible")', nota='Fechas y escenario son los guardados. Editar índices no genera una proyección.')
        reproducible = all(isinstance(ws[c].value, (int, float)) and math.isfinite(ws[c].value) for c in entradas)
        if reproducible:
            reproducible = ws[p].value > 0 and (acero or 0 <= ws[a].value < ws[p].value)
            if caso in ('solo_iccp', 'empalme_completo'):
                reproducible = reproducible and ws[i0c].value != 0
            if caso in ('solo_icociv', 'empalme_completo'):
                reproducible = reproducible and (ws[i0].value != 0 if empalme else ws[i0].value > 0)
        def calculo(nombre, clave, expresion, precision=2):
            return campo(nombre, clave, f'={expresion}' if reproducible else None, precision=precision)
        base = calculo('Base (COP)', 'base_ajustable', p if acero else f'{p}-{a}')
        if empalme:
            r1 = calculo('R1 (COP)', 'r1', f'{base}*({ic}/{i0c}-1)' if caso in ('solo_iccp', 'empalme_completo') else '0')
            r2 = calculo('R2 (COP)', 'r2', f'({base}+{r1})*({indice}/{i0}-1)' if caso in ('solo_icociv', 'empalme_completo') else '0')
            ajuste = calculo('R (COP)', 'r_total', f'{r1}+{r2}')
        else:
            calculo('Factor ICOCIV', 'factor_icociv', f'{indice}/{i0}', precision=6)
            ajuste = calculo('R (COP)', 'r_total', f'{base}*({indice}/{i0}-1)')
        calculo('Valor actualizado (COP)', 'valor_actualizado' if empalme else 'valor_proyectado', f'{base}+{ajuste}')
        calculo('Diferencia (%)', 'diferencia_porcentual', f'({ajuste}/{base})*100', precision=3)
        if empalme:
            variacion = calculo('Variación del índice (%)', None, f'({indice}/{i0}-1)*100' if caso != 'solo_iccp' else '""', precision=3)
            ws.cell(int(variacion[1:]), 3, variacion_indice(registro.get('i0_icociv'), registro.get('i_icociv')))
        if acero:
            ix = campo('Ix (COP/kg)', 'ix', editable=True)
            q = campo('q (kg)', 'q', editable=True, precision=4)
            facturado = campo('Ix × q (COP)', 'valor_facturado_total', f'=IF(OR({ix}="",{q}=""),"",{ix}*{q})')
            campo('Z (COP)', 'z', f'=IF({facturado}="","",{facturado}-({ajuste}+{p}))', nota='La falta de Ix/q deja Z vacío; no afecta R.')
        fila += 2
    ajustar_presentacion(ws)
    # Las notas largas no ensanchan las entradas ni las fórmulas.
    for col, width in (('A', 30), ('B', 30), ('C', 30), ('D', 46)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[2].height = 45
    ws.row_dimensions[3].height = 60
    for cell in (ws['A2'], ws['A3']):
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.merge_cells('A2:D2');ws.merge_cells('A3:D3')
    return ws


def agregar_resultados_empalme(wb, registros):
    from app_icociv.interfaz.widgets.empalme_iccp_icociv import _escribir_tabla_excel
    ws = wb.create_sheet('Resultados registrados')
    campos = [('N.º', None), ('Ítem', 'item'), ('Escenario', 'caso'), ('P / P0 (COP)', 'precio_base'), ('A (COP)', 'anticipo_amortizado'),
              ('Cantidad', 'cantidad'), ('Base (COP)', 'base_ajustable'), ('R1 (COP)', 'r1'), ('R2 (COP)', 'r2'), ('R (COP)', 'r_total'),
              ('Valor actualizado (COP)', 'valor_actualizado'), ('Diferencia (%)', 'diferencia_porcentual'), ('Ix × q (COP)', 'valor_facturado_total'), ('Z (COP)', 'z')]
    rows = [[i, r.get('item'), r.get('caso')] + [numero_original(r.get(k)) for _, k in campos[3:]] + [variacion_indice(r.get('i0_icociv'), r.get('i_icociv')), r.get('tipo_componente'), r.get('nombre_componente', '')] for i, r in enumerate(registros, 1)]
    _escribir_tabla_excel(ws, 1, 'Resultados registrados por SAVIP sin fórmulas ni edición de entradas', [h for h, _ in campos] + ['Variación del índice (%)', 'Tipo de componente', 'Nombre del componente'], rows)
    ws.freeze_panes = 'B3';ws.auto_filter.ref = f'A2:Q{2+len(registros)}'
    for row in ws.iter_rows(min_row=3):
        for col in (12, 15):
            row[col-1].number_format = '#,##0.000'
    return ws


def ajustar_presentacion(ws):
    encabezados = {}
    ancho_entrada = max((len(f'{c.value:,.4f}') for row in ws for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)), default=12)
    for row in ws:
        for cell in row:
            if cell.value is None:
                continue
            if cell.fill.fgColor.rgb == '00475569':
                encabezados[cell.column] = str(cell.value).strip().lower()
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                titulo = encabezados.get(cell.column, '')
                if titulo in ('n°', 'n.º', 'nº', 'parte', 'registros', 'horizonte', 'horizonte solicitado') or titulo.startswith('n.º en '):
                    cell.number_format = '0'
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = 'center'
            if cell.data_type in ('n', 'f'):
                alignment.horizontal = 'right'
                if ws.title in ('Valores proyectados', 'Empalme ICCP-ICOCIV') and cell.fill.fgColor.rgb in ('00FFF2CC', '00E7E6E6'):
                    alignment.wrap_text = False
            cell.alignment = alignment
    for column in ws.columns:
        cells = [c for c in column if c.value is not None and c.coordinate not in ws.merged_cells]
        if not cells:
            continue
        headers = [str(c.value) for c in cells if c.fill.fgColor.rgb == '00475569']
        body = [c for c in cells if c.fill.fgColor.rgb != '00475569']
        numeric = bool(body) and all(c.data_type in ('n', 'f', 'b') for c in body)
        if numeric:
            lengths = [len(f'{c.value:,.4f}') if isinstance(c.value, (int, float)) else 14 for c in body]
            if any(c.data_type == 'f' and '$' in c.number_format for c in body):
                lengths.append(ancho_entrada + 2)
            width = min(32, max(10, max(lengths, default=10) + 2))
        else:
            cap = 60 if any('JSON' in h for h in headers) else 42
            width = min(cap, max(14, max(len(line) for c in cells for line in str(c.value).splitlines() or ['']) + 2))
        if headers and all(h.lower() in ('n°', 'n.º', 'nº', 'parte') for h in headers):
            width = 7
        ws.column_dimensions[get_column_letter(column[0].column)].width = width
    if ws.title == 'Cálculos editables':
        for col, width in (('A', 30), ('B', 30), ('C', 30), ('D', 46)):
            ws.column_dimensions[col].width = width
    for row in ws:
        lines = 1
        for c in row:
            if c.value is None or c.coordinate in ws.merged_cells or c.data_type in ('f', 'n'):
                continue
            width = ws.column_dimensions[c.column_letter].width or 14
            lines = max(lines, sum(max(1, math.ceil(len(t) / max(1, width - 2))) for t in str(c.value).splitlines()))
        if lines > 1:
            ws.row_dimensions[row[0].row].height = max(ws.row_dimensions[row[0].row].height or 18, min(120, lines * 15 + 6))


def validar_cantidad_excel(ws, columna, ultima_fila):
    if ultima_fila<3:return
    from openpyxl.worksheet.datavalidation import DataValidation
    validacion=DataValidation(type='decimal',operator='greaterThan',formula1='0',allow_blank=False)
    validacion.showErrorMessage=True;validacion.errorStyle='stop'
    validacion.errorTitle='Cantidad presupuestal'
    validacion.error='Use un número mayor que cero para convertir entre importes unitarios y totales.'
    ws.add_data_validation(validacion);validacion.add(f'{columna}3:{columna}{ultima_fila}')
