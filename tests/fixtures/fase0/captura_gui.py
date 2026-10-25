"""Fase 0: escenario de sesión y exportables sobre la ventana real (Qt offscreen) con un anexo sintético.

Recorrido: cargar anexo → proyectar serie 1 a H11 y serie 2 a H23 (motor real) → actividades unitaria,
total y de la serie 2 → consultar H11→H23→H11 sin motor → deshacer/rehacer → Guardar como, Ctrl+S,
cancelación y error de escritura → reabrir sin motor → exportar Excel, PDF y DOCX de la sesión, y CSV,
DOCX y PDF del informe de proyección. Devuelve un resumen normalizado (sin fechas, IDs ni rutas)."""

from __future__ import annotations

import os
import re
import time
from csv import reader as lector_csv
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from app_icociv.interfaz.controladores.controlador_principal import ControladorPrincipal
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.reportes import generador_reportes as gr
from app_icociv.reportes.modelo import ConfiguracionInforme
from app_icociv.servicios import fecha_objetivo as fo
from tests.fixtures.fase0 import generadores as gen
from tests.fixtures.fase0.captura import json_estable, normalizar_texto

APP = QApplication.instance() or QApplication([])
_MOTOR = ControladorPrincipal.ejecutar_analisis
HOJAS_EXCEL = ["Resumen", "Valores proyectados", "Proyecciones ICOCIV", "Trazabilidad", "Metodología y fórmulas"]
CLAVES_VOLATILES = {"fecha_calculo", "fecha_guardado", "creado", "fecha", "fecha_exportacion", "actualizado"}


class Contador:
    """Espía del motor real (cuenta llamadas y delega en el método original)."""

    def __init__(self):
        self.n = 0

        def metodo(controlador, *a, **k):
            self.n += 1
            return _MOTOR(controlador, *a, **k)
        self.metodo = metodo


def _esperar(v, condicion, limite=180):
    fin = time.monotonic() + limite
    while time.monotonic() < fin:
        APP.processEvents()
        if condicion() and (v.trabajador is None or not v.trabajador.isRunning()) and not v._operacion_sesion:
            return
        QTest.qWait(20)
    raise AssertionError("No terminó: " + v.etiqueta_estado.text())


def _cargar(v, ruta):
    v._archivo_cargado(v.controlador.cargar_archivo(str(ruta))); APP.processEvents()


def _ejecutar(v, grupo, horizonte):
    """«Ejecutar proyección» con horizonte personalizado; prohíbe ventanas modales. Devuelve análisis nuevos."""
    w = v.widget_proyecciones; n = len(w.historial)
    v.combo_grupo.setCurrentIndex(grupo)
    v.combo_horizonte.setCurrentIndex(v.combo_horizonte.findData(None)); v.spin_horizonte_personalizado.setValue(horizonte)
    with patch.object(QMessageBox, "exec", side_effect=AssertionError("ventana modal")), \
         patch.object(QMessageBox, "question", side_effect=AssertionError("ventana modal")), \
         patch.object(v.notificaciones, "mostrar"):
        v.ejecutar_proyeccion()
        if v.trabajador is not None and v.trabajador.isRunning():
            _esperar(v, lambda: len(w.historial) > n)
    for _ in range(5):
        APP.processEvents()
    return len(w.historial) - n


def _calcular(w, nombre, modalidad, cantidad, precio, anticipo):
    w.tipo_componente.setCurrentText("Actividad"); w.nombre_componente.setText(nombre); w.item.setText(nombre)
    w._item_editado = True; w.unidad.setCurrentIndex(1); w.cantidad.setValue(cantidad); w.precio_base.setValue(precio)
    w.anticipo.setValue(anticipo); w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modalidad))
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


def _guardar(v, destino, como=False, atajo=False):
    with patch.object(v, "_solicitar_usuario", return_value="QA"), \
         patch.object(QFileDialog, "getSaveFileName", return_value=(str(destino), "")) as dialogo:
        if atajo:
            v.activateWindow(); v.widget_proyecciones.tabla_valores.setFocus(); APP.processEvents()
            QTest.keyClick(v.widget_proyecciones.tabla_valores, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
            _esperar(v, lambda: True); return None, dialogo.call_count
        resultado = v.guardar_rapido(como=como); _esperar(v, lambda: True)
        return resultado, dialogo.call_count


# ----------------------------------------------------------------------------------- normalización
CLAVES_ID = {"id", "calculo_id", "analisis_id", "registros_ids", "revision_de", "analisis_activo_id"}


def _es_hex_id(valor: str) -> bool:
    return len(valor) in (8, 32) and all(c in "0123456789abcdef" for c in valor)


def reemplazos_ids(datos: dict) -> tuple[tuple[str, str], ...]:
    """IDs aleatorios de la sesión (análisis uuid4, cálculos y filas técnicas) → fichas estables por orden de
    aparición: primero los análisis, luego un recorrido determinista de toda la sesión."""
    ids = [a["id"] for a in datos["proyecciones_icociv"].get("historial_analisis", [])]

    def recorrer(x):
        if isinstance(x, dict):
            for k, v in x.items():
                if k in CLAVES_ID and isinstance(v, str) and _es_hex_id(v):
                    ids.append(v)
                elif k in CLAVES_ID and isinstance(v, list):
                    ids.extend(i for i in v if isinstance(i, str) and _es_hex_id(i))
                recorrer(v)
        elif isinstance(x, list):
            for v in x:
                recorrer(v)
    recorrer(datos)
    unicos = list(dict.fromkeys(ids))
    fichas = {i: f"<ID{n}>" for n, i in enumerate(unicos, 1)}
    return tuple(sorted(fichas.items(), key=lambda x: -len(x[0])))  # primero los largos (uuid) que los cortos


def _normalizar(objeto, rutas, reemplazos):
    if isinstance(objeto, dict):
        return {normalizar_texto(str(k), rutas, reemplazos): ("<VOLATIL>" if k in CLAVES_VOLATILES and objeto[k] else _normalizar(v, rutas, reemplazos))
                for k, v in objeto.items()}
    if isinstance(objeto, list):
        return [_normalizar(x, rutas, reemplazos) for x in objeto]
    if isinstance(objeto, str):
        return normalizar_texto(objeto, rutas, reemplazos)
    return objeto


def resumen_sesion(datos: dict, rutas: tuple[str, ...], reemplazos) -> dict:
    """Lo que el Golden Master fija de una sesión guardada: estructura, presupuesto, consultas y contexto.

    El ``resultado`` completo del motor (varios MB) no se versiona: de él se fijan sus claves y las cifras
    que la interfaz publica; el motor ya tiene su propio Golden Master."""
    p = datos["proyecciones_icociv"]
    historial = []
    for a in p.get("historial_analisis", []):
        res = a.get("resultado") or {}
        historial.append({
            "id": a["id"], "nombre": a["nombre"], "contexto": a["contexto"], "parametros": a["parametros"],
            "seleccion": a["seleccion"], "consulta_horizonte": a.get("consulta_horizonte"), "registros_ids": a.get("registros_ids"),
            "trayectoria_24": fo.trayectoria(a), "origen_analisis": fo.origen_analisis(a),
            "resultado_claves": sorted(res), "resultado_modelo": res.get("modelo_codigo"),
            "resultado_y_proj": res.get("y_proj"), "resultado_rmse_seleccion": res.get("rmse_seleccion_oos"),
        })
    resumen = {
        "claves_sesion": sorted(datos), "claves_proyecciones_icociv": sorted(p),
        "corte_dane": datos.get("corte_dane"), "usuario": datos.get("usuario"),
        "parametros_proyeccion": datos.get("parametros_proyeccion"), "seleccion": datos.get("seleccion"),
        "proyecto": p.get("proyecto"), "valores": p.get("valores"), "proyecciones": p.get("proyecciones"),
        "contexto": p.get("contexto"), "calculadora": p.get("calculadora"), "analisis_activo_id": p.get("analisis_activo_id"),
        "filtrar_analisis": p.get("filtrar_analisis"), "comparacion_historica": p.get("comparacion_historica"),
        "fecha_objetivo_presupuesto": p.get("fecha_objetivo_presupuesto"), "excepciones_fecha": p.get("excepciones_fecha"),
        "historial": historial,
    }
    return json_estable(_normalizar(resumen, rutas, reemplazos))


def leer_excel(ruta, rutas, reemplazos=()) -> dict:
    libro = load_workbook(ruta)  # fórmulas tal cual (sin valores en caché)
    hojas = {}
    for hoja in libro.worksheets:
        celdas = [[c.coordinate, normalizar_texto(c.value, rutas, reemplazos) if isinstance(c.value, str) else c.value]
                  for fila in hoja.iter_rows() for c in fila if c.value is not None]
        hojas[hoja.title] = {"dimension": hoja.dimensions, "celdas": celdas}
    return json_estable({"hojas": libro.sheetnames, "contenido": hojas})


def _partible(*piezas: str) -> str:
    """Patrón que admite saltos de línea entre caracteres: las celdas estrechas del PDF parten los IDs y las
    horas en un punto que depende del ancho de cada carácter (distinto en cada ejecución)."""
    return r"\s*".join(piezas)


_FECHA_PDF = re.compile(_partible(*[r"\d"] * 4, "-", *[r"\d"] * 2, "-", *[r"\d"] * 2, "[T ]", *[r"\d"] * 2, ":", *[r"\d"] * 2, ":", *[r"\d"] * 2))
_INFORME_PDF = re.compile(_partible(*"SAVIP-INF-", *[r"\d"] * 8, "-", *[r"\d"] * 6))


def leer_pdf(ruta, rutas, reemplazos=()) -> dict:
    lector = PdfReader(str(ruta))
    paginas = []
    for pagina in lector.pages:
        texto = pagina.extract_text() or ""
        for viejo, nuevo in reemplazos:
            texto = re.sub(_partible(*map(re.escape, viejo)), nuevo, texto)
        texto = _INFORME_PDF.sub("SAVIP-INF-<FECHA-HORA>", _FECHA_PDF.sub("<FECHA-HORA>", texto))
        paginas.append(normalizar_texto(texto, rutas))
    return {"paginas": len(paginas), "texto": [l for p in paginas for l in p.splitlines() if l.strip()]}


def leer_docx(ruta, rutas, reemplazos=()) -> dict:
    doc = Document(str(ruta))
    parrafos = [normalizar_texto(p.text, rutas, reemplazos) for p in doc.paragraphs if p.text.strip()]
    tablas = [[[normalizar_texto(c.text, rutas, reemplazos) for c in f.cells] for f in t.rows] for t in doc.tables]
    return {"parrafos": parrafos, "tablas": tablas}


# ----------------------------------------------------------------------------------- escenario
def escenario(tmp: Path) -> dict:
    tmp = Path(tmp); rutas = (str(tmp),)
    anexo = gen.anexo_sintetico(tmp / "anex-ICOCIV-sintetico-dic2026.xlsx")
    v = VentanaPrincipal(); v.resize(1366, 900); v.show(); w = v.widget_proyecciones
    motor = Contador(); salida: dict = {"invariantes": {}}
    inv = salida["invariantes"]
    otra = None
    try:
        with patch.object(ControladorPrincipal, "ejecutar_analisis", motor.metodo):
            _cargar(v, anexo)
            inv["corte_anexo"] = v._corte_dane_cargado()
            inv["serie1_h11_nuevos"] = _ejecutar(v, 1, 11)
            a = w.historial[0]
            _calcular(w, "Actividad unitaria", "unitario", 2, 1000, 100)
            _calcular(w, "Actividad total", "total", 3, 2500.55, 0)
            inv["serie2_h23_nuevos"] = _ejecutar(v, 2, 23)
            _calcular(w, "Actividad serie 2", "unitario", 1.5, 98765.43, 1234.5)
            inv["motor_tras_dos_series"] = motor.n
            valores = deepcopy(w.valores)
            consultas = []
            for h in (23, 11):  # H11 → H23 → H11 sobre la serie 1: trayectoria guardada, sin motor
                inv[f"serie1_h{h}_reconsulta_nuevos"] = _ejecutar(v, 1, h)
                c = w.historial[0]["consulta_horizonte"]
                consultas.append({"horizonte": c["horizonte"], "periodo": c["periodo"], "indice": c["indice"],
                                  "igual_a_trayectoria": c["indice"] == fo.trayectoria(a)[h - 1]})
            salida["consultas_serie1"] = consultas
            inv["motor_tras_reconsultas"] = motor.n
            inv["valores_intactos_tras_reconsultas"] = w.valores == valores
            # Deshacer/rehacer de un reordenamiento.
            orden = [r["calculo_id"] for r in w.valores]
            w._registrar_cambio("reordenar", list(reversed(w.valores)), w.valores[0]["calculo_id"]); APP.processEvents()
            invertido = [r["calculo_id"] for r in w.valores]
            w.pila_deshacer.undo(); APP.processEvents(); tras_deshacer = [r["calculo_id"] for r in w.valores]
            w.pila_deshacer.redo(); APP.processEvents(); tras_rehacer = [r["calculo_id"] for r in w.valores]
            w.pila_deshacer.undo(); APP.processEvents()
            inv["deshacer_restaura_orden"] = tras_deshacer == orden and invertido == orden[::-1]
            inv["rehacer_reaplica"] = tras_rehacer == invertido
            inv["orden_final_original"] = [r["calculo_id"] for r in w.valores] == orden

            # Guardar como (primera vez), Ctrl+S sobre la misma ruta, cancelación y error de escritura.
            ruta = tmp / "sesion_fase0.json"
            ok, dialogos = _guardar(v, ruta, como=True)
            inv["guardar_como"] = [ok, dialogos, ruta.exists(), v.ruta_sesion_actual == ruta]
            w.nombre_proyecto.setText("Presupuesto sintético Fase 0"); APP.processEvents()
            _, dialogos = _guardar(v, tmp / "no_debe_usarse.json", atajo=True)
            inv["ctrl_s_misma_ruta"] = [dialogos, not (tmp / "no_debe_usarse.json").exists(), v._firma_trabajo() == v._firma_guardada]
            bytes_guardados = ruta.read_bytes()
            w.nombre_proyecto.setText("Cambio no guardado"); APP.processEvents()
            ok, _ = _guardar(v, "", como=True)
            inv["guardar_como_cancelado"] = [ok, ruta.read_bytes() == bytes_guardados, v.ruta_sesion_actual == ruta]
            with patch("os.replace", side_effect=PermissionError("Acceso denegado (prueba)")), \
                 patch.object(QMessageBox, "warning") as aviso:
                ok, _ = _guardar(v, ruta)
            inv["error_escritura"] = [ok, ruta.read_bytes() == bytes_guardados, "siguen en memoria" in aviso.call_args[0][2],
                                      sorted(p.name for p in tmp.iterdir())]
            w.nombre_proyecto.setText("Presupuesto sintético Fase 0"); APP.processEvents()
            import json
            guardado = json.loads(ruta.read_text(encoding="utf-8"))
            reemplazos = reemplazos_ids(guardado)
            salida["sesion"] = resumen_sesion(guardado, rutas, reemplazos)

        # Reabrir en otra ventana con el motor prohibido: misma vista, sin recalcular.
        otra = VentanaPrincipal(); otra.resize(1366, 900); otra.show()
        _cargar(otra, anexo)
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(ruta), "")), \
             patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")), \
             patch.object(VentanaPrincipal, "_cargar_archivo", side_effect=AssertionError("no se abre el Excel de la sesión")):
            otra.abrir_sesion(); _esperar(otra, lambda: True)
        o = otra.widget_proyecciones
        inv["reabrir_valores_iguales"] = o.valores == guardado["proyecciones_icociv"]["valores"]
        inv["reabrir_consultas_iguales"] = [x.get("consulta_horizonte") for x in o.historial] == \
            [x.get("consulta_horizonte") for x in guardado["proyecciones_icociv"]["historial_analisis"]]
        inv["reabrir_contexto_igual"] = o.contexto == guardado["proyecciones_icociv"]["contexto"]
        inv["reabrir_analisis_activo"] = o.analisis_activo_id == guardado["proyecciones_icociv"]["analisis_activo_id"]

        # Exportables de la sesión reabierta.
        alcance = o.alcance_exportacion("sesion")
        o._generar_excel(str(tmp / "sesion.xlsx"), alcance)
        o._generar_informe(str(tmp / "sesion.pdf"), "pdf", alcance)
        o._generar_informe(str(tmp / "sesion.docx"), "docx", alcance)
        salida["excel_sesion"] = leer_excel(tmp / "sesion.xlsx", rutas, reemplazos)
        salida["pdf_sesion"] = leer_pdf(tmp / "sesion.pdf", rutas, reemplazos)
        salida["docx_sesion"] = leer_docx(tmp / "sesion.docx", rutas, reemplazos)
        salida["factor_legado"] = escenario_legado(tmp, guardado, reemplazos)
    finally:
        for ventana in (v, otra):
            if ventana is not None:
                with patch.object(ventana, "_decidir_cierre", return_value="descartar"):
                    ventana.close()
        APP.processEvents()
    salida["informe_proyeccion"] = informe_proyeccion(tmp)
    return json_estable(salida)


def sesion_legado(guardado: dict) -> dict:
    """Variante sintética anterior a la regla 4I-C: la primera actividad conserva F = I/I0 sin redondear
    (sin ``criterio_factor``) y sus importes R = (P − A)·(F − 1), Valor = (P − A) + R."""
    legado = deepcopy(guardado)
    r = legado["proyecciones_icociv"]["valores"][0]
    r.pop("criterio_factor", None); r.pop("factor_icociv_indices", None)
    r["factor_icociv"] = r["i_icociv"] / r["i0_icociv"]
    r["r_total"] = r["diferencia_absoluta"] = r["base_ajustable"] * (r["factor_icociv"] - 1)
    r["valor_proyectado"] = r["base_ajustable"] + r["r_total"]
    r["diferencia_porcentual"] = (r["factor_icociv"] - 1) * 100
    return legado


def escenario_legado(tmp: Path, guardado: dict, reemplazos) -> dict:
    from openpyxl import load_workbook as cargar
    from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
    from app_icociv.servicios.actualizacion_icociv import es_factor_legado
    rutas = (str(tmp),)
    w = WidgetProyeccionesICOCIV()
    try:
        w.restaurar_datos_sesion(deepcopy(sesion_legado(guardado)["proyecciones_icociv"]))
        antes = deepcopy(w.valores)
        w._generar_excel(str(tmp / "legado.xlsx"))
        hoja = cargar(tmp / "legado.xlsx")["Valores proyectados"]
        columnas, _ = w.plan_excel_valores(); n = {k: i for i, (k, _) in enumerate(columnas, 1)}
        formulas = [[hoja.cell(fila, n[k]).value for k in ("factor", "diferencia")] for fila in (3, 4, 5)]
        resultado = {"legados": [es_factor_legado(r) for r in w.valores], "aviso": normalizar_texto(w.lbl_legados.text()),
                     "aviso_visible": w.aviso_legados.isVisibleTo(w), "formulas_excel_filas_3_4_5": formulas,
                     "exportar_no_modifica": w.valores == antes}
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=False):
            resultado["cancelar_no_cambia"] = (not w.actualizar_factor_legado(antes[0]["calculo_id"])) and w.valores == antes
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=True),              patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")):
            resultado["actualizacion_aceptada"] = w.actualizar_factor_legado(antes[0]["calculo_id"])
        nuevo = w.valores[0]
        resultado["registro_actualizado"] = {k: nuevo.get(k) for k in ("criterio_factor", "factor_icociv", "diferencia_porcentual",
                                                                      "r_total", "valor_proyectado", "i0_icociv", "i_icociv")}
        w.pila_deshacer.undo()
        resultado["deshacer_restaura_legado"] = w.valores == antes
        return json_estable(_normalizar(resultado, rutas, reemplazos))
    finally:
        w.close()


def informe_proyeccion(tmp: Path) -> dict:
    """CSV reproducible, DOCX y PDF del informe de proyección (generador_reportes) para una serie sintética."""
    from app_icociv.proyeccion.servicio_proyeccion import ejecutar_proyeccion
    rutas = (str(tmp),)
    serie = gen.serie("tendencia_lineal_ruido")
    anio, mes = gen.objetivo(serie, 11)
    resultado = ejecutar_proyeccion(serie, anio, mes, 2020)
    csv = Path(gr.generar_csv_reproducibilidad(tmp / "reproducibilidad.csv", serie, resultado, None))
    comunes = dict(usuario="QA", archivo_excel="serie_sintetica_fase0", seleccion={}, parametros_proyeccion={},
                   ruta_jerarquica=None, fuente_label="Serie sintética Fase 0", fila=serie.head(1), serie_df=serie,
                   resultado_proyeccion=resultado, year_month=list(serie["Periodo"]),
                   configuracion=ConfiguracionInforme.desde_tipo("ejecutivo"))
    docx = gr.generar_reporte_proyeccion(tmp / "informe.docx", **comunes)
    pdf = gr.generar_reporte_pdf(tmp / "informe.pdf", **comunes)
    with csv.open(encoding="utf-8-sig", newline="") as archivo:
        filas = [[_numero_o_texto(normalizar_texto(x, rutas)) for x in fila] for fila in lector_csv(archivo)]
    return json_estable({"csv": filas, "docx": leer_docx(docx, rutas), "pdf": leer_pdf(pdf, rutas)})


def _numero_o_texto(valor: str):
    """Los números del CSV se comparan como flotantes (con tolerancia), no como texto con 17 cifras."""
    try:
        return float(valor) if valor.strip() and valor.strip().lower() not in ("nan", "inf", "-inf", "true", "false") else valor
    except ValueError:
        return valor
