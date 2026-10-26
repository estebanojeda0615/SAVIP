"""4E-A: dos pares tipificados, principales intactos y referencias por registro."""
import json
import math
from copy import deepcopy
from dataclasses import asdict
from unittest.mock import patch
from PySide6.QtCore import Qt,QTimer,QUrl
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QComboBox,QDialogButtonBox,QScrollArea
from openpyxl import load_workbook
from tests.test_importes_presupuestales_fase_4e import APP,reales,sintetico,cargar,registrar_seis,bloque,guardar,abrir
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV,COLUMNAS_VALORES,COLUMNAS_PROYECCIONES
from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV,COLUMNAS_OPERATIVAS
from app_icociv.utilidades.importes_presupuestales import desglose_presupuestal,texto_desglose
from app_icociv.servicios.actualizacion_icociv import calcular_actualizacion_icociv
from app_icociv.reportes.sesion_modulos import construir_proyecciones,construir_empalme
from app_icociv.persistencia.gestor_sesiones import cargar_sesion

def comprobar_modelo(modelo,fila,registro,valor,emp=False):
    d=desglose_presupuestal(registro,valor)
    # Fase 4G: en ICOCIV las columnas se localizan por clave (puede haber columnas de fecha auxiliares).
    for c,k in zip((7,8,13,14) if emp else [modelo.claves.index(k) for k in d],d):
        ix=modelo.index(fila,c)
        assert ix.data(Qt.UserRole)==d[k],(c,k,ix.data(Qt.UserRole),d[k])
        if d[k] is None:assert ix.data()=='—'
        else:assert isinstance(ix.data(Qt.UserRole),(int,float)) and '$' not in ix.data()

def editar(w,modo,aceptar=True):
    def actuar():
        d=QApplication.activeModalWidget();c=d.findChild(QComboBox);c.setCurrentIndex(c.findData(modo))
        QTest.mouseClick(d.findChild(QDialogButtonBox).button(QDialogButtonBox.Save if aceptar else QDialogButtonBox.Cancel),Qt.LeftButton)
    QTimer.singleShot(0,actuar);w._editar_interpretacion(QUrl('interpretacion'))

def test_01_principales_900_990_tres_modos_cantidades_y_extremos():
    r=calcular_actualizacion_icociv({'precio_base':1000.,'anticipo_amortizado':100.,'i0_icociv':100.,'i_icociv':110.,'cantidad':2.})
    assert r['base_ajustable']==900 and math.isclose(r['valor_proyectado'],990)
    for modo in ('heredado','unitario','total'):
        for q in (2.,2.5,1.,None,0.,-2.,float('nan'),float('inf'),1e308):
            registro={**r,'cantidad':q,'modalidad_importe':modo};antes=json.dumps(registro,sort_keys=True)
            d=desglose_presupuestal(registro,r['valor_proyectado'])
            assert json.dumps(registro,sort_keys=True)==antes
            if modo=='heredado':assert all(x is None for x in d.values());continue
            lado=('base_unitaria','valor_unitario') if modo=='unitario' else ('base_total','valor_total')
            assert (d[lado[0]],d[lado[1]])==(900,r['valor_proyectado'])
            otro=('base_total','valor_total') if modo=='unitario' else ('base_unitaria','valor_unitario')
            for k,principal in zip(otro,(900,r['valor_proyectado'])):
                if q is not None and math.isfinite(q) and q>0:
                    esperado=principal*q if modo=='unitario' else principal/q
                    assert d[k]==(esperado if math.isfinite(esperado) else None)
                else:assert d[k] is None
    for monto in (float('nan'),float('inf'),-900.,1e308,1e-308):
        d=desglose_presupuestal({'base_ajustable':monto,'cantidad':1e308,'modalidad_importe':'unitario'},monto)
        assert all(x is None or math.isfinite(x) for x in d.values())
    texto=texto_desglose(r,r['valor_proyectado'],'Original')
    assert 'Importe original sin clasificar' in texto and 'Base original: 900' in texto and 'Valor principal: 990' in texto

def test_02_tablas_17_18_roles_fuente_declarada_y_acero():
    w=sintetico();emp=WidgetEmpalmeICCPICOCIV()
    try:
        registrar_seis(emp);original=deepcopy(emp.calculos)
        assert len(COLUMNAS_VALORES)==17 and len(COLUMNAS_PROYECCIONES)==21 and len(COLUMNAS_OPERATIVAS)==18
        for modo in ('heredado','unitario','total'):
            w.valores[0]['modalidad_importe']=modo;w._refrescar_tabla_valores();r=w.valores[0]
            comprobar_modelo(w.modelo_valores,0,r,r['valor_proyectado'])
            for r in emp.calculos:r['modalidad_importe']=modo
            emp._actualizar_tabla()
            for i,r in enumerate(emp.calculos):comprobar_modelo(emp.modelo_operativo,i,r,r['valor_actualizado'],True)
            for modelo,fila,cols in ((w.modelo_valores,0,(8,9,13,14)),(emp.modelo_operativo,3,(7,8,13,14))):
                for posicion,c in enumerate(cols):
                    font=modelo.index(fila,c).data(Qt.FontRole)
                    declarado=modo!='heredado' and posicion%2==(0 if modo=='unitario' else 1)
                    assert bool(font and font.bold())==declarado
            assert all({k:v for k,v in r.items() if k!='modalidad_importe'}==old for r,old in zip(emp.calculos,original))
        acero=emp.calculos[3];assert acero['cantidad']!=acero['q'] and acero['anticipo_amortizado']==0
        d=desglose_presupuestal({**acero,'modalidad_importe':'unitario'},acero['valor_actualizado'])
        assert d['base_unitaria']==acero['p0'] and d['base_total']==acero['p0']*acero['cantidad']
        assert all(x is None for x in desglose_presupuestal({**acero,'p0':None},acero['valor_actualizado']).values())
    finally:w.close();emp.close()

def test_03_edicion_real_actualiza_pares_y_cancelar_conserva_principal():
    w=sintetico()
    try:
        w.tabla_valores.selectRow(0);w._detalle_valor();r=w.valores[0];original=deepcopy(r)
        with patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')):
            for modo in ('unitario','total','heredado'):
                editar(w,modo);assert r=={**original,'modalidad_importe':modo}
                comprobar_modelo(w.modelo_valores,0,r,r['valor_proyectado'])
                assert 'Base original: 900' in w.resumen_presupuestal.text() and 'Valor principal: 990' in w.resultado.toPlainText()
            previo=deepcopy(w.datos_sesion());editar(w,'total',False);assert w.datos_sesion()==previo
        assert 'Importe original sin clasificar' in w.resumen_comparacion_historica()
    finally:w.close()

def test_04_slider_tres_analisis_fechas_y_retorno_exactos():
    v=VentanaPrincipal();w=cargar(v)
    try:
        for r,m in zip(w.valores,('unitario','total','heredado')):r['modalidad_importe']=m
        v._activar_analisis(w.historial[0]['id']);antes=deepcopy(w.valores);hist=deepcopy(w.historial)
        with patch.object(v.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor')),patch.object(v.controlador,'cargar_archivo',side_effect=AssertionError('DANE')),patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')):
            assert w.seleccionar_referencia_historica('2024-02')
            assert [w.escenario_registro(r)[0] for r in w.valores][1:]==antes[1:]
            w.aplicar_global.setChecked(True);assert w.seleccionar_referencia_historica('2024-05')
            for i,r in enumerate(w.valores):
                e=w.escenario_registro(r)[0];comprobar_modelo(w.modelo_valores,i,r,e['valor_proyectado'])
                assert w.modelo_valores.index(i,10).data(Qt.UserRole)==e['i0_icociv']
                assert all(p in w.modelo_valores.index(i,11).data(Qt.ToolTipRole) for p in w.fechas_registro(r))
            assert 'Base original:' in v.lbl_comparacion_resultados.text()
            assert w.seleccionar_referencia_historica('2026-12') and [w.escenario_registro(r)[0] for r in w.valores]==antes
        assert w.valores==antes and w.historial==hist
    finally:v.close()

def test_05_excel_mixto_formulas_nulas_entradas_invalidas_y_sin_inyeccion(tmp_path):
    w=sintetico();emp=WidgetEmpalmeICCPICOCIV()
    try:
        r=deepcopy(w.valores[0]);w.valores=[{**r,'calculo_id':str(i),'modalidad_importe':m,'cantidad':q} for i,(m,q) in enumerate((('unitario',2.),('total',2.),('heredado',2.),('unitario',0.),('total',None)))]
        w.valores[0]['observacion_tecnica']='=1+1';w._refrescar_tabla_valores();antes=deepcopy(w.valores);ruta=tmp_path/'ico.xlsx';w._generar_excel(str(ruta));b=load_workbook(ruta)
        try:
            assert b.sheetnames==['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
            s=b['Valores proyectados'];t=b['Trazabilidad'];f=bloque(t)
            assert s.max_column==17 and b['Proyecciones ICOCIV'].max_column==21 and s.auto_filter.ref=='A2:Q7'
            # Fase 1A (AS-007): expectativa anterior a la regla 4I-C (factor I/I0 sin redondear). Contrato vigente: el Excel
            # escribe F = 1+ROUND((I/I0-1)*100,2)/100 para registros calculados con el criterio vigente (este lo crea el calculador).
            assert [s[c].value for c in ('I3','J3','M3','N3','O3','P3')]==['=G3-H3','=I3*F3','=1+ROUND((L3/K3-1)*100,2)/100','=I3+I3*(M3-1)','=N3*F3','=ROUND((M3-1)*100,2)']
            assert [s[c].value for c in ('I4','J4','N4','O4')]==['=J4/F4','=G4-H4','=O4/F4','=J4+J4*(M4-1)']
            assert all(s.cell(5,c).value is None for c in (9,10,14,15))
            assert s['J6'].value is s['O6'].value is s['I7'].value is s['N7'].value is None
            assert all(s[c].data_type=='f' for c in ('I6','N6','J7','O7'))
            assert t.cell(f+2,6).value==900 and t.cell(f+2,7).data_type=='f'
            assert all(t.cell(f+2,c).value is None for c in (8,9,10,11))
            assert s['Q3'].value=='=1+1' and s['Q3'].data_type=='s'
            dv=list(s.data_validations.dataValidation);assert len(dv)==1 and str(dv[0].sqref)=='F3:F7' and dv[0].errorStyle=='stop'
            assert not b._external_links and not any(c.comment for h in b for row in h for c in row)
        finally:b.close()
        assert w.valores==antes
        registrar_seis(emp)
        for r,m in zip(emp.calculos,('unitario','total','heredado','unitario','total','heredado')):r['modalidad_importe']=m
        emp._actualizar_tabla();antes=deepcopy(emp.datos_sesion());ruta=tmp_path/'emp.xlsx';emp._generar_excel_empalme(str(ruta));b=load_workbook(ruta)
        try:
            s=b['Resumen operativo Empalme'];t=b['Resultados registrados'];f=bloque(t)
            assert s.max_column==18 and len(b.sheetnames)==5 and s.protection.sheet
            assert [s[c].value for c in ('I3','O3','H4','N4')]==['=H3*G3','=N3*G3','=I4/G4','=O4/G4']
            assert all(s.cell(5,c).value is None for c in (8,9,14,15))
            for i,r in enumerate(emp.calculos,3):
                assert s.cell(i,18).value==r['calculo_id'] and not s.cell(i,7).protection.locked
                assert t.cell(f+i-3,6).value==r['base_ajustable'] and t.cell(f+i-3,7).data_type=='f'
            assert not any(c.comment for h in b for row in h for c in row)
        finally:b.close()
        assert emp.datos_sesion()==antes
    finally:w.close();emp.close()

def test_06_pdf_word_modelos_principales_y_cuatro_importes_sin_expandir_tecnica():
    w=sintetico();emp=WidgetEmpalmeICCPICOCIV()
    try:
        registrar_seis(emp)
        for modo in ('heredado','unitario','total'):
            w.valores[0]['modalidad_importe']=modo;emp.calculos[0]['modalidad_importe']=modo
            ico=asdict(construir_proyecciones(w));em=asdict(construir_empalme(emp.calculos,{},historial=emp.historial_empalme))
            seccion=next(s for s in ico['secciones'] if s['clave']=='valores');texto=json.dumps(seccion,ensure_ascii=False)
            assert [len(b['encabezados']) for b in seccion['bloques'] if 'encabezados' in b]==[7,7,3]
            for token in ('Base original: 900','Valor principal: 990','Base ajustable unitaria:','Base ajustable total:','Valor ajustado unitario:','Valor ajustado total:'):assert token in texto
            texto_emp=json.dumps(em,ensure_ascii=False,default=str)
            for token in ('Base original: 900','Valor principal: 945','Base ajustable unitaria:','Base ajustable total:'):assert token in texto_emp
            if modo=='heredado':assert 'Importe original sin clasificar' in texto and 'Importe original sin clasificar' in texto_emp
    finally:w.close();emp.close()

def test_07_sesion_sin_dane_pares_identicos_y_ausencia_metadata_legada(tmp_path):
    v=VentanaPrincipal();w=cargar(v);otro=None
    try:
        registrar_seis(v.widget_empalme)
        for r,m in zip(w.valores,('unitario','total','heredado')):
            if m!='heredado':r['modalidad_importe']=m
        for i,r in enumerate(v.widget_empalme.calculos):
            if i%3!=2:r['modalidad_importe']='unitario' if i%3==0 else 'total'
        ausente=tmp_path/'DANE_NO_EXISTE.xlsx'
        for a in w.historial:a['archivo_excel']=str(ausente)
        v.controlador.ruta_archivo=ausente;v._activar_analisis(w.historial[0]['id']);w.aplicar_global.setChecked(True);w.seleccionar_referencia_historica('2024-05')
        archivo=tmp_path/'mixta.json';guardar(v,archivo);d=cargar_sesion(archivo);otro=VentanaPrincipal()
        with patch.object(otro.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor')),patch.object(otro.controlador,'cargar_archivo',side_effect=AssertionError('DANE')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Servicio')):
            abrir(otro,archivo);n=otro.widget_proyecciones;e=otro.widget_empalme
            assert n.datos_sesion()==d['proyecciones_icociv'] and e.datos_sesion()==d['empalme_iccp_icociv']
            for i,r in enumerate(n.valores):comprobar_modelo(n.modelo_valores,i,r,n.escenario_registro(r)[0]['valor_proyectado'])
            for i,r in enumerate(e.calculos):comprobar_modelo(e.modelo_operativo,i,r,r['valor_actualizado'],True)
            assert 'modalidad_importe' not in n.valores[2] and 'modalidad_importe' not in e.calculos[2]
        assert not any(k in json.dumps(d) for k in ('base_unitaria','base_total','valor_unitario','valor_total'))
    finally:
        if otro:otro.close()
        v.close()

def test_08_geometria_copia_foco_y_seleccion_por_id():
    v=VentanaPrincipal();w=cargar(v);e=v.widget_empalme
    try:
        registrar_seis(e)
        for r,m in zip(w.valores,('unitario','total','heredado')):r['modalidad_importe']=m
        w._refrescar_tabla_valores();e.calculos[0]['modalidad_importe']='total';e._actualizar_tabla()
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for ancho,alto in ((1100,700),(1366,768),(1920,1080)):
                v.resize(ancho,alto);v.show();APP.processEvents()
                # Fase 4G: ICOCIV = 17 + columnas de fecha auxiliares que correspondan a las filas visibles.
                for seccion,widget,tabla,cantidad in ((1,w,w.tabla_valores,w.modelo_valores.columnCount()),(2,e,e.tabla_operativa,18)):
                    v._cambiar_seccion(seccion);APP.processEvents();APP.processEvents();tabla.selectRow(0);tabla.setFocus()
                    assert widget.findChild(QScrollArea).horizontalScrollBar().maximum()==0
                    QTest.keyClick(tabla,Qt.Key_C,Qt.ControlModifier);QTest.keyRelease(tabla,Qt.Key_Control)
                    assert len(APP.clipboard().text().split('\t'))==cantidad
                    assert tabla.hasFocus() and tabla.horizontalScrollBar().maximum()>0
        ident=e.calculos[0]['calculo_id'];e.calculos.reverse();e._actualizar_tabla()
        assert e.calculos[e._fila_seleccionada()]['calculo_id']==ident
        w.tabla_valores.clearSelection();assert 'Seleccione una sola fila' in w.resumen_presupuestal.text()
    finally:v.close()
