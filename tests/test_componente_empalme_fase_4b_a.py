"""Clasificación libre: paridad económica, consulta, edición y compatibilidad."""
import json
from copy import deepcopy
from unittest.mock import patch
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QMessageBox, QApplication, QGroupBox, QScrollArea
from PySide6.QtTest import QTest
from PySide6.QtGui import QFontDatabase,QFont
from openpyxl import load_workbook
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.persistencia.historial_empalme import validar_sesion_empalme, campos_componente
from app_icociv.servicios.empalme_iccp_icociv import calcular_empalme_iccp_icociv
from tests.test_historial_empalme_fase_4b import registrar_casos, guardar, abrir, ICCP, ROOT, PREVIO

COMPONENTES=[('Capítulo','Estructuras de concreto'),('Subclase de obra','Puentes y vías — Ñandú'),
             ('Actividad','Fundir losa'),('Insumo','Acero corrugado Ø ½'),('Otro','Frente norte')]

def registrar_componentes(w):
    finalizar=w._finalizar_calculo
    posiciones=iter(COMPONENTES)
    def con_componente(entrada,opcion):
        tipo,nombre=next(posiciones)
        w.tipo_componente.setCurrentText(tipo);w.nombre_componente.setText(nombre)
        finalizar({**entrada,'tipo_componente':tipo,'nombre_componente':nombre},opcion)
    with patch.object(w,'_finalizar_calculo',side_effect=con_componente):
        return registrar_casos(w)

def test_ui_rotulos_compatibles_tab_unicode_y_sin_inferencia():
    fuente_anterior=QApplication.instance().font()
    for nombre in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):assert QFontDatabase.addApplicationFont('C:/Windows/Fonts/'+nombre)>=0
    QApplication.instance().setFont(QFont('Segoe UI',10))
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        assert w.tipo_componente.currentData() is None and w.nombre_componente.text()==''
        nombres={'Grupo de obra':'Nombre del grupo','Subclase de obra':'Nombre de la subclase',
                 'Tipología':'Nombre de la tipología','Capítulo':'Nombre del capítulo','Actividad':'Nombre de la actividad',
                 'Ítem':'Nombre del ítem','Insumo':'Nombre del insumo','Otro':'Nombre del componente'}
        nombre='Frente Ñandú — Puente Ø ½ / 漢字 '+('nombre extenso '*30)
        w.nombre_componente.setText(nombre)
        for tipo,etiqueta in nombres.items():
            w.tipo_componente.setCurrentText(tipo)
            assert w.tipo_componente.currentData()==tipo and w.lbl_nombre_componente.text()==etiqueta
            assert w.nombre_componente.text()==nombre and w.nombre_componente.accessibleName()==etiqueta
        w.tipo_componente.setCurrentText('Grupo');assert w.tipo_componente.currentData()=='Grupo de obra'
        w.tipo_componente.setCurrentText('Subclase');assert w.tipo_componente.currentData()=='Subclase de obra'
        grupo=next(g for g in w.findChildren(QGroupBox) if g.title().startswith('3.'))
        grid=grupo.layout()
        assert grid.getItemPosition(grid.indexOf(w.tipo_componente))[:2]==(1,0)
        assert grid.getItemPosition(grid.indexOf(w.nombre_componente))[:2]==(1,1)
        assert grid.itemAtPosition(2,1).widget().isAncestorOf(w.item)
        v._cambiar_seccion(2);v.show();QApplication.processEvents()
        w.tipo_componente.setFocus();QTest.keyClick(w.tipo_componente,Qt.Key_Tab)
        assert w.nombre_componente.hasFocus()
        portapapeles=QApplication.clipboard();texto_anterior=portapapeles.text();portapapeles.setText(nombre)
        w.nombre_componente.selectAll();QTest.keyClick(w.nombre_componente,Qt.Key_V,Qt.ControlModifier)
        assert w.nombre_componente.text()==nombre
        portapapeles.setText(texto_anterior)
        QTest.keyClick(w.nombre_componente,Qt.Key_Tab);assert w.item.hasFocus()
        w.tipo_componente.setFocus();QTest.keyClick(w.tipo_componente,Qt.Key_Down)
        assert w.lbl_nombre_componente.text()=='Nombre de la tipología' and w.nombre_componente.text()==nombre
        QTest.keyClick(w.tipo_componente,Qt.Key_Up)
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for ancho,alto in ((1100,700),(1366,768),(1920,1080)):
                v.resize(ancho,alto);QApplication.processEvents()
                assert w.findChild(QScrollArea).horizontalScrollBar().maximum()==0,(tema,ancho,w.findChild(QScrollArea).horizontalScrollBar().maximum())
                assert w.tipo_componente.geometry().right()<w.nombre_componente.geometry().left()
        w._actualizar_jerarquia_icociv();assert w.tipo_componente.currentData()=='Subclase de obra'
        w.limpiar_formulario();assert w.tipo_componente.currentData() is None and not w.nombre_componente.text()
    finally:
        v.close();QApplication.instance().setFont(fuente_anterior)

def test_cinco_ejecuciones_paridad_y_dos_selectores_cero_servicios():
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        entradas=registrar_componentes(w)
        for r,(e,o),(tipo,nombre) in zip(w.calculos,entradas,COMPONENTES):
            real=calcular_empalme_iccp_icociv(e,o['indices'],ICCP)
            for k in real:
                if k!='fecha_calculo':assert r[k]==real[k],k
            assert (r['tipo_componente'],r['nombre_componente'])==(tipo,nombre)
        ico=deepcopy(v.widget_proyecciones.datos_sesion());before=deepcopy(w.datos_sesion())
        with patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor')) as motor,patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Economía')) as servicio:
            for a,(tipo,nombre) in zip(w.historial_empalme,COMPONENTES):
                w.selector_analisis_empalme.setCurrentIndex(w.selector_analisis_empalme.findData(a['id']))
                assert w.tipo_componente.currentData()==tipo and w.nombre_componente.text()==nombre
                v.selector_tipo_resultado.setCurrentIndex(1)
                v.selector_resultados_empalme.setCurrentIndex(v.selector_resultados_empalme.findData(a['id']))
                assert nombre in v.texto_resultados_empalme.toPlainText()
                assert all(f'{k}' in w.resultado.toPlainText() and val in w.resultado.toPlainText() for k,val in campos_componente(w.resultado_analisis_empalme(a)))
            assert motor.call_count==servicio.call_count==0
        assert w.calculos==before['calculos'] and w.historial_empalme==before['historial_empalme']
        assert v.widget_proyecciones.datos_sesion()==ico
    finally:v.close()

def test_editar_borrar_y_cancelar_clasificacion_original(tmp_path):
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        entradas=registrar_componentes(w);original=deepcopy(w.calculos[0]);a=w.historial_empalme[0]
        w.nombre_componente.setText('Borrador')
        with patch.object(QMessageBox,'question',return_value=QMessageBox.No):assert not w.activar_analisis_empalme(a['id'])
        assert w.nombre_componente.text()=='Borrador'
        w.tabla.selectRow(0);w.editar_seleccionado()
        assert w.nombre_componente.text()==original['nombre_componente']
        e,o=deepcopy(entradas[0]);e.update(tipo_componente='Subclase de obra',nombre_componente='Nueva subclase')
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):w._finalizar_calculo(e,o)
        actual=w.calculos[0]
        assert actual['calculo_id']==original['calculo_id'] and len(w.historial_empalme)==5
        assert actual['nombre_componente']=='Nueva subclase' and a['resultado_archivado']==original
        assert actual['valor_actualizado']==original['valor_actualizado']
        out=tmp_path/'editado.xlsx';w._generar_excel_empalme(str(out));b=load_workbook(out)
        assert b['Resultados registrados'].cell(3,16).value=='Subclase de obra' and b['Resultados registrados'].cell(3,17).value=='Nueva subclase'
        assert b['Historial de Empalme'].cell(3,14).value=='Capítulo' and b['Historial de Empalme'].cell(3,15).value==original['nombre_componente']
        b.close()
        from app_icociv.reportes.sesion_modulos import construir_empalme
        from app_icociv.reportes.modelo import Parrafo
        informe=construir_empalme(w.calculos,w._datos_generales(),historial=w.historial_empalme)
        textos='\n'.join(b.texto for s in informe.secciones for b in s.bloques if isinstance(b,Parrafo))
        assert 'Nombre de la subclase: Nueva subclase' in textos
        assert f"Análisis 1 · ID {a['id']} · Tipo de componente: Capítulo" in textos and original['nombre_componente'] in textos
        w.tabla.selectRow(0);w.eliminar_seleccionado();assert len(w.historial_empalme)==5
        assert w.activar_analisis_empalme(a['id'],confirmar=False)
        assert w.tipo_componente.currentData()=='Capítulo' and w.nombre_componente.text()==original['nombre_componente']
        w.validar_datos_sesion(w.datos_sesion())
    finally:v.close()

def test_sesion_mixta_y_metadata_corrupta_no_destructiva(tmp_path):
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        v.sesion_pendiente={'proyecciones_icociv':deepcopy(PREVIO['proyecciones'])};v._restaurar_sesion_pendiente()
        registrar_componentes(w);v.selector_tipo_resultado.setCurrentIndex(1)
        path=tmp_path/'mixta.json';guardar(v,path);data=cargar_sesion(path)
        otro=VentanaPrincipal()
        try:
            with patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.cargar_archivo',side_effect=AssertionError('Leer fuente')),patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Economía')):
                abrir(otro,path)
                assert otro.widget_empalme.datos_sesion()==data['empalme_iccp_icociv']
                assert otro.widget_proyecciones.datos_sesion()==data['proyecciones_icociv']
                guardar(otro,tmp_path/'reabierta.json');reabierta=cargar_sesion(tmp_path/'reabierta.json')
                for k in ('empalme_iccp_icociv','proyecciones_icociv'):assert reabierta[k]==data[k]
            before=deepcopy(otro.widget_empalme.datos_sesion());ico=deepcopy(otro.widget_proyecciones.datos_sesion())
            for campo,valor in (('tipo_componente','Inventado'),('tipo_componente',[]),('nombre_componente',None),('nombre_componente',3),('analisis_empalme_id','huérfano')):
                bad=deepcopy(data);bad['empalme_iccp_icociv']['calculos'][0][campo]=valor
                p=tmp_path/'corrupta.json';p.write_text(json.dumps(bad),encoding='utf-8')
                with patch.object(QMessageBox,'warning') as warning:abrir(otro,p);assert warning.call_count==1
                assert otro.widget_empalme.datos_sesion()==before and otro.widget_proyecciones.datos_sesion()==ico
            bad=deepcopy(before);bad['historial_empalme'][0]['resultado_archivado']=deepcopy(bad['calculos'][0]);bad['historial_empalme'][0]['resultado_archivado']['nombre_componente']={}
            try:validar_sesion_empalme(bad)
            except ValueError:pass
            else:raise AssertionError('Archivo corrupto aceptado')
        finally:otro.close()
    finally:v.close()

def test_sesion_anterior_neutra_edicion_y_exportacion(tmp_path):
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        datos=json.loads((ROOT/'docs/fase_4b/FIXTURE_DATOS.json').read_text(encoding='utf-8'))
        p=tmp_path/'anterior.json';p.write_text(json.dumps(datos),encoding='utf-8');abrir(v,p)
        assert w.datos_sesion()==datos['empalme_iccp_icociv'] and w.tipo_componente.currentData() is None
        assert w.nombre_componente.text()==''
        for a in w.historial_empalme:
            w.activar_analisis_empalme(a['id'],confirmar=False)
            assert w.tipo_componente.currentData() is None and not w.nombre_componente.text()
        w.activar_analisis_empalme(datos['empalme_iccp_icociv']['analisis_empalme_activo_id'],confirmar=False)
        w.tabla.selectRow(0);w.editar_seleccionado();assert w.tipo_componente.currentData() is None
        out=tmp_path/'anterior.xlsx';w._generar_excel_empalme(str(out));b=load_workbook(out)
        assert b['Resultados registrados'].max_column==17
        assert all(b['Resultados registrados'].cell(r,16).value is None and b['Resultados registrados'].cell(r,17).value is None for r in range(3,8))
        b.close();guardar(v,tmp_path/'guardada.json')
        # Fase 4I-I: la edición abierta de la fila viaja como borrador; el resto de la sesión queda idéntico.
        guardada=cargar_sesion(tmp_path/'guardada.json')['empalme_iccp_icociv'];borrador=guardada.pop('borrador_formulario')
        assert guardada==datos['empalme_iccp_icociv'] and borrador['calculo_id_en_edicion']==w.calculos[0]['calculo_id']
    finally:v.close()

def test_entrada_real_y_captura_metadata_antes_de_callback():
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('unitario'));w.cantidad.setValue(2)
        w.item.setText('Concreto de 3.000 psi');w.codigo.setText('C-001');w.precio_base.setValue(1000)
        w._establecer_combo_por_data(w.unidad,'m3')
        w.tipo_componente.setCurrentText('Capítulo');w.nombre_componente.setText('Estructuras de concreto')
        w._establecer_periodo(w.fecha_inicial_anio,w.fecha_inicial_mes,'2021_1')
        w._establecer_periodo(w.fecha_final_anio,w.fecha_final_mes,'2021_6')
        w._establecer_combo_por_data(w.combo_tipo_iccp,'total_iccp');w._actualizar_series_iccp()
        w._establecer_combo_por_data(w.combo_serie_iccp,'Total ICCP')
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)) as servicio,patch.object(QMessageBox,'warning') as warning:
            w.calcular();assert not warning.called and servicio.call_count==1
        assert 'tipo_componente' not in servicio.call_args.args[0] and 'nombre_componente' not in servicio.call_args.args[0]
        assert w.calculos[0]['codigo_item']=='C-001' and w.calculos[0]['nombre_componente']=='Estructuras de concreto'
        w._establecer_periodo(w.fecha_final_anio,w.fecha_final_mes,'2022_9')
        opcion={'indices':{'2021_12':100.,'2022_6':150.},'ruta':'Ruta sintética','seleccion':{},'ruta_estructurada':[]}
        futura={**opcion,'indices':{**opcion['indices'],'2022_9':180.}}
        pendientes=[];w.callback_proyeccion_icociv=lambda seleccion,anio,mes,fn:pendientes.append(fn)
        with patch.object(w,'_opcion_icociv_o_vacia',return_value=opcion),patch.object(w,'_aplicar_resultado_proyeccion',return_value=futura),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):
            w.calcular();assert len(pendientes)==1
            w.tipo_componente.setCurrentText('Otro');w.nombre_componente.setText('Texto posterior')
            pendientes[0]({},None)
        assert w.calculos[-1]['tipo_componente']=='Capítulo' and w.calculos[-1]['nombre_componente']=='Estructuras de concreto'
        w.limpiar_formulario();assert w.tipo_componente.currentData() is None
    finally:v.close()
