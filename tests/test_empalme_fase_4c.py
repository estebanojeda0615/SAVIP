"""Vistas derivadas, selección por ID y exportación sin mutar los registros."""
import json,time,math
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch
from PySide6.QtCore import Qt,QSettings
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog,QPushButton,QScrollArea
from PySide6.QtGui import QFontDatabase,QFont
from openpyxl import load_workbook
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets import empalme_iccp_icociv as modulo
from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV,COLUMNAS_OPERATIVAS,_fila_operativa
from app_icociv.servicios.empalme_iccp_icociv import calcular_empalme_iccp_icociv
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.reportes.sesion_modulos import construir_empalme
from tests.test_componente_empalme_fase_4b_a import registrar_componentes
from tests.test_historial_empalme_fase_4b import ICCP,PREVIO,ROOT,guardar,abrir
APP=QApplication.instance() or QApplication([])

def registrar_seis(w):
    finalizar=w._finalizar_calculo
    entradas=[]
    def nuevo(e,o):
        numero=len(w.calculos)+1;w.contrato.setText(f'Proyecto guardado {numero}')
        entrada={**e,'codigo_item':f'C-{numero:03d}'}
        finalizar(entrada,o);entradas.append((entrada,o))
    with patch.object(w,'_finalizar_calculo',side_effect=nuevo):registrar_componentes(w)
    e,o=deepcopy(entradas[1]);e.update(item='Decremento',codigo_item='D-006',tipo_componente='Grupo de obra',nombre_componente='Frente Ñandú')
    o['indices']={'2022_1':150.,'2022_6':120.}
    w.contrato.setText('Proyecto guardado 6');w.tipo_componente.setCurrentText('Grupo');w.nombre_componente.setText('Frente Ñandú')
    with patch.object(modulo,'calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):finalizar(e,o)
    w.contrato.setText('Proyecto actual posterior')
    assert len(w.calculos)==len(w.historial_empalme)==6
    return entradas+[(e,o)]

def confirmar_todo(w,texto):
    with patch.object(QMessageBox,'exec',lambda box:0),patch.object(QMessageBox,'clickedButton',lambda box:next(b for b in box.buttons() if b.text()==texto)):w.eliminar_todo()

def test_seis_calculos_reales_columnas_valores_y_proyecto_original():
    w=WidgetEmpalmeICCPICOCIV()
    try:
        entradas=registrar_seis(w);original=deepcopy(w.datos_sesion())
        assert w.vistas_tablas.currentIndex()==0 and w.modelo_operativo.rowCount()==6 and w.modelo_operativo.columnCount()==18
        assert len(w.barra.findChildren(QPushButton))==4 and len(w.barra.tablas)==3
        for i,(r,(e,o)) in enumerate(zip(w.calculos,entradas)):
            real=calcular_empalme_iccp_icociv(e,o['indices'],ICCP)
            for k in real:
                if k!='fecha_calculo':assert r[k]==real[k],k
            m=w.modelo_operativo
            assert m.index(i,0).data() == str(r['numero_calculo']) and type(m.index(i,0).data(Qt.UserRole)) is int
            assert m.index(i,1).data()==f'Proyecto guardado {i+1}' and m.index(i,17).data(Qt.UserRole+1)==r['calculo_id']
            for c,k in ((6,'cantidad'),(12,'r_total')):
                assert m.index(i,c).data(Qt.UserRole)==r[k]
                assert '$' not in m.index(i,c).data() and '%' not in m.index(i,c).data()
            assert all(m.index(i,c).data(Qt.UserRole) is None for c in (7,8,13,14))
            w.tabla_operativa.selectRow(i);assert 'Base original:' in w.contexto_registro.text() and 'Valor principal:' in w.contexto_registro.text()
            assert m.index(i,15).data(Qt.UserRole)==(r['r_total']/r['base_ajustable'])*100
            assert m.index(i,9).data()==_fila_operativa(r)['fecha_inicial']
        assert 'Acero' in w.modelo_operativo.index(3,11).data() and w.calculos[3]['anticipo_amortizado']==0
        assert w.modelo_operativo.index(5,15).data(Qt.UserRole)<0
        assert _fila_operativa({'base_ajustable':0,'r_total':10})['variacion_valor'] is None
        assert w.datos_sesion()==original
    finally:w.close()

def test_seleccion_por_id_celda_copiable_cambio_vista_y_reordenacion():
    w=WidgetEmpalmeICCPICOCIV()
    try:
        registrar_seis(w);original=deepcopy(w.datos_sesion())
        with patch.object(modulo,'calcular_empalme_iccp_icociv',side_effect=AssertionError('Recalculo')) as servicio,patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor')) as motor:
            for i,r in enumerate(w.calculos):
                w.tabla_operativa.selectRow(i);clave=r['calculo_id']
                for vista in (1,0,1):
                    w.vistas_tablas.setCurrentIndex(vista)
                    assert w.calculos[w._fila_seleccionada()]['calculo_id']==clave
                    assert all({x.row() for x in t.selectedIndexes()}=={i} for t,_ in w.barra.tablas)
                assert r['nombre_componente'] in w.contexto_registro.text()
            assert motor.call_count==servicio.call_count==0
        w.tabla_operativa.selectRow(2);objetivo=w.calculos[2]
        w.calculos.reverse()
        assert w.calculos[w._fila_seleccionada()] is objetivo
        with patch.object(w,'_dialogo_detalle') as detalle:w.boton_detalle.click();assert detalle.call_args.args[0] is objetivo
        w.calculos.reverse();assert w.datos_sesion()==original
        cell=w.modelo_valor_ajustado.index(1,7);w.tabla.setCurrentIndex(cell)
        assert len(w.tabla.selectedIndexes())==1
        w.tabla.clearSelection();assert not w.boton_detalle.isEnabled() and not w.boton_eliminar.isEnabled()
    finally:w.close()

def test_edicion_borrado_snapshot_y_borrar_todo_confirmado():
    w=WidgetEmpalmeICCPICOCIV()
    try:
        entradas=registrar_seis(w);r=deepcopy(w.calculos[0]);h=deepcopy(w.historial_empalme)
        w.tabla_operativa.selectRow(0);w.accion_editar.trigger()
        assert w.indice_edicion==0 and w.nombre_componente.text()==r['nombre_componente']
        w.nombre_componente.setText('Borrador')
        with patch.object(QMessageBox,'question',return_value=QMessageBox.No):assert not w.activar_analisis_empalme(w.historial_empalme[1]['id'])
        assert w.nombre_componente.text()=='Borrador'
        e,o=deepcopy(entradas[0]);e.update(precio_base=1200.,nombre_componente='Componente editado',tipo_componente='Actividad')
        with patch.object(modulo,'calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):w._finalizar_calculo(e,o)
        assert w.calculos[0]['calculo_id']==r['calculo_id'] and len(w.historial_empalme)==6
        assert w.historial_empalme[0]['resultado_archivado']==r and w.modelo_operativo.index(0,13).data(Qt.UserRole) is None and 'Valor principal:' in w.contexto_registro.text()
        w.vistas_tablas.setCurrentIndex(1);assert w.calculos[w._fila_seleccionada()]['calculo_id']==r['calculo_id']
        w.boton_eliminar.click();assert w.modelo_operativo.rowCount()==w.modelo_equivalencias.rowCount()==w.modelo_valor_ajustado.rowCount()==5
        assert not w.boton_detalle.isEnabled() and not w.boton_eliminar.isEnabled()
        assert w.activar_analisis_empalme(h[0]['id'],confirmar=False) and w.nombre_componente.text()==r['nombre_componente']
        antes=deepcopy(w.datos_sesion());confirmar_todo(w,'Cancelar');assert w.datos_sesion()==antes
        confirmar_todo(w,'Eliminar');assert not w.calculos and len(w.historial_empalme)==6
        assert w.modelo_operativo.rowCount()==0 and all('resultado_archivado' in a for a in w.historial_empalme)
        w.validar_datos_sesion(w.datos_sesion())
        informe=construir_empalme([w.resultado_analisis_empalme(a) for a in w.historial_empalme],w._datos_generales(),historial=w.historial_empalme,operativos=w.calculos)
        assert not any(s.clave=='resumen_operativo_empalme' for s in informe.secciones)
    finally:w.close()

def test_sesion_mixta_sin_fuente_no_recalcula_y_compatibilidad(tmp_path):
    v=VentanaPrincipal()
    try:
        v.sesion_pendiente={'proyecciones_icociv':deepcopy(PREVIO['proyecciones'])};v._restaurar_sesion_pendiente()
        registrar_seis(v.widget_empalme);ico=deepcopy(v.widget_proyecciones.datos_sesion())
        v.widget_empalme.tabla_operativa.selectRow(4);path=tmp_path/'mixta.json';guardar(v,path);data=cargar_sesion(path)
        assert not any(k in data['empalme_iccp_icociv'] for k in ('vista_operativa','vistas_tablas','filas_operativas'))
        otro=VentanaPrincipal()
        try:
            with patch.object(modulo,'calcular_empalme_iccp_icociv',side_effect=AssertionError('Economico')) as servicio,patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor')) as motor:
                abrir(otro,path);w=otro.widget_empalme
                assert w.datos_sesion()==data['empalme_iccp_icociv'] and otro.widget_proyecciones.datos_sesion()==ico
                assert w.modelo_operativo.rowCount()==6 and len(otro.widget_proyecciones.historial)==3
                for a in w.historial_empalme:
                    assert w.activar_analisis_empalme(a['id'],confirmar=False)
                    otro.selector_tipo_resultado.setCurrentIndex(1)
                    otro.selector_resultados_empalme.setCurrentIndex(otro.selector_resultados_empalme.findData(a['id']))
                assert motor.call_count==servicio.call_count==0
                assert otro.widget_proyecciones.datos_sesion()==ico
            guardada=tmp_path/'otra.json';guardar(otro,guardada);nueva=cargar_sesion(guardada)
            assert nueva['empalme_iccp_icociv']['calculos']==data['empalme_iccp_icociv']['calculos']
            assert nueva['empalme_iccp_icociv']['historial_empalme']==data['empalme_iccp_icociv']['historial_empalme']
            legacy=deepcopy(PREVIO['empalme'])
            for r in legacy:r.pop('calculo_id',None);r.pop('tipo_componente',None);r.pop('nombre_componente',None)
            bloque={'calculos':legacy,'generales':{'contrato':'Actual'}};w.restaurar_datos_sesion(bloque)
            w.tabla_operativa.selectRow(2);w.vistas_tablas.setCurrentIndex(1)
            assert w.calculos[w._fila_seleccionada()] is legacy[2] or w.calculos[w._fila_seleccionada()]==legacy[2]
            assert w.modelo_operativo.index(2,2).data()=='' and w.calculos==legacy
            assert all('calculo_id' not in r for r in w.datos_sesion()['calculos'])
        finally:otro.close()
    finally:v.close()

def test_excel_resumen_vinculos_congelados_y_heredados(tmp_path):
    w=WidgetEmpalmeICCPICOCIV()
    try:
        registrar_seis(w)
        for r in w.calculos:r['modalidad_importe']='unitario'  # Fuente declarada; legado sin modo se prueba en las otras vistas.
        w._actualizar_tabla();antes=deepcopy(w.datos_sesion());w.tabla_operativa.selectRow(1)
        path=tmp_path/'seis.xlsx';w._generar_excel_empalme(str(path));b=load_workbook(path)
        assert b.sheetnames==['Empalme ICCP-ICOCIV','Resultados registrados','Cálculos editables','Historial de Empalme','Resumen operativo Empalme']
        resumen=b['Resumen operativo Empalme'];assert resumen.protection.sheet and resumen.max_row==8 and resumen.max_column==18
        assert [c.value for c in resumen[2]]==[h for _,h in COLUMNAS_OPERATIVAS]
        for i,r in enumerate(w.calculos):
            fila=3+i;origen=w._fila_datos_calculo+i
            assert resumen.cell(fila,18).value==r['calculo_id'] and resumen.cell(fila,2).value==r['contrato']
            assert resumen.cell(fila,14).value==f"='Empalme ICCP-ICOCIV'!K{origen}"
            assert resumen.cell(fila,16).value==f"=M{fila}/('Empalme ICCP-ICOCIV'!C{origen}-'Empalme ICCP-ICOCIV'!L{origen})*100"
            assert all(resumen.cell(fila,c).data_type=='f' and 'IF(' not in resumen.cell(fila,c).value for c in (8,9,13,14,15,16))
            # La serialización XLSX conserva 15 cifras; el registro/JSON se compara sin tolerancia.
            assert math.isclose(b['Resultados registrados'].cell(fila,11).value,r['valor_actualizado'],rel_tol=1e-14,abs_tol=1e-10)
        assert not b._external_links and all(c.comment is None for row in resumen for c in row);b.close()
        assert w.datos_sesion()==antes
        legacy=deepcopy(w.calculos[0]);legacy.pop('calculo_id');legacy['nombre_componente']='=1+1'
        w.calculos=[legacy];w.historial_empalme=[];w._actualizar_tabla()
        path=tmp_path/'legacy.xlsx';w._generar_excel_empalme(str(path));b=load_workbook(path)
        s=b['Resumen operativo Empalme'];assert s['N3'].data_type=='n' and s['N3'].value==legacy['valor_actualizado']
        assert s['C3'].value=='=1+1' and s['C3'].data_type=='s';b.close()
    finally:w.close()

def test_rendimiento_1_10_100_1000_scroll_y_temas():
    for nombre in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):QFontDatabase.addApplicationFont('C:/Windows/Fonts/'+nombre)
    APP.setFont(QFont('Segoe UI',10));v=VentanaPrincipal();w=v.widget_empalme
    try:
        registrar_seis(w);base=deepcopy(w.calculos[0]);base['observacion_tecnica']='Texto extenso — Ñandú '*500
        w.historial_empalme=[];base.pop('analisis_empalme_id',None)
        for cantidad in (1,10,100,1000):
            w.calculos=[{**deepcopy(base),'calculo_id':f'carga-{i}','numero_calculo':i+1} for i in range(cantidad)]
            original=deepcopy(w.calculos);inicio=time.perf_counter();w._actualizar_tabla();duracion=time.perf_counter()-inicio
            assert duracion<10,(cantidad,duracion)
            assert w.modelo_operativo.rowCount()==w.modelo_equivalencias.rowCount()==w.modelo_valor_ajustado.rowCount()==cantidad
            w.tabla_operativa.selectRow(cantidad-1);w.vistas_tablas.setCurrentIndex(1)
            assert w.calculos[w._fila_seleccionada()]['calculo_id']==f'carga-{cantidad-1}' and w.calculos==original
        v._cambiar_seccion(2);v.show()
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for ancho,alto in ((1100,700),(1366,768),(1920,1080)):
                v.resize(ancho,alto);APP.processEvents()
                for vista in (0,1):
                    w.vistas_tablas.setCurrentIndex(vista);APP.processEvents()
                    assert w.findChild(QScrollArea).horizontalScrollBar().maximum()==0
                assert w.tabla_operativa.horizontalScrollBar().maximum()>0 and w.tabla_operativa.rowHeight(0)<60
                fuente=w.tabla_operativa.horizontalHeader().fontMetrics()
                for col,(_,titulo) in enumerate(COLUMNAS_OPERATIVAS):assert w.tabla_operativa.columnWidth(col)>=fuente.horizontalAdvance(titulo)+12
                assert w.barra.minimumSizeHint().width()<w.findChild(QScrollArea).viewport().width()
    finally:v.close()

def test_periodo_historico_largo_y_fuente_sintetica_disponible(tmp_path):
    from tests.test_proyecciones_fase_3 import anexo_sintetico
    path=tmp_path/'anexo_sintetico.xlsx';anexo_sintetico(path)
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        v.controlador.cargar_archivo(path);w.actualizar_icociv(v.controlador.tablas,v.controlador.periodos,path)
        registrar_seis(w)
        with patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.cargar_archivo',side_effect=AssertionError('Lectura al consultar')),patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor')):
            for a in w.historial_empalme:assert w.activar_analisis_empalme(a['id'],confirmar=False)
        w.limpiar_formulario();w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('unitario'));w.cantidad.setValue(2);w.item.setText('ICCP histórico largo');w.precio_base.setValue(1000);w.anticipo.setValue(100)
        w._establecer_combo_por_data(w.unidad,'m3');w._establecer_combo_por_data(w.combo_tipo_iccp,'total_iccp');w._actualizar_series_iccp();w._establecer_combo_por_data(w.combo_serie_iccp,'Total ICCP')
        w._establecer_periodo(w.fecha_inicial_anio,w.fecha_inicial_mes,'2000_1');w._establecer_periodo(w.fecha_final_anio,w.fecha_final_mes,'2021_12')
        datos={'Total ICCP':{'2000_1':10.,'2021_12':110.}}
        with patch.object(modulo,'calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,datos)),patch.object(QMessageBox,'warning') as aviso:w.calcular();assert not aviso.called
        assert len(w.calculos)==7 and w.calculos[-1]['valor_actualizado']==9900
        assert w.modelo_operativo.index(6,9).data()=='2000-01' and w.modelo_operativo.index(6,10).data()=='2021-12'
    finally:v.close()
