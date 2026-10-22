"""4E-B: entradas diferentes reales, pares simultáneos y guardas de registro."""
import json,math
from copy import deepcopy
from contextlib import ExitStack
from unittest.mock import patch
from PySide6.QtCore import Qt,QTimer
from PySide6.QtWidgets import QMessageBox
from openpyxl import load_workbook
from docs.fase_4e_b_unificada.escenarios import APP,contexto,entrada,ejecutar_emp,fila,VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV
from app_icociv.utilidades.importes_presupuestales import desglose_presupuestal,texto_desglose
from app_icociv.servicios.actualizacion_icociv import calcular_actualizacion_icociv
from tests.test_historial_empalme_fase_4b import guardar,abrir

def near(a,b):
    # §1: como máximo 16 ULP por operaciones IEEE-754, nunca redondear originales.
    assert abs(a-b)<=16*max(math.ulp(a),math.ulp(b)),(a,b,a-b)

def crear(w,emp,modo,p=100000.,a=0.,q=100.,indice=110.):
    entrada(w,modo,p,a,q,emp)
    n=len(w.calculos if emp else w.valores)
    with patch.object(QMessageBox,'warning',side_effect=AssertionError('Aviso inesperado')):
        if emp:assert ejecutar_emp(w,indice)==1
        else:
            with patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',wraps=calcular_actualizacion_icociv) as s:w.boton_calcular.click();assert s.call_count==1
    assert len(w.calculos if emp else w.valores)==n+1
    r=(w.calculos if emp else w.valores)[-1];d=fila(w,emp)['UserRole']
    assert all(type(x) in (float,int) for x in d)
    assert r['precio_base']==p and r['anticipo_amortizado']==a and r['cantidad']==q
    b=p-a;v=r['valor_actualizado' if emp else 'valor_proyectado']
    assert d==([b,b*q,v,v*q] if modo=='unitario' else [b/q,b,v/q,v])
    tabla=w.tabla_operativa if emp else w.tabla_valores
    assert tabla.selectionModel().selectedRows()[0].row()==n
    return deepcopy(r),d

def test_A_B_C_entradas_equivalentes_anticipo_fraccion_decremento_contraste():
    for emp in (False,True):
        for q,a,indice in ((100.,0.,110.),(100.,10000.,110.),(2.5,10000.,90.)):
            w=WidgetEmpalmeICCPICOCIV() if emp else WidgetProyeccionesICOCIV()
            try:
                if not emp:contexto(w,indice)
                u,du=crear(w,emp,'unitario',a=a,q=q,indice=indice)
                t,dt=crear(w,emp,'total',p=100000.*q,a=a*q,q=q,indice=indice)
                assert u['calculo_id']!=t['calculo_id'] and u['precio_base']!=t['precio_base']
                if indice<100:assert u['r_total']<0 and t['r_total']<0
                for x,y in zip(du,dt):near(x,y)
                for x,y in zip(du,[(100000.-a),(100000.-a)*q,(100000.-a)*indice/100.,(100000.-a)*indice/100.*q]):near(x,y)
                _,distinto=crear(w,emp,'total',a=a,q=q,indice=indice)
                assert distinto!=du
            finally:w.close()

def test_D_dialogo_real_cancelar_elegir_un_solo_registro_sin_inferencia():
    for emp in (False,True):
        w=WidgetEmpalmeICCPICOCIV() if emp else WidgetProyeccionesICOCIV()
        try:
            if not emp:contexto(w)
            for decision in ('Cancelar','Unitario','Total'):
                entrada(w,'heredado',emp=emp);previo=deepcopy(w.datos_sesion());n=len(w.calculos if emp else w.valores)
                def actuar():
                    d=APP.activeModalWidget();assert isinstance(d,QMessageBox)
                    assert 'P y A' in d.text() and not w.boton_calcular.isEnabled()
                    # Reentrada programática durante el modal tampoco registra dos veces.
                    (w.calcular if emp else w._calcular)()
                    next(b for b in d.buttons() if b.text()==decision).click()
                QTimer.singleShot(0,actuar)
                if emp:llamadas=ejecutar_emp(w)
                else:
                    with patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',wraps=calcular_actualizacion_icociv) as s:w.boton_calcular.click()
                    llamadas=s.call_count
                assert llamadas==int(decision!='Cancelar')
                assert len(w.calculos if emp else w.valores)==n+int(decision!='Cancelar')
                if decision=='Cancelar':assert w.datos_sesion()==previo
                else:assert all(x is not None for x in fila(w,emp)['UserRole'])
                assert w.precio_base.value()==100000 and w.anticipo.value()==0
        finally:w.close()

def test_E_cantidad_invalida_nueva_no_registra_ni_servicio():
    for emp in (False,True):
        w=WidgetEmpalmeICCPICOCIV() if emp else WidgetProyeccionesICOCIV()
        try:
            if not emp:contexto(w)
            for q in (0.,-1.,float('nan'),float('inf')):
                entrada(w,'unitario',emp=emp);previo=deepcopy(w.datos_sesion())
                with patch.object(w.cantidad,'value',return_value=q),patch.object(QMessageBox,'warning') as aviso,patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')):
                    if emp:assert ejecutar_emp(w)==0
                    else:w.boton_calcular.click()
                assert aviso.call_count==1 and 'mayor que cero' in aviso.call_args.args[2]
                assert w.datos_sesion()==previo
        finally:w.close()

def test_E_legado_clasificacion_por_id_y_cantidad_invalida():
    from tests.test_base_valor_presupuestal_fase_4e_a import editar
    for emp in (False,True):
        w=WidgetEmpalmeICCPICOCIV() if emp else WidgetProyeccionesICOCIV()
        try:
            if not emp:contexto(w)
            crear(w,emp,'unitario');r=(w.calculos if emp else w.valores)[0];r.pop('modalidad_importe');original=deepcopy(r)
            if emp:w._actualizar_tabla();w.tabla_operativa.selectRow(0)
            else:w._refrescar_tabla_valores();w.tabla_valores.selectRow(0);w._detalle_valor()
            assert 'Importe original sin clasificar' in texto_desglose(r,r['valor_actualizado' if emp else 'valor_proyectado'],'Original')
            with patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Servicio')):
                if emp:w.editar_seleccionado();w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('total'));w.boton_calcular.click()
                else:editar(w,'total')
            r=(w.calculos if emp else w.valores)[0]
            assert r=={**original,'modalidad_importe':'total'} and all(x is not None for x in fila(w,emp)['UserRole'])
            if emp:assert w.historial_empalme[0]['resultado_archivado']==original
            for q in (0.,-1.,None,float('nan'),float('inf'),1e-308):
                old={**r,'cantidad':q};antes=json.dumps(old,sort_keys=True)
                d=desglose_presupuestal(old,old['valor_actualizado' if emp else 'valor_proyectado'])
                assert d['base_total']==r['base_ajustable'] and d['valor_total']==r['valor_actualizado' if emp else 'valor_proyectado']
                assert d['base_unitaria'] is d['valor_unitario'] is None
                assert json.dumps(old,sort_keys=True)==antes and 'No disponible' in texto_desglose(old,d['valor_total'],'Original')
        finally:w.close()

def test_F_G_series_propias_slider_originales_y_equivalencia():
    w=WidgetProyeccionesICOCIV()
    try:
        contexto(w);crear(w,False,'unitario');crear(w,False,'total',p=10000000.)
        contexto(w,indice=260.,base=200.);crear(w,False,'total',p=10000000.)
        originales=deepcopy(w.valores);hist=deepcopy(w.historial)
        ids=[r['analisis_id'] for r in w.valores];assert ids[0]==ids[1]!=ids[2]
        assert w.seleccionar_referencia_historica('2024-02')
        assert w.escenario_registro(w.valores[0])[0]==originales[0]
        assert w.escenario_registro(w.valores[2])[0]['i0_icociv']==180.
        w.aplicar_global.setChecked(True);assert w.seleccionar_referencia_historica('2024-05')
        ds=[]
        for i,r in enumerate(w.valores):
            efectivo=w.escenario_registro(r)[0];d=desglose_presupuestal(r,efectivo['valor_proyectado']);ds.append(list(d.values()))
            assert efectivo['i0_icociv']==(190. if i==2 else 95.)
            assert fila(w,i=i)['UserRole']==list(d.values())
        for a,b in zip(ds[0],ds[1]):near(a,b)
        assert w.valores==originales and w.historial==hist
        assert w.seleccionar_referencia_historica('2026-12')
        assert [w.escenario_registro(r)[0] for r in w.valores]==originales
    finally:w.close()

def test_H_empalme_tres_escenarios_acero_Q_distinta_q():
    for inicio,fin,caso in ((2021,2021,'solo_iccp'),(2022,2022,'solo_icociv'),(2021,2022,'empalme_completo')):
        w=WidgetEmpalmeICCPICOCIV()
        try:
            for modo in ('unitario','total'):
                entrada(w,modo,emp=True);w.fecha_inicial_anio.setValue(inicio);w.fecha_final_anio.setValue(fin)
                assert ejecutar_emp(w)==1;r=deepcopy(w.calculos[-1]);assert r['caso']==caso
                assert all(x is not None for x in fila(w,True)['UserRole'])
                assert r['valor_actualizado']==r['base_ajustable']+r['r_total']
            entrada(w,'unitario',emp=True);w.chk_acero.setChecked(True);w.p0.setValue(100000.);w.ix.setValue(1000.);w.q.setValue(200.)
            assert ejecutar_emp(w)==1;r=deepcopy(w.calculos[-1]);assert r['cantidad']==100 and r['q']==200 and r['anticipo_amortizado']==0
            assert r['z']==r['ix']*r['q']-r['valor_actualizado']
            assert fila(w,True)['UserRole'][3]==r['valor_actualizado']*100
        finally:w.close()

def test_J_K_exportadores_pares_sesion_mixta_exacta_sin_fuente(tmp_path):
    v=VentanaPrincipal();otro=VentanaPrincipal();w=v.widget_proyecciones;e=v.widget_empalme
    try:
        contexto(w)
        for emp,widget in ((False,w),(True,e)):
            crear(widget,emp,'unitario');crear(widget,emp,'total',p=10000000.)
            path=tmp_path/('emp.xlsx' if emp else 'ico.xlsx')
            (widget._generar_excel_empalme if emp else widget._generar_excel)(str(path));b=load_workbook(path)
            try:
                assert len(b.sheetnames)==5;s=b['Resumen operativo Empalme' if emp else 'Valores proyectados'];assert s.max_column==(18 if emp else 17)
                assert all(s.cell(row,col).data_type=='f' for row in (3,4) for col in ((8,9,14,15) if emp else (9,10,14,15)))
                assert not any(c.comment for sheet in b for row in sheet for c in row)
                # La hoja histórica Cálculos editables conserva guardas de disponibilidad.
                assert not any(any(f in str(c.value).upper() for f in ('IF(','AND(','ISNUMBER(')) for row in s for c in row if c.data_type=='f')
            finally:b.close()
        v._activar_analisis(w.historial[0]['id']);w.seleccionar_referencia_historica('2024-05');antes=deepcopy(w.valores);empantes=deepcopy(e.datos_sesion());p=tmp_path/'sin_dane.json';guardar(v,p)
        with ExitStack() as stack:
            spies=[stack.enter_context(patch(target,side_effect=AssertionError(target))) for target in ('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv','app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv','app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis','app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.cargar_archivo')]
            abrir(otro,p);assert otro.widget_proyecciones.valores==antes and otro.widget_empalme.datos_sesion()==empantes
            for i in range(2):
                assert fila(otro.widget_proyecciones,i=i)['UserRole']==fila(w,i=i)['UserRole']
                assert fila(otro.widget_empalme,True,i)['UserRole']==fila(e,True,i)['UserRole']
            assert all(s.call_count==0 for s in spies)
    finally:otro.close();v.close()

def test_K_json_corrupto_y_desbordamiento_no_reemplazan_estado(tmp_path):
    v=VentanaPrincipal();w=v.widget_proyecciones;e=v.widget_empalme
    try:
        contexto(w);crear(w,False,'unitario');crear(e,True,'unitario')
        antes=deepcopy(w.datos_sesion());empantes=deepcopy(e.datos_sesion())
        p=tmp_path/'corrupta.json';p.write_text('{"proyecciones_icociv":',encoding='utf8')
        with patch.object(QMessageBox,'warning') as aviso:abrir(v,p)
        assert aviso.called and w.datos_sesion()==antes and e.datos_sesion()==empantes
        entrada(w,'unitario');antes=deepcopy(w.datos_sesion())
        with patch.object(w.cantidad,'value',return_value=1e308),patch.object(QMessageBox,'warning') as aviso:w.boton_calcular.click()
        assert aviso.called and 'rango numérico' in aviso.call_args.args[2]
        assert w.datos_sesion()==antes
    finally:v.close()
