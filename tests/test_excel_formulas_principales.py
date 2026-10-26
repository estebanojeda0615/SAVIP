import json
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV, COLUMNAS_VALORES
from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV
from tests.fixtures.sesiones import cargar as cargar_sesion_sintetica, ruta_temporal as ruta_sesion_sintetica  # Fase 1A (AS-007)

ROOT=Path(__file__).resolve().parents[1]
APP=QApplication.instance() or QApplication([])

def test_icociv_principal_formulas_fuentes_auditoria_y_sesion(tmp_path):
    fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(deepcopy(fixture['proyecciones']))
    try:
        for r in w.valores:r['modalidad_importe']='unitario'
        w.valores[0]['proyecto']='=1+1';w.valores[0]['item']='=1+1';w.valores[0]['observacion_tecnica']='=1+1'
        before=deepcopy(w.datos_sesion());path=tmp_path/'icociv.xlsx';w._generar_excel(str(path))
        with_book=load_workbook(path);ws=with_book['Valores proyectados'];audit=with_book['Trazabilidad']
        # Fase 4G: las fechas de I0 o de I pueden añadir 1-2 columnas auxiliares; las coordenadas se toman del plan.
        from openpyxl.utils import get_column_letter
        columnas,_=w.plan_excel_valores();C={k:i for i,(k,_) in enumerate(columnas,1)};L={k:get_column_letter(i) for k,i in C.items()}
        assert ws.max_column==len(columnas) in (17,18,19) and [c.value for c in ws[2]][:9]==[t for _,t in COLUMNAS_VALORES][:9]
        assert [ws.cell(2,C[k]).value.split('\n')[0] for k in ('i0_icociv','i_icociv')]==['I0 ICOCIV','I ICOCIV proyectado']
        for row,r in enumerate(w.valores,3):
            for col,key in ((6,'cantidad'),(7,'precio_base'),(8,'anticipo_amortizado'),(C['i0_icociv'],'i0_icociv'),(C['i_icociv'],'i_icociv')):
                expected=r.get(key)
                if isinstance(expected,float):expected=float(format(expected,'.16g'))
                assert ws.cell(row,col).value==expected and ws.cell(row,col).data_type=='n', (row,col)
                assert ws.cell(row,col).fill.fgColor.rgb.endswith('FFF2CC')
            for col,key,audit_col in ((9,'base_ajustable',None),(C['factor'],'factor_icociv',14),(C['valor_unitario'],'valor_proyectado',16),(C['diferencia'],'diferencia_porcentual',17)):
                c=ws.cell(row,col);assert c.data_type=='f' and f'{row}' in c.value and '[' not in c.value
                assert c.comment is None and not c.alignment.wrap_text
                if audit_col:assert audit.cell(row,audit_col).value==float(format(r[key],'.16g')) and audit.cell(row,audit_col).data_type=='n'
                else:assert r[key]==r['precio_base']-r['anticipo_amortizado']
            assert ws.cell(row,C['diferencia']).number_format=='#,##0.00'  # Fase 4I: variación con 2 decimales
            I,I0,F=L['i_icociv'],L['i0_icociv'],L['factor']
            assert [ws.cell(row,C[k]).value for k in ('base_unitaria','factor','valor_unitario','diferencia')]==[f'=G{row}-H{row}',f'={I}{row}/{I0}{row}',f'=I{row}+I{row}*({F}{row}-1)',f'=({F}{row}-1)*100']
            assert ws.cell(row,C['factor']).number_format=='#,##0.00000000'
        assert all(ws.cell(3,col).data_type=='s' for col in (2,4,C['observacion']))
        assert not any(c.data_type=='f' for row in audit.iter_rows(max_row=2+len(w.valores)) for c in row)  # Auditoría original fija; 4E añade fórmulas en otro bloque.
        assert with_book.calculation.calcMode=='auto' and not with_book._external_links
        assert w.datos_sesion()==before
        with_book.close()
    finally:w.close()

def test_empalme_principal_escenarios_acero_y_heredados(tmp_path):
    fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
    w=WidgetEmpalmeICCPICOCIV();w.calculos=deepcopy(fixture['empalme'])
    inherited=deepcopy(w.calculos[0]);inherited.pop('caso');inherited['item']='=1+1'
    missing=deepcopy(w.calculos[0]);missing.pop('i0_iccp')
    w.calculos += [inherited,missing]
    try:
        before=deepcopy(w.calculos);path=tmp_path/'empalme.xlsx';w._generar_excel_empalme(str(path));b=load_workbook(path);s=b['Empalme ICCP-ICOCIV']
        start=next(c.row+1 for row in s for c in row if c.column==1 and c.value=='INSUMO CONTRACTUAL' and s.cell(c.row,2).value=='UNIDAD')
        assert not s.column_dimensions['L'].hidden
        for row,r in enumerate(w.calculos,start):
            if row>=start+6:
                assert s.cell(row,11).value==r['valor_actualizado'] and 'No reproducible' in s.cell(row,1).comment.text
                assert all(s.cell(row,col).data_type!='f' for col in (6,9,10,11));continue
            assert all(s.cell(row,col).data_type=='f' for col in (6,9,10,11))
            assert s.cell(row,10).value==f'=F{row}+I{row}' and s.cell(row,11).value==f'=(C{row}-L{row})+J{row}'
            assert all(s.cell(row,col).comment is None and 'IF(' not in s.cell(row,col).value for col in (6,9,10,11))
            if r['caso']=='solo_iccp':assert s.cell(row,9).value=='=0'
            if r['caso']=='solo_icociv':assert s.cell(row,6).value=='=0'
            if r['tipo_calculo']=='Cálculo especial acero':assert s.cell(row,12).value=='=0'
        assert s.cell(start+6,1).data_type=='s'
        steel=next(c.row+1 for row in s for c in row if c.column==1 and c.value=='INSUMO CONTRACTUAL' and s.cell(c.row,2).value=='P0')
        assert s.cell(steel,2).value.startswith('=C') and s.cell(steel,10).value==f'=IF(E{steel}="","No calculado",E{steel}-(H{steel}+B{steel}))'
        assert not any(c.data_type=='f' for row in b['Resultados registrados'].iter_rows(max_row=2+len(w.calculos)) for c in row)
        assert w.calculos==before;b.close()
    finally:w.close()

def test_excel_cancelacion_permiso_y_escritura_atomica_principales(tmp_path):
    from unittest.mock import patch
    from PySide6.QtWidgets import QFileDialog,QMessageBox
    from openpyxl import Workbook
    fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
    p=WidgetProyeccionesICOCIV();p.restaurar_datos_sesion(deepcopy(fixture['proyecciones']))
    e=WidgetEmpalmeICCPICOCIV();e.calculos=deepcopy(fixture['empalme'])
    try:
        for name,w,action,writer in (('icociv',p,p._exportar_excel,'_generar_excel'),('empalme',e,e.exportar_excel,'_generar_excel_empalme')):
            before=deepcopy(p.datos_sesion() if w is p else e.calculos)
            with patch.object(QFileDialog,'getSaveFileName',return_value=('','')),patch.object(w,writer) as generate:
                action();assert not generate.called
            path=tmp_path/(name+'.xlsx');path.write_bytes(b'archivo anterior intacto')
            def failure(book,temporary):
                Path(temporary).write_bytes(b'escritura incompleta');raise PermissionError('Fallo de escritura simulado')
            with patch.object(QFileDialog,'getSaveFileName',return_value=(str(path),'')),patch.object(Workbook,'save',failure),patch.object(QMessageBox,'critical') as error,patch.object(QMessageBox,'information') as ok:
                action();assert error.call_count==1 and not ok.called
            assert path.read_bytes()==b'archivo anterior intacto'
            assert (p.datos_sesion() if w is p else e.calculos)==before
        assert sorted(p.name for p in tmp_path.iterdir())==['empalme.xlsx','icociv.xlsx']
    finally:p.close();e.close()

def test_registros_incompletos_no_inventan_formulas_ni_notas_numericas(tmp_path):
    fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
    p=WidgetProyeccionesICOCIV();p.restaurar_datos_sesion(deepcopy(fixture['proyecciones']))
    e=WidgetEmpalmeICCPICOCIV();e.calculos=deepcopy(fixture['empalme'])
    p.valores[0]['i0_icociv']=0;p.valores[1]['precio_base']=None
    e.calculos[0]['i0_iccp']=0
    try:
        for name,w in (('icociv',p),('empalme',e)):
            path=tmp_path/(name+'.xlsx');(w._generar_excel if w is p else w._generar_excel_empalme)(str(path));b=load_workbook(path)
            s=b['Valores proyectados' if w is p else 'Empalme ICCP-ICOCIV']
            start=3 if w is p else next(c.row+1 for row in s for c in row if c.column==1 and c.value=='INSUMO CONTRACTUAL' and s.cell(c.row,2).value=='UNIDAD')
            for row in range(start,start+(2 if w is p else 1)):
                assert s.cell(row,1).comment is None if w is p else 'No reproducible' in s.cell(row,1).comment.text
                for col in ((9,13,14,16) if w is p else (6,9,10,11)):
                    assert s.cell(row,col).data_type!='f' and s.cell(row,col).comment is None
            if w is p:
                assert list(b['Trazabilidad'].values) and b['Trazabilidad'].cell(start,2).value==p.valores[0]['calculo_id']
            b.close()
    finally:p.close();e.close()
