"""Contratos de UI y Excel de la corrección 4A-D."""
import os,sys,json
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from PySide6.QtCore import Qt,QPoint
from PySide6.QtWidgets import QApplication,QLabel,QScrollArea,QStyle,QStyleOptionSpinBox
from PySide6.QtTest import QTest,QSignalSpy
from PySide6.QtGui import QFont,QFontDatabase
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
app=QApplication.instance() or QApplication([])
for filename in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):
    QFontDatabase.addApplicationFont(str(Path('C:/Windows/Fonts')/filename))
app.setFont(QFont('Segoe UI',10))
def ventana():
    with patch('app_icociv.interfaz.ventana_principal.aplicar_preferencia_sistema',return_value=True):return VentanaPrincipal()
def flecha(spin):
    option=QStyleOptionSpinBox();spin.initStyleOption(option)
    return spin.style().subControlRect(QStyle.CC_SpinBox,option,QStyle.SC_SpinBoxUp,spin).center()

def test_a_fechas_layout_controles_y_senales():
    v=ventana();e=v.widget_empalme
    try:
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for size in ((1100,700),(1366,768),(1920,1080)):
                v.resize(*size);v.show();v._cambiar_seccion(2);app.processEvents();QTest.qWait(40)
                scroll=e.findChild(QScrollArea);scroll.ensureWidgetVisible(e.fecha_final_mes);app.processEvents()
                labels=[l for l in e.findChildren(QLabel) if l.text() in ('Fecha inicial','Fecha final')]
                assert labels[0].mapTo(v,QPoint()).x()==labels[1].mapTo(v,QPoint()).x()
                for inicial,final in ((e.fecha_inicial_anio,e.fecha_final_anio),(e.fecha_inicial_mes,e.fecha_final_mes)):
                    assert inicial.mapTo(v,QPoint()).x()==final.mapTo(v,QPoint()).x()
                for label,spin in zip(labels,(e.fecha_inicial_anio,e.fecha_final_anio)):
                    assert abs(label.mapTo(v,label.rect().center()).y()-spin.mapTo(v,spin.rect().center()).y())<=2
                assert scroll.horizontalScrollBar().maximum()==0,(tema,size,scroll.horizontalScrollBar().maximum(),v.size().width())
        spin=e.fecha_inicial_mes;before=spin.value();spy=QSignalSpy(spin.valueChanged)
        plus=next(b for b in spin.parentWidget().findChildren(__import__('PySide6.QtWidgets',fromlist=['QPushButton']).QPushButton) if b.text()=='+')
        QTest.mouseClick(plus,Qt.LeftButton);assert spin.value()==before+1 and spy.count()==1
        spin.setValue(before);spin.setFocus();QTest.keyClick(spin,Qt.Key_Up);assert spin.value()==before+1
        QTest.keyClick(spin,Qt.Key_Tab);assert not spin.hasFocus()
        assert (e.fecha_inicial_anio.minimum(),e.fecha_inicial_anio.maximum())==(2000,2100)
    finally:v.close()

def test_b_intentos_reales_h24_y_asignaciones_silenciosas():
    from tests.test_horizonte_editable import ventana_horizonte
    v=ventana_horizonte()
    try:
        v.resize(1100,700);v.show();v.activateWindow();v._cambiar_seccion(1);s=v.spin_horizonte_personalizado;s.setEnabled(True)
        with patch.object(v,'_alertar_horizonte_excedido') as aviso:
            s.setValue(12);app.processEvents();QTest.mouseClick(s,Qt.LeftButton,pos=flecha(s));assert s.value()==13 and aviso.call_count==0
            s.setValue(24);before=(v.spin_anio.value(),v.spin_mes.value());spy=QSignalSpy(s.valueChanged)
            QTest.mouseClick(s,Qt.LeftButton,pos=flecha(s));assert aviso.call_count==1 and s.value()==24 and spy.count()==0
            s.setFocus();QTest.keyClick(s,Qt.Key_Up);assert aviso.call_count==2 and s.value()==24
            s.lineEdit().selectAll();QTest.keyClicks(s.lineEdit(),'25');assert aviso.call_count==3 and s.value()==24
            app.clipboard().setText('25');s.lineEdit().selectAll();QTest.keyClick(s.lineEdit(),Qt.Key_V,Qt.ControlModifier)
            assert aviso.call_count==4 and s.value()==24 and (v.spin_anio.value(),v.spin_mes.value())==before
            s.lineEdit().selectAll();menu=s.lineEdit().createStandardContextMenu();acciones=menu.actions()
            pegar=next(a for a in acciones if 'paste' in a.text().lower() or 'pegar' in a.text().lower())
            pegar.trigger();app.processEvents();QTest.qWait(30)
            assert aviso.call_count==5 and s.value()==24
            assert v.combo_horizonte.currentData() is None and s.maximum()==24
            s.setValue(12);s.setValue(24);s.setValue(25);v._sincronizar_periodo_desde_horizonte();app.processEvents();assert aviso.call_count==5
            v._restaurar_seleccion_y_parametros(v._seleccion_actual(),{'anio':2028,'mes':5,'horizonte':24})
            assert aviso.call_count==5 and s.value()==24
            v.spin_anio.setValue(v.spin_anio.value()+1);assert aviso.call_count==6
            s.setValue(12);assert v._horizonte_solicitado()==12
    finally:v.close()

def test_d_resultados_texto_tarjetas_y_historial_sin_motor():
    v=ventana()
    fixture=json.loads((ROOT/'docs/fase_4a_c/FIXTURE_DATOS.json').read_text(encoding='utf-8'))
    v.widget_proyecciones.restaurar_datos_sesion(deepcopy(fixture['proyecciones']))
    before=v.widget_proyecciones.datos_sesion()
    try:
        with patch.object(v.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor al recuperar')):
            for a in v.widget_proyecciones.historial:
                v._activar_analisis(a['id']);assert v.selector_resultados_icociv.currentData()==a['id']
        v._cambiar_seccion(3)
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for size in ((1100,700),(1366,768),(1920,1080)):
                v.resize(*size);v.show();v.valores_kpi['modelo'].setText('Fourier K=1 + Holt lineal');v._ajustar_tarjetas_resultados();app.processEvents();QTest.qWait(50)
                assert len(v.tarjetas_kpi)==6 and v.scroll_resultados.horizontalScrollBar().maximum()==0
                for key,card in v.tarjetas_kpi.items():
                    value=v.valores_kpi[key];action=next(l for l in card.findChildren(QLabel) if l.text()=='Ver explicación')
                    assert value.height()>=value.heightForWidth(value.width()),(size,key,value.size(),value.heightForWidth(value.width()))
                    assert value.y()+value.height()<=action.y(),(size,key)
                assert v.texto_resultados.height()>=360
                assert 'Resultado del horizonte solicitado' in v.texto_resultados.toPlainText()
                v.scroll_resultados.ensureWidgetVisible(v.texto_resultados);app.processEvents()
                v.texto_resultados.verticalScrollBar().setValue(v.texto_resultados.verticalScrollBar().maximum())
                assert v.texto_resultados.verticalScrollBar().value()>0
                with patch.object(v,'_mostrar_explicacion_tarjeta') as abrir:
                    card=v.tarjetas_kpi['modelo'];card.setFocus();QTest.keyClick(card,Qt.Key_Return)
                    assert abrir.call_count==1 and abrir.call_args.args[0]=='modelo'
        assert [a['resultado'] for a in before['historial_analisis']]==[a['resultado'] for a in v.widget_proyecciones.datos_sesion()['historial_analisis']]
    finally:v.close()

def test_c_excel_formulas_estaticos_inyeccion_y_escenario(tmp_path):
    from openpyxl import Workbook,load_workbook
    from app_icociv.exportables.excel_modulos import agregar_calculos_editables
    from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV
    from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
    fixture=json.loads((ROOT/'docs/fase_4a_d/FIXTURE_DATOS.json').read_text(encoding='utf-8'))
    p=WidgetProyeccionesICOCIV();p.restaurar_datos_sesion(deepcopy(fixture['proyecciones']))
    e=WidgetEmpalmeICCPICOCIV();e.calculos=deepcopy(fixture['empalme'])
    p.valores[0]['item']='=1+1';e.calculos[0]['item']='=1+1'
    e.observacion_general.setPlainText('=HYPERLINK("https://example.invalid","No ejecutar")')
    try:
        for name,w in (('ICOCIV',p),('EMPALME',e)):
            original=p.datos_sesion() if w is p else deepcopy(e.calculos)
            path=tmp_path/(name+'.xlsx');(w._generar_excel if w is p else w._generar_excel_empalme)(str(path))
            b=load_workbook(path);cached=load_workbook(path,data_only=True)
            assert ('Cálculos editables' in b.sheetnames if w is e else b.sheetnames==['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']) and not b._external_links
            formula=b['Valores proyectados' if w is p else 'Cálculos editables'];formulas=[c for row in formula for c in row if c.data_type=='f']
            assert formulas and all('!' not in c.value and '[' not in c.value for c in formulas)
            import re
            for c in formulas:
                for ref in re.findall(r'\b([A-Z]{1,3}[1-9][0-9]*)\b',c.value):assert formula[ref].value is not None,(c.coordinate,ref)
            for sheet in b:
                for row in sheet:
                    for c in row:
                        if c.value=='=1+1' or isinstance(c.value,str) and c.value.startswith('=HYPERLINK'):assert c.data_type=='s'
                assert all(6<=d.width<=60 for d in sheet.column_dimensions.values())
            static=cached['Valores proyectados' if w is p else 'Resultados registrados']
            assert any(isinstance(c.value,(int,float)) and c.value!=0 for row in static for c in row)
            assert static['A3'].value==1 and b[static.title]['A3'].number_format=='0'
            assert (p.datos_sesion() if w is p else e.calculos)==original
            b.close();cached.close()
        b=Workbook();s=agregar_calculos_editables(b,[{'item':'Sin escenario','r_total':123}],empalme=True)
        assert not any(c.data_type=='f' for row in s for c in row)
        assert any('falta un escenario' in str(c.value) for row in s for c in row)
    finally:p.close();e.close()
