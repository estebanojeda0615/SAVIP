"""Fase 4G-B: ámbito → proyección, calculadora sin autoselección y exportación por alcance.

Integración Qt real con motor real sobre anexos DANE SINTÉTICOS (tests/fixtures_dane_4g.py).
Sesión base: tipología del grupo 1 con dos capítulos (Pavimentos con 2 revisiones y
Bases), 5 actividades; capítulo homónimo «Pavimentos» en el grupo 2 con 1 actividad.
"""
import json
import os
import tempfile
from copy import deepcopy
from functools import lru_cache
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QMessageBox
from openpyxl import load_workbook

from tests.test_fase_4g_historial_dane_fechas import APP, esperar, cargar_anexo, proyectar, Contador
from tests.test_historial_empalme_fase_4b import guardar, abrir
from tests.fixtures_dane_4g import anexo_dane_sintetico, PAVIMENTOS, BASES
from tests.test_slider_comparacion_historica_fase_4d import sin_motor
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.widgets.historial_series import resolver_alcance, analisis_en_ambito

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/fase_4g_b_consulta_exportacion"
HOJAS = ["Resumen", "Valores proyectados", "Proyecciones ICOCIV", "Trazabilidad", "Metodología y fórmulas"]


def evidencia(nombre, datos):
    try:
        (OUT / nombre).write_text(json.dumps(datos, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    except PermissionError:
        pass


def _ruta(v, grupo, capitulo, objetivo):
    v.combo_grupo.setCurrentIndex(grupo); APP.processEvents()
    for combo in (v.combo_nivel2, v.combo_nivel3):
        combo.setCurrentIndex(1); APP.processEvents()
    v.combo_nivel4.setCurrentIndex(v.combo_nivel4.findText(capitulo)); APP.processEvents()
    v.spin_anio.setValue(objetivo[0]); v.spin_mes.setValue(objetivo[1]); APP.processEvents()


def _actividad(w, item, modalidad, cantidad, p, a=0):
    w._formulario_guardado = None
    w.item.setText(item); w._item_editado = True; w.unidad.setCurrentIndex(1)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modalidad))
    w.cantidad.setValue(cantidad); w.precio_base.setValue(p); w.anticipo.setValue(a)
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True), \
         patch.object(w, "_confirmar_vinculo_fuera_de_ambito", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


@lru_cache(maxsize=1)
def _sesion_base_json():
    tmp = Path(tempfile.mkdtemp(prefix="savip_4gb_"))
    anexo = anexo_dane_sintetico(tmp / "anexo_4gb.xlsx", "2026-08", mismo_rotulo=True, capitulo_extra=True)
    v = VentanaPrincipal(); w = v.widget_proyecciones
    try:
        cargar_anexo(v, anexo)
        _ruta(v, 1, PAVIMENTOS, (2027, 6)); assert proyectar(v) == 1
        _actividad(w, "Base asfáltica", "unitario", 100, 100000)
        _actividad(w, "Carpeta asfáltica", "total", 50, 5000000)
        _ruta(v, 1, PAVIMENTOS, (2027, 3)); assert proyectar(v) == 1
        _actividad(w, "Base asfáltica", "unitario", 100, 100000)  # misma actividad aparente en otra revisión
        _ruta(v, 1, BASES, (2027, 6)); assert proyectar(v) == 1
        _actividad(w, "Subbase granular", "unitario", 30, 45000)
        _actividad(w, "Base granular", "total", 20, 900000, 50000)
        _ruta(v, 2, PAVIMENTOS, (2027, 6)); assert proyectar(v) == 1
        _actividad(w, "Pavimento grupo 2", "unitario", 10, 70000)
        _ruta(v, 1, PAVIMENTOS, (2027, 9)); assert proyectar(v) == 1  # proyección sin actividades
        datos = deepcopy(w.datos_sesion())
        ruta_sesion = tmp / "sesion_4gb.json"; guardar(v, ruta_sesion)
        return json.dumps(datos), str(ruta_sesion), str(anexo)
    finally:
        v.close()


def base():
    texto, ruta, anexo = _sesion_base_json()
    return json.loads(texto), Path(ruta), Path(anexo)


def widget():
    datos, _, _ = base(); w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(datos); return w


def ambito(w, nivel, valor, grupo_texto=None):
    candidatos = [a for a in w._ambitos if a["nivel"] == nivel and a["valor"] == valor and (grupo_texto is None or grupo_texto in a["etiqueta"] + a["contexto"])]
    assert len(candidatos) == 1, [a["etiqueta"] for a in candidatos]
    return candidatos[0]


def elegir_ambito(w, a):
    w.selector_ambito.setCurrentIndex(w.selector_ambito.findData(a["clave"] if a else None)); APP.processEvents()


def formulario(w):
    return json.dumps(w._datos_calculadora(), sort_keys=True)


def borrador(w):
    w.item.setText("Nueva actividad"); w._item_editado = True; w.unidad.setCurrentIndex(2)
    w.cantidad.setValue(100); w.precio_base.setValue(120000); w.anticipo.setValue(0)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData("unitario"))
    w.nombre_componente.setText("Componente en borrador")


# ---------------------------------------------------------------- ámbitos y homónimos
def test_01_homonimos_y_tipologia_con_dos_series():
    w = widget()
    try:
        paviment = [a for a in w._ambitos if a["tipo"] == "categoria" and a["valor"] == PAVIMENTOS]
        assert len(paviment) == 2 and len({a["clave"] for a in paviment}) == 2
        assert len({a["etiqueta"] for a in paviment}) == 2 and all("—" in a["etiqueta"] for a in paviment)
        tipologia = ambito(w, w._fichas[0]["ruta"][2][0], w._fichas[0]["ruta"][2][1])
        ids = analisis_en_ambito(w._fichas, tipologia)
        series = {f["serie_final"] for f in w._fichas if f["id"] in ids}
        assert series == {PAVIMENTOS, BASES} and len(ids) == 4
        assert all(f["ruta"][0][1].endswith("1") for f in w._fichas if f["id"] in ids)
        g1, g2 = sorted(paviment, key=lambda a: a["contexto"])
        assert set(analisis_en_ambito(w._fichas, g1)).isdisjoint(analisis_en_ambito(w._fichas, g2))
        evidencia("MATRIZ_AMBITOS_Y_SERIES.json", {
            "datos": "Anexo DANE sintético", "ambitos": [{k: a.get(k) for k in ("tipo", "clave", "etiqueta", "contexto")}
                                                          | {"analisis": analisis_en_ambito(w._fichas, a)} for a in w._ambitos],
            "fichas": [{k: f[k] for k in ("id", "numero", "revision", "serie", "serie_final", "objetivo", "horizonte")}
                       | {"ruta": f["ruta"]} for f in w._fichas]})
    finally:
        w.close()


def test_02_checkbox_desmarcado_y_filtros_por_ambito_y_proyeccion():
    w = widget()
    try:
        assert not w.solo_analisis_activo.isChecked() and w.solo_analisis_activo.text() == "Mostrar solo registros de la selección"
        total = len(w.valores); assert total == 6
        tipologia = ambito(w, w._fichas[0]["ruta"][2][0], w._fichas[0]["ruta"][2][1])
        elegir_ambito(w, tipologia); assert len(w._valores_visibles) == total  # casilla OFF: toda la sesión
        w.solo_analisis_activo.setChecked(True); assert len(w._valores_visibles) == 5
        assert "5 de 6 actividades" in w.lbl_filtro_valores.text()
        revision = w._fichas[1]["id"]
        w.selector_analisis_valores.setCurrentIndex(w.selector_analisis_valores.findData(revision)); APP.processEvents()
        assert [r["analisis_id"] for r in w._valores_visibles] == [revision]
        sin = w._fichas[-1]["id"]
        w.selector_analisis_valores.setCurrentIndex(w.selector_analisis_valores.findData(sin)); APP.processEvents()
        assert not w._valores_visibles and w.sin_valores.text() == "No hay actividades guardadas para esta selección."
        assert len(w.valores) == total
        w.solo_analisis_activo.setChecked(False); assert len(w._valores_visibles) == total
    finally:
        w.close()


# ---------------------------------------------------------------- no autoselección
def test_03_navegar_no_toca_calculadora_ni_analisis_activo():
    w = widget()
    try:
        borrador(w); antes = formulario(w); activo = w.analisis_activo_id
        tipologia = ambito(w, w._fichas[0]["ruta"][2][0], w._fichas[0]["ruta"][2][1])
        for accion in (lambda: elegir_ambito(w, tipologia), lambda: w.solo_analisis_activo.setChecked(True),
                       lambda: w.selector_analisis_valores.setCurrentIndex(0),
                       lambda: elegir_ambito(w, ambito(w, "Serie exacta", BASES)),
                       lambda: elegir_ambito(w, None)):
            accion(); APP.processEvents()
            assert formulario(w) == antes and w.analisis_activo_id == activo
            assert not w.tabla_valores.selectionModel().selectedRows()
        # Ámbito con UNA sola actividad (proyección sin otras): tampoco se autocompleta.
        w.solo_analisis_activo.setChecked(True); elegir_ambito(w, ambito(w, w._fichas[0]["ruta"][-1][0], PAVIMENTOS, "sintético 2"))
        assert len(w._valores_visibles) == 1 and formulario(w) == antes and w.analisis_activo_id == activo
        # Revisión concreta: activa SOLO su analisis_id; no carga actividades.
        destino = w.selector_analisis_valores.itemData(1)
        from PySide6.QtTest import QSignalSpy
        espia = QSignalSpy(w.analisis_solicitado)
        w.selector_analisis_valores.setCurrentIndex(1); APP.processEvents()
        assert espia.count() == 1 and espia.at(0)[0] == destino
        w.establecer_analisis_activo(destino)  # lo que hace la ventana principal al recibir la señal
        assert w.analisis_activo_id == destino and formulario(w) == antes
        assert "Para nuevos cálculos:" in w.lbl_calculo_activo.text()
        # Clic de fila y Ver detalle: consulta pura.
        w._seleccionar_fila_valores(0); w.barra.activa = 1; w._accion_tabla("detalle")
        assert formulario(w) == antes and w.analisis_activo_id == destino
    finally:
        w.close()


def test_04_editar_explicito_cancelar_y_guardar_cambios():
    v = VentanaPrincipal(); w = v.widget_proyecciones
    try:
        datos, _, _ = base(); v.sesion_pendiente = {"proyecciones_icociv": datos}; v._restaurar_sesion_pendiente()
        borrador(w); antes = formulario(w); objetivo = w.valores[1]; total = len(w.valores)
        fila = next(i for i, r in enumerate(w._valores_visibles) if r["calculo_id"] == objetivo["calculo_id"])
        w._seleccionar_fila_valores(fila)
        with patch.object(w, "_confirmar_descartar_borrador", return_value=False):
            assert not w.editar_seleccionado()
        assert formulario(w) == antes and w._edicion_calculo_id is None
        with patch.object(w, "_confirmar_descartar_borrador", return_value=True), sin_motor() as motor:
            assert w.editar_seleccionado()
        assert motor.call_count == 0 and w.analisis_activo_id == objetivo["analisis_id"]
        assert (w.precio_base.value(), w.cantidad.value(), w.modalidad_importe.currentData(), w.item.text()) == \
            (objetivo["precio_base"], objetivo["cantidad"], "total", objetivo["item"])
        w.precio_base.setValue(6000000)
        with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
            w.boton_calcular.click()
        assert len(w.valores) == total and w.valores[1]["calculo_id"] == objetivo["calculo_id"]
        assert w.valores[1]["precio_base"] == 6000000 and w.valores[1]["analisis_id"] == objetivo["analisis_id"]
        assert w._edicion_calculo_id is None and w.boton_calcular.text() == "Calcular y agregar a la tabla"
        assert [r["calculo_id"] for r in w.valores] == [r["calculo_id"] for r in datos["valores"]]
    finally:
        v.close()


def test_05_sin_analisis_activo_y_ambito_distinto_del_activo():
    v = VentanaPrincipal(); w = v.widget_proyecciones
    try:
        datos, _, _ = base(); v.sesion_pendiente = {"proyecciones_icociv": datos}; v._restaurar_sesion_pendiente()
        borrador(w); antes = formulario(w); valores = deepcopy(w.valores)
        w.solo_analisis_activo.setChecked(True); tipologia = ambito(w, w._fichas[0]["ruta"][2][0], w._fichas[0]["ruta"][2][1])
        elegir_ambito(w, tipologia); visibles = [r["calculo_id"] for r in w._valores_visibles]
        w.selector_analisis.setCurrentIndex(0); APP.processEvents()
        assert w.analisis_activo_id is None and not w.boton_calcular.isEnabled()
        assert "Ninguna proyección activa" in w.lbl_calculo_activo.text()
        assert formulario(w) == antes and w.valores == valores and [r["calculo_id"] for r in w._valores_visibles] == visibles
        # Activo fuera del ámbito consultado: advertencia visible y confirmación antes de vincular.
        otro = next(f["id"] for f in w._fichas if f["id"] not in analisis_en_ambito(w._fichas, tipologia))
        v._activar_analisis(otro); assert "Estás consultando" in w.lbl_calculo_activo.text() and formulario(w) == antes
        with patch.object(w, "_confirmar_vinculo_fuera_de_ambito", return_value=False) as confirma:
            w.boton_calcular.click()
        assert confirma.call_count == 1 and w.valores == valores
        with patch.object(w, "_confirmar_vinculo_fuera_de_ambito", return_value=True), \
             patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
            w.boton_calcular.click()
        assert len(w.valores) == len(valores) + 1 and w.valores[-1]["analisis_id"] == otro
    finally:
        v.close()


# ---------------------------------------------------------------- exportación
def _exportar(w, alcance, carpeta, nombre):
    rutas = {}
    for formato in ("xlsx", "pdf", "docx"):
        ruta = carpeta / f"{nombre}.{formato}"
        w._generar_excel(str(ruta), alcance) if formato == "xlsx" else w._generar_informe(str(ruta), formato, alcance)
        rutas[formato] = ruta
    return rutas


def _texto_docx(ruta):
    from docx import Document
    return "\n".join(p.text for p in Document(str(ruta)).paragraphs)


def _texto_pdf(ruta):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)


def test_06_alcances_independientes_del_filtro_y_del_activo_en_tres_formatos(tmp_path):
    w = widget()
    try:
        borrador(w)
        tipologia = ambito(w, w._fichas[0]["ruta"][2][0], w._fichas[0]["ruta"][2][1])
        elegir_ambito(w, tipologia); w.solo_analisis_activo.setChecked(True)  # pantalla parcial
        antes = (formulario(w), w.analisis_activo_id, w.solo_analisis_activo.isChecked(), [r["calculo_id"] for r in w._valores_visibles])
        serie = ambito(w, "Serie exacta", PAVIMENTOS, "sintético 1")
        alcances = {"sesion": w.alcance_exportacion("sesion"), "categoria": w.alcance_exportacion("categoria"),
                    "serie": w.alcance_exportacion("serie", serie["clave"]),
                    "proyeccion": w.alcance_exportacion("proyeccion", analisis_id=w._fichas[1]["id"])}
        assert len(alcances["sesion"]["calculos"]) == 6 and len(alcances["categoria"]["calculos"]) == 5
        assert len(alcances["serie"]["calculos"]) == 3 and len(alcances["serie"]["analisis"]) == 3
        assert len(alcances["proyeccion"]["calculos"]) == 1 and alcances["proyeccion"]["analisis"] == [w._fichas[1]["id"]]
        assert alcances["categoria"]["posibles_duplicados"] and not alcances["proyeccion"]["posibles_duplicados"]
        matriz = {}
        for tipo, alcance in alcances.items():
            rutas = _exportar(w, alcance, tmp_path, f"QA_{tipo}")
            libro = load_workbook(rutas["xlsx"])
            try:
                s = libro["Valores proyectados"]; t = libro["Trazabilidad"]; columnas = s.max_column
                assert libro.sheetnames == HOJAS and s.max_row == 2 + len(alcance["calculos"])
                assert [t.cell(f, 2).value for f in range(3, 3 + len(alcance["calculos"]))] == alcance["calculos"]
                assert {s.cell(f, c).data_type for f in range(3, 3 + len(alcance["calculos"])) for c in (9, 10) if s.cell(f, c).value is not None} <= {"f"}
                resumen = " ".join(str(c.value) for fila in libro["Resumen"] for c in fila if c.value)
                assert ("Alcance exportado" in resumen) == (tipo != "sesion")
                proy = libro["Proyecciones ICOCIV"]
                ids_proy = [c.value for c in proy["A"] if isinstance(c.value, str) and len(c.value) == 32]
                assert set(ids_proy) == set(alcance["analisis"])
            finally:
                libro.close()
            for formato in ("pdf", "docx"):
                texto = _texto_pdf(rutas["pdf"]) if formato == "pdf" else _texto_docx(rutas["docx"])
                assert ("Alcance exportado" in texto) == (tipo != "sesion")
                fuera = [c for c in alcances["sesion"]["calculos"] if c not in alcance["calculos"]]
                assert all(c in texto for c in alcance["calculos"]) and not any(c in texto for c in fuera)
            matriz[tipo] = {"descripcion": alcance["descripcion"], "calculos": alcance["calculos"], "analisis": alcance["analisis"],
                            "excluidos": [c for c in alcances["sesion"]["calculos"] if c not in alcance["calculos"]],
                            "motivo_exclusion": "Fuera del alcance elegido (por IDs)", "posibles_duplicados": alcance["posibles_duplicados"],
                            "columnas_xlsx": columnas}
        assert (formulario(w), w.analisis_activo_id, w.solo_analisis_activo.isChecked(), [r["calculo_id"] for r in w._valores_visibles]) == antes
        evidencia("MATRIZ_EXPORTACION_POR_ALCANCE.json", {"datos": "Anexo DANE sintético", "alcances": matriz})
    finally:
        w.close()


def test_07_flujo_de_exportacion_cancelar_y_no_mutar_estado(tmp_path):
    w = widget()
    try:
        borrador(w); estado = lambda: (formulario(w), w.analisis_activo_id, w.solo_analisis_activo.isChecked(),
                                       json.dumps(w.valores, sort_keys=True, default=str))
        antes = estado()
        with patch.object(w, "_elegir_alcance_exportacion", return_value=None), \
             patch.object(QFileDialog, "getSaveFileName", side_effect=AssertionError("No debe pedir ruta")):
            w._exportar_formato("xlsx")
        alcance = w.alcance_exportacion("proyeccion", analisis_id=w._fichas[-1]["id"])  # sin actividades
        assert alcance["solo_tecnico"]
        with patch.object(w, "_elegir_alcance_exportacion", return_value=alcance), \
             patch.object(w, "_confirmar_avisos_alcance", return_value=False), \
             patch.object(QFileDialog, "getSaveFileName", side_effect=AssertionError("No debe pedir ruta")):
            w._exportar_formato("xlsx")
        destino = tmp_path / "tecnico.xlsx"
        with patch.object(w, "_elegir_alcance_exportacion", return_value=alcance), \
             patch.object(w, "_confirmar_avisos_alcance", return_value=True), \
             patch.object(QFileDialog, "getSaveFileName", return_value=(str(destino), "")), \
             patch.object(QMessageBox, "information"):
            w._exportar_formato("xlsx")
        libro = load_workbook(destino)
        assert libro.sheetnames == HOJAS and libro["Valores proyectados"].max_row == 2 and libro["Valores proyectados"].max_column == 17
        libro.close()
        with patch.object(w, "_elegir_alcance_exportacion", return_value=w.alcance_exportacion("sesion")), \
             patch.object(w, "_confirmar_avisos_alcance", return_value=True), \
             patch.object(QFileDialog, "getSaveFileName", return_value=("", "")):
            w._exportar_formato("pdf")
        assert estado() == antes and w._alcance_pendiente is None
    finally:
        w.close()


def test_08_legado_sin_vinculo_y_duplicados_aviso():
    datos, _, _ = base(); w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(datos)
    try:
        legado = deepcopy(w.valores[0]); legado["calculo_id"] = "legado01"; legado.pop("analisis_id")
        valores = w.valores + [legado]
        sesion = resolver_alcance("sesion", w._fichas, valores)
        assert "legado01" in sesion["calculos"] and not sesion["no_atribuibles"]
        categoria = resolver_alcance("categoria", w._fichas, valores, w._ambitos[0])
        assert "legado01" not in categoria["calculos"] and categoria["no_atribuibles"] == ["legado01"]
        assert len(sesion["calculos"]) == len(set(sesion["calculos"]))
        try:
            resolver_alcance("proyeccion", w._fichas, valores, None, None)
        except ValueError:
            pass
        else:
            raise AssertionError("Proyección vacía aceptada")
    finally:
        w.close()


def test_09_sesion_reabierta_sin_dane_conserva_consulta_y_formulario(tmp_path):
    datos, ruta, anexo = base()
    retirado = anexo.with_suffix(".retirado")
    v = None
    try:
        if anexo.exists():anexo.rename(retirado)
        v = VentanaPrincipal(); w = v.widget_proyecciones
        with sin_motor() as motor:
            abrir(v, ruta)
        assert motor.call_count == 0 and len(w.valores) == 6 and not w.solo_analisis_activo.isChecked()
        borrador(w); antes = formulario(w)
        for i in range(w.selector_ambito.count()):
            w.selector_ambito.setCurrentIndex(i); APP.processEvents(); assert formulario(w) == antes
        for tema in ("claro", "oscuro"):
            v._aplicar_tema(tema, persistir=False); APP.processEvents(); assert formulario(w) == antes
        alcance = w.alcance_exportacion("sesion"); assert len(alcance["calculos"]) == 6
    finally:
        if v is not None:v.close()
        if retirado.exists():retirado.rename(anexo)
