"""Fase 4G: historial Serie → Análisis → Actividades, DANE agosto→octubre y fechas de índices.

Integración Qt real. Los anexos DANE son SINTÉTICOS (tests/fixtures_dane_4g.py)
y las sesiones de referencia son fixtures de fases anteriores; ninguna cifra
corresponde a una publicación oficial.
"""
import json
import os
import tempfile
import time
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtWidgets import QMessageBox
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from openpyxl import load_workbook

from tests.test_slider_comparacion_historica_fase_4d import APP, reales, sintetico, sin_motor
from tests.test_historial_empalme_fase_4b import guardar, abrir
from tests.fixtures_dane_4g import anexo_dane_sintetico, valor_sintetico, PAVIMENTOS
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV, COLUMNAS_VALORES, plan_columnas_valores
from app_icociv.interfaz.widgets.historial_series import identidad_serie, huella_datos, firma_ejecucion
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.servicios.actualizacion_icociv import factor_aplicado
from tests.fixtures.salidas import carpeta_salida

ROOT = Path(__file__).resolve().parents[1]
OUT = carpeta_salida("fase_4g")  # Fase 1A (AS-007): antes en docs/, no publicado
TITULOS = [t for _, t in COLUMNAS_VALORES]


def evidencia(nombre, datos):
    """Guarda evidencia legible en una carpeta temporal (antes docs/fase_4g)."""
    try:
        (OUT / nombre).write_text(json.dumps(datos, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    except PermissionError:
        pass


def esperar(v, condicion, limite=180):
    fin = time.monotonic() + limite
    while time.monotonic() < fin:
        APP.processEvents()
        if condicion() and (v.trabajador is None or not v.trabajador.isRunning()):
            return
        QTest.qWait(20)
    raise AssertionError("Tiempo agotado: " + v.etiqueta_estado.text())


def cargar_anexo(v, ruta):
    v._cargar_archivo(str(ruta))
    # También se espera a que la ventana procese la carga (señal en cola): si no, rellenaría los selectores
    # DESPUÉS de que la prueba eligiera la ruta y la selección quedaría vacía.
    esperar(v, lambda: v.combo_grupo.count() == 3 and v.combo_grupo.isEnabled() and v.controlador.ruta_archivo == Path(ruta)
            and v.etiqueta_archivo.text() == Path(ruta).name)


def ruta_pavimentos(v, grupo=1, objetivo=(2027, 6)):
    """Capítulo Pavimentos y fecha objetivo (2027-06 = H10 desde 2026-08, H8 desde 2026-10)."""
    v.combo_grupo.setCurrentIndex(grupo); APP.processEvents()
    for combo in (v.combo_nivel2, v.combo_nivel3, v.combo_nivel4):
        combo.setCurrentIndex(1); APP.processEvents()
    v.spin_anio.setValue(objetivo[0]); v.spin_mes.setValue(objetivo[1]); APP.processEvents()


def proyectar(v):
    antes = v.resultado_ui_actual; n = len(v.widget_proyecciones.historial)
    v.ejecutar_proyeccion()
    if v.trabajador is not None and v.trabajador.isRunning():
        esperar(v, lambda: v.resultado_ui_actual is not antes)
    for _ in range(5): APP.processEvents()
    return len(v.widget_proyecciones.historial) - n


def actividad(w, nombre, modalidad, cantidad, p, a=0):
    w.item.setText(nombre); w._item_editado = True; w.unidad.setCurrentIndex(1)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modalidad))
    w.cantidad.setValue(cantidad); w.precio_base.setValue(p); w.anticipo.setValue(a)
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


class Contador:
    def __init__(self, v):
        self.n = 0; original = v.controlador.ejecutar_analisis

        def contado(*a, **k):
            self.n += 1; return original(*a, **k)
        v.controlador.ejecutar_analisis = contado


# ---------------------------------------------------------------- A. historial
def test_A_serie_analisis_actividades_reutilizar_revision_y_rotulos():
    with tempfile.TemporaryDirectory(prefix="savip_4g_") as tmp:
        tmp = Path(tmp); v = VentanaPrincipal(); w = v.widget_proyecciones; motor = Contador(v)
        try:
            cargar_anexo(v, anexo_dane_sintetico(tmp / "a.xlsx", "2026-08")); ruta_pavimentos(v)
            assert proyectar(v) == 1 and motor.n == 1
            a = actividad(w, "Actividad A", "unitario", 100, 100000)
            b = actividad(w, "Actividad A", "total", 25, 2500000)  # misma descripción, otro ID
            assert motor.n == 1 and len(w.historial) == 1 and a["calculo_id"] != b["calculo_id"]
            assert a["analisis_id"] == b["analisis_id"] == w.historial[0]["id"]
            grupos = w.grupos_historial()
            assert len(grupos) == 1 and len(grupos[0]["analisis"]) == 1 and len(grupos[0]["analisis"][0]["actividades"]) == 2
            # Segundo clic equivalente (Fase 4I-J): reutiliza la trayectoria guardada, sin motor, sin entrada nueva
            # y sin ventana de confirmación.
            assert proyectar(v) == 0 and motor.n == 1 and w.analisis_activo_id == w.historial[0]["id"]
            # Publicación revisada hasta 2026-10 (datos distintos, misma serie): «Ejecutar proyección» no crea
            # revisiones; la revisión nace solo del botón expreso «Actualizar análisis con DANE más reciente».
            cargar_anexo(v, anexo_dane_sintetico(tmp / "b.xlsx", "2026-10", revision=0.05)); ruta_pavimentos(v)
            assert proyectar(v) == 0 and motor.n == 1
            with patch.object(v, "_confirmar_reanalisis", return_value=True), patch.object(QMessageBox, "information"):
                v.actualizar_analisis_dane(); esperar(v, lambda: len(w.historial) == 2)
            assert motor.n == 2 and w.historial[1]["revision_de"] == w.historial[0]["id"]
            grupos = w.grupos_historial()
            assert len(grupos) == 1 and len(grupos[0]["analisis"]) == 2
            assert len({x["huella_datos"] for x in grupos[0]["analisis"]}) == 2
            # Mismo rótulo visible en otro grupo: identidad distinta → grupo nuevo (primera vez ⇒ motor).
            cargar_anexo(v, anexo_dane_sintetico(tmp / "c.xlsx", "2026-08", mismo_rotulo=True)); ruta_pavimentos(v, grupo=2)
            assert proyectar(v) == 1
            grupos = w.grupos_historial()
            assert len(grupos) == 2 and all(PAVIMENTOS in g["etiqueta"] for g in grupos)
            assert all(g["verificable"] for g in grupos)
            # Árbol: 2 series, actividades bajo su análisis; localizar por ID.
            assert w.arbol_historial.topLevelItemCount() == 2
            assert w.localizar_actividad(b["calculo_id"])
            fila = w.tabla_valores.selectionModel().selectedRows()[0].row()
            assert w._valores_visibles[fila]["calculo_id"] == b["calculo_id"]
            evidencia("MATRIZ_TRAZABILIDAD.json", {
                "datos": "Anexos DANE sintéticos (no oficiales); motor real",
                "grupos": [{**{k: g[k] for k in ("clave", "etiqueta", "tabla", "verificable")},
                            "analisis": [{**{k: x[k] for k in ("id", "etiqueta", "revision", "ultimo_observado", "huella_datos")},
                                          "periodo_I0_original": next((h["contexto"]["periodo_base"] for h in w.historial if h["id"] == x["id"])),
                                          "I0_original": next((h["contexto"]["indice_base"] for h in w.historial if h["id"] == x["id"])),
                                          "archivo": Path(next(h["archivo_excel"] for h in w.historial if h["id"] == x["id"])).name,
                                          "actividades": x["actividades"]} for x in g["analisis"]]} for g in grupos]})
        finally:
            v.close()


def test_A_sin_analisis_activo_real_en_los_tres_selectores_y_sesion(tmp_path):
    v = VentanaPrincipal(); w = v.widget_proyecciones
    try:
        v.sesion_pendiente = {"proyecciones_icociv": reales()}; v._restaurar_sesion_pendiente()
        historial = deepcopy(w.historial); valores = deepcopy(w.valores)
        # Fase 4G-B: el desplegable inferior pasó a ser «Proyección» de consulta; «Sin análisis activo» vive arriba y en Resultados.
        for selector in (w.selector_analisis, v.selector_resultados_icociv):
            v._activar_analisis(historial[1]["id"]); w.solo_analisis_activo.setChecked(True)
            assert w.analisis_activo_id == historial[1]["id"] and len(w._valores_visibles) == 1
            with sin_motor() as motor:
                selector.setCurrentIndex(0); APP.processEvents()
            assert motor.call_count == 0
            assert w.analisis_activo_id is None
            assert w.selector_analisis.currentData() is None and v.selector_resultados_icociv.currentData() is None
            # 4G-B: la consulta (casilla y proyección consultada) se conserva al desactivar.
            assert w.solo_analisis_activo.isChecked() and len(w._valores_visibles) == 1
            assert not w.boton_calcular.isEnabled() and "No hay análisis seleccionado" in w.lbl_contexto.text()
            assert "Ninguna proyección activa" in w.lbl_calculo_activo.text()
            w.solo_analisis_activo.setChecked(False)
            assert w.historial == historial and w.valores == valores
        p = tmp_path / "sin_activo.json"; guardar(v, p)
        assert cargar_sesion(p)["proyecciones_icociv"]["analisis_activo_id"] is None
        otro = VentanaPrincipal()
        try:
            with sin_motor() as motor:
                abrir(otro, p)
            assert otro.widget_proyecciones.analisis_activo_id is None and motor.call_count == 0
            assert otro.widget_proyecciones.valores == valores
        finally:
            otro.close()
    finally:
        v.close()


# ---------------------------------------------------------------- B. DANE agosto → octubre
def _canonico(w):
    return json.dumps({"datos": w.datos_sesion(), "escenarios": [w.escenario_registro(r) for r in w.valores]},
                      sort_keys=True, ensure_ascii=False, default=str)


def test_B_sesion_agosto_con_dane_octubre_invariante_actualizacion_y_restauracion():
    with tempfile.TemporaryDirectory(prefix="savip_4g_dane_") as tmp:
        tmp = Path(tmp); v = VentanaPrincipal(); w = v.widget_proyecciones; escenarios = {}
        otro = tercero = None
        try:
            cargar_anexo(v, anexo_dane_sintetico(tmp / "agosto.xlsx", "2026-08")); ruta_pavimentos(v)
            assert proyectar(v) == 1
            actividad(w, "Base 100 m3 unitario", "unitario", 100, 100000)
            actividad(w, "Base 100 m3 total", "total", 100, 10000000)
            sesion = tmp / "agosto.json"; guardar(v, sesion)
            antes = _canonico(w); escenarios["A_guardado"] = json.loads(antes)
            original = deepcopy(w.valores); analisis = deepcopy(w.historial[0])
            assert w._periodo_por_analisis[analisis["id"]] == "2026-08"
            v.close(); v = None
            otro = VentanaPrincipal(); n = otro.widget_proyecciones; motor = Contador(otro)
            with patch("app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv",
                       side_effect=AssertionError("Servicio económico")) as servicio:
                abrir(otro, sesion)  # recarga el anexo original de agosto y restaura sin recalcular
                esperar(otro, lambda: otro.sesion_pendiente is None and bool(n.valores))
                reabierto = _canonico(n)
                assert json.loads(reabierto)["datos"]["valores"] == json.loads(antes)["datos"]["valores"]
                assert json.loads(reabierto)["escenarios"] == json.loads(antes)["escenarios"]
                cargar_anexo(otro, anexo_dane_sintetico(tmp / "octubre.xlsx", "2026-10", revision=0.05))
                # B: cargar octubre no cambia NADA: slider en agosto, I0 original exacto, motor y servicio en cero.
                assert _canonico(n) == reabierto and motor.n == 0 and servicio.call_count == 0
                assert n.selector_referencia.currentText() == "2026-08" and n._periodos_slider[-1] == "2026-08"
                assert n.valores == original and n.historial[0]["resultado"] == analisis["resultado"]
                escenarios["B_tras_cargar_octubre"] = {"identico": True, "slider": "2026-08"}
                # C: actualización expresa con el último índice DANE disponible.
                with patch.object(n, "_confirmar_actualizacion_dane", return_value=True) as confirmar:
                    assert n.usar_ultimo_indice_dane()
                assert confirmar.call_count == 1 and motor.n == 0 and servicio.call_count == 0
                # Índice exacto tal como lo leyó el cargador del anexo de octubre (2016-01 + 129 meses, sintético).
                octubre = n._referencias_externas[analisis["id"]]["2026-10"]["indice"]
                assert abs(octubre - (valor_sintetico(1, 129) + 0.05)) < 1e-9
                assert n._periodo_por_analisis[analisis["id"]] == "2026-10" and n._overrides[analisis["id"]]
                for r in n.valores:
                    e, estado = n.escenario_registro(r)
                    assert e["i0_icociv"] == octubre and e["i_icociv"] == r["i_icociv"] and estado.startswith("Referencia modificada")
                    # Fase 1A (AS-007): la expectativa era el factor I/I0 sin redondear, anterior a la regla 4I-C.
                    # Contrato vigente (comparacion_historica.escenario_registro -> aplicar_factor, golden economico):
                    # F = 1 + ROUND_HALF_UP(((I/I0) - 1) * 100; 2) / 100.
                    assert e["factor_icociv"] == float(factor_aplicado(r["i_icociv"], octubre)[1])
                assert n.valores == original and n.historial[0]["resultado"] == analisis["resultado"]
                escenarios["C_octubre_expreso"] = {"I0": octubre, "escenarios": [n.escenario_registro(r)[0] for r in n.valores]}
                sesion_oct = tmp / "octubre.json"; guardar(otro, sesion_oct)
                guardado = cargar_sesion(sesion_oct)["proyecciones_icociv"]["comparacion_historica"]
                assert guardado["version"] == 3 and guardado["referencias_externas"][analisis["id"]]["2026-10"]["indice"] == octubre
            con_octubre = _canonico(n)
            for anexo in (tmp / "agosto.xlsx", tmp / "octubre.xlsx"):
                anexo.rename(anexo.with_suffix(".no_disponible"))  # reabrir SIN ningún anexo DANE
            tercero = VentanaPrincipal()
            with sin_motor() as m3:
                abrir(tercero, sesion_oct)
            t = tercero.widget_proyecciones
            assert m3.call_count == 0 and _canonico(t) == con_octubre and t.selector_referencia.currentText() == "2026-10"
            assert "anexo posterior" in t.lbl_referencia_comparacion.text()
            # D: restaurar = agosto y valores registrados exactos.
            t.restaurar_referencia_analisis()
            assert t._periodo_por_analisis[analisis["id"]] == "2026-08" and not t._overrides[analisis["id"]]
            assert [t.escenario_registro(r)[0] for r in t.valores] == original
            escenarios["D_restaurado"] = {"identico_al_original": True}
            evidencia("ESCENARIOS_AGOSTO_OCTUBRE.json", escenarios)
        finally:
            for x in (v, otro, tercero):
                if x is not None:x.close()


def test_F_rechazos_no_destructivos_y_global_atomico():
    with tempfile.TemporaryDirectory(prefix="savip_4g_rechazo_") as tmp:
        tmp = Path(tmp); v = VentanaPrincipal(); w = v.widget_proyecciones
        try:
            cargar_anexo(v, anexo_dane_sintetico(tmp / "agosto.xlsx", "2026-08")); ruta_pavimentos(v)
            assert proyectar(v) == 1; actividad(w, "Actividad", "unitario", 2, 1000)
            antes = _canonico(w)
            casos = {"sin_ruta": anexo_dane_sintetico(tmp / "sin_ruta.xlsx", "2026-10", sin_ruta=True),
                     "mismo_mes_revisado": anexo_dane_sintetico(tmp / "revisado.xlsx", "2026-08", revision=0.05),
                     "mas_alla_del_objetivo": anexo_dane_sintetico(tmp / "largo.xlsx", "2027-07")}
            for nombre, ruta in casos.items():
                cargar_anexo(v, ruta)
                with patch.object(w, "_confirmar_actualizacion_dane", side_effect=AssertionError("No debe preguntar")):
                    assert not w.usar_ultimo_indice_dane(), nombre
                assert _canonico(w) == antes and "Ningún escenario cambió" in w.lbl_estado_comparacion.text(), nombre
            cargar_anexo(v, casos["mismo_mes_revisado"])
            assert not w.usar_ultimo_indice_dane() and "revisó 2026-08" in w.lbl_estado_comparacion.text()
            # Global con un análisis cuya serie no existe en el anexo: no se aplica a ninguno.
            cargar_anexo(v, tmp / "agosto.xlsx"); ruta_pavimentos(v, grupo=2)
            assert proyectar(v) == 1
            cargar_anexo(v, casos["sin_ruta"])
            antes = _canonico(w); w.aplicar_global.setChecked(True); assert w.aplicar_global.isChecked()
            with patch.object(w, "_confirmar_actualizacion_dane", side_effect=AssertionError("No debe preguntar")):
                assert not w.usar_ultimo_indice_dane()
            w.aplicar_global.setChecked(False); assert _canonico(w) == antes
            # Sin anexo cargado: tampoco cambia nada.
            w.fuente_dane = lambda: None
            assert not w.usar_ultimo_indice_dane() and _canonico(w) == antes
        finally:
            v.close()


def test_G_sesiones_legado_v1_v2_v3_y_corrupcion():
    w = WidgetProyeccionesICOCIV()
    try:
        datos = reales(); w.restaurar_datos_sesion(datos)  # V2 sin referencias externas
        assert w.datos_sesion()["comparacion_historica"]["version"] == 2 and not w._referencias_externas
        identificador = w.historial[0]["id"]
        destino = next(a["id"] for a in w.historial if a["contexto"]["periodo_proyectado"][:7] == "2027-03")
        v1 = deepcopy(datos); v1["comparacion_historica"] = {"version": 1, "periodo_por_analisis": {identificador: "2024-02"}}
        w.restaurar_datos_sesion(v1); assert w._consultas_v1 == {identificador: "2024-02"} and not w._overrides.get(identificador)
        sin_bloque = deepcopy(datos); sin_bloque.pop("comparacion_historica", None); w.restaurar_datos_sesion(sin_bloque)
        valido = deepcopy(w.datos_sesion())
        # Análisis con objetivo 2027-03 y snapshot hasta 2026-12: solo 2027-01 y 2027-02 son admisibles.
        v3 = deepcopy(valido); bloque = v3["comparacion_historica"]; bloque["version"] = 3
        externa = {"indice": 120.5, "archivo": "anexo_sintetico.xlsx", "tabla": "T_16", "huella_fuente": "x" * 64, "aplicado": "2026-10-13T00:00:00"}
        for corrupto in ({"2027-00": externa}, {"2027-03": externa}, {"2026-11": externa}, {"2027-01": {**externa, "indice": "120"}},
                         {"2027-01": {**externa, "indice": -1.0}}, {"2027-01": {**externa, "archivo": ""}}):
            bloque["referencias_externas"] = {destino: corrupto}
            try:
                w.restaurar_datos_sesion(v3)
            except ValueError:
                pass
            else:
                raise AssertionError(f"Se aceptó {corrupto}")
            assert w.datos_sesion() == valido
        bloque["referencias_externas"] = {destino: {"2027-01": externa}}
        w.restaurar_datos_sesion(v3)
        assert "2027-01" in w._historicos_comparacion[destino]
        assert w._periodo_por_analisis[destino] == valido["comparacion_historica"]["selecciones"][destino]["periodo"]
        assert w.datos_sesion()["comparacion_historica"]["referencias_externas"] == {destino: {"2027-01": externa}}
    finally:
        w.close()


# ---------------------------------------------------------------- C. fechas 0/1/2 columnas
def _copia_analisis(w, origen, sufijo, objetivo=None):
    """Copia de prueba de un análisis (otro ID); ``objetivo`` cambia su periodo proyectado y el de sus filas."""
    a = deepcopy(next(x for x in w.historial if x["id"] == origen)); a["id"] = origen + sufijo
    a["nombre"] = a["nombre"] + " (copia de prueba)"
    registros = [deepcopy(r) for r in w.valores if r.get("analisis_id") == origen]
    if objetivo:
        a["contexto"]["periodo_proyectado"] = objetivo
        for r in registros: r["periodo_proyectado"] = objetivo
    for r in registros: r["analisis_id"] = a["id"]; r["calculo_id"] = r["calculo_id"] + sufijo
    a["registros_ids"] = [r["calculo_id"] for r in registros]
    w.historial.append(a); w.valores.extend(registros); w._preparar_comparacion_historica(); w._refrescar_historial(); w._refrescar_tabla_valores()
    return a["id"]


def _casos_fechas():
    w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(reales())
    primero = w.historial[0]["id"]
    w.historial = [a for a in w.historial if a["id"] == primero]; w.valores = [r for r in w.valores if r.get("analisis_id") == primero]
    for a in w.historial: a["registros_ids"] = [r["calculo_id"] for r in w.valores]
    w._preparar_comparacion_historica(); w._refrescar_historial()
    copia = _copia_analisis(w, primero, "_b")  # mismo I0 y mismo objetivo
    yield "00_homogeneas", w, 17
    w.establecer_analisis_activo(copia); w.seleccionar_referencia_historica("2024-02")
    yield "10_solo_I0", w, 18
    w.close()
    w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(reales())
    yield "01_solo_I", w, 18
    w.establecer_analisis_activo(w.historial[0]["id"]); w.seleccionar_referencia_historica("2024-02")
    yield "11_ambas", w, 19
    w.close()


def test_C_cuatro_combinaciones_en_qt_y_excel(tmp_path):
    resumen = {}
    for nombre, w, esperado in _casos_fechas():
        m = w.modelo_valores; fechas = [w.fechas_registro(r) for r in w._valores_visibles]
        assert m.columnCount() == esperado, (nombre, m.columnCount())
        titulos = [m.headerData(c, Qt.Horizontal) for c in range(m.columnCount())]
        base = [t for t in m._dataframe.columns if t not in ("Fecha I0 (AAAA-MM)", "Fecha de I proyectado (AAAA-MM)")]
        assert base == TITULOS, nombre
        i0, i = m.claves.index("i0_icociv"), m.claves.index("i_icociv")
        assert titulos[i0].startswith("I0 ICOCIV") and titulos[i].startswith("I ICOCIV proyectado")
        for posicion, (clave, auxiliar) in enumerate((("i0_icociv", "fecha_i0"), ("i_icociv", "fecha_i"))):
            unicas = {f[posicion] for f in fechas}
            if len(unicas) == 1:
                assert auxiliar not in m.claves and titulos[m.claves.index(clave)].endswith("\n" + next(iter(unicas)))
            else:
                assert m.claves.index(auxiliar) == m.claves.index(clave) + 1
                assert [m.index(f, m.claves.index(auxiliar)).data() for f in range(m.rowCount())] == [x[posicion] for x in fechas]
        # Copia TSV y valores numéricos de los índices intactos.
        w._seleccionar_fila_valores(0); QTest.keyClick(w.tabla_valores, Qt.Key_C, Qt.ControlModifier)
        celdas = APP.clipboard().text().split("\t"); assert len(celdas) == esperado
        assert float(celdas[i0]) == w.escenario_registro(w._valores_visibles[0])[0]["i0_icociv"]
        # Excel: evaluado sobre todos los registros exportados.
        for r in w.valores: r["modalidad_importe"] = "unitario"
        p = tmp_path / f"{nombre}.xlsx"; w._generar_excel(str(p)); b = load_workbook(p)
        try:
            s = b["Valores proyectados"]; cols, _ = w.plan_excel_valores(); C = {k: n for n, (k, _) in enumerate(cols, 1)}
            assert b.sheetnames == ["Resumen", "Valores proyectados", "Proyecciones ICOCIV", "Trazabilidad", "Metodología y fórmulas"]
            assert s.max_column == esperado and s.max_row == 2 + len(w.valores) and s.auto_filter.ref.endswith(f"{2 + len(w.valores)}")
            assert not any(c.comment for h in b for row in h for c in row)
            from openpyxl.utils import get_column_letter as L
            for f in range(3, 3 + len(w.valores)):
                factor = s.cell(f, C["factor"])
                assert factor.data_type == "f" and factor.value == f"={L(C['i_icociv'])}{f}/{L(C['i0_icociv'])}{f}"
                for clave in ("base_unitaria", "valor_unitario", "diferencia"):
                    assert s.cell(f, C[clave]).data_type == "f"
                    assert not any(L(C[k]) + str(f) in s.cell(f, C[clave]).value for k in ("fecha_i0", "fecha_i") if k in C)
                assert s.cell(f, C["i0_icociv"]).data_type == "n" and s.cell(f, C["i_icociv"]).data_type == "n"
            resumen[nombre] = {"qt_columnas": esperado, "xlsx_columnas": s.max_column, "encabezados_qt": titulos,
                               "encabezados_xlsx": [c.value for c in s[2]], "auxiliares": [k for k in C if k.startswith("fecha_")],
                               "formulas_fila_3": {k: s.cell(3, C[k]).value for k in ("base_unitaria", "factor", "valor_unitario", "valor_total", "diferencia")},
                               "excel_nativo": "No ejecutado en esta suite (openpyxl no recalcula fórmulas)"}
        finally:
            b.close()
    evidencia("FECHAS_COLUMNAS_4_CASOS.json", resumen)


def test_C_filtro_cambia_solo_columnas_y_conserva_seleccion_y_cero_filas():
    w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(reales())
    try:
        m = w.modelo_valores; objetivo = w.valores[0]["calculo_id"]
        w.establecer_analisis_activo(w.valores[0]["analisis_id"]); w.localizar_actividad(objetivo)
        assert m.columnCount() == 18
        w.solo_analisis_activo.setChecked(True)
        assert m.columnCount() == 17 and w._valores_visibles[w.tabla_valores.selectionModel().selectedRows()[0].row()]["calculo_id"] == objetivo
        w.solo_analisis_activo.setChecked(False)
        assert m.columnCount() == 18 and w._valores_visibles[w.tabla_valores.selectionModel().selectedRows()[0].row()]["calculo_id"] == objetivo
        assert plan_columnas_valores([]) == (list(COLUMNAS_VALORES), {"i0_icociv": "I0 ICOCIV", "i_icociv": "I ICOCIV proyectado"})
        assert len(plan_columnas_valores([("2026-08", "2027-08"), ("", "2027-08")])[0]) == 18  # fecha faltante → auxiliar
    finally:
        w.close()


# ---------------------------------------------------------------- D. equivalencia económica
def test_D_100_m3_ambas_modalidades_con_0_1_2_auxiliares():
    w = sintetico()
    try:
        w.valores.clear(); w._refrescar_tabla_valores()
        unitario = actividad(w, "100 m3 unitario", "unitario", 100, 100000)
        total = actividad(w, "100 m3 total", "total", 100, 10000000)
        esperado = {"base_unitaria": 100000.0, "base_total": 10000000.0, "valor_unitario": 110000.0, "valor_total": 11000000.0}
        for configuracion in ("0 auxiliares", "1 auxiliar", "2 auxiliares"):
            if configuracion == "1 auxiliar":
                _copia_analisis(w, w.historial[0]["id"], "_c"); w.establecer_analisis_activo(w.historial[0]["id"] + "_c")
                w.seleccionar_referencia_historica("2024-02")
            if configuracion == "2 auxiliares":
                _copia_analisis(w, w.historial[0]["id"], "_d", objetivo="2030-01")
            m = w.modelo_valores
            assert m.columnCount() == {"0 auxiliares": 17, "1 auxiliar": 18, "2 auxiliares": 19}[configuracion]
            for r in (unitario, total):
                fila = next(i for i, x in enumerate(w._valores_visibles) if x["calculo_id"] == r["calculo_id"])
                for clave, valor in esperado.items():
                    assert abs(m.index(fila, m.claves.index(clave)).data(Qt.UserRole) - valor) < 1e-6, (configuracion, clave)
    finally:
        w.close()


# ---------------------------------------------------------------- E. nueva proyección expresa
def test_E_nueva_proyeccion_desde_octubre_es_revision_y_conserva_agosto():
    with tempfile.TemporaryDirectory(prefix="savip_4g_e_") as tmp:
        tmp = Path(tmp); v = VentanaPrincipal(); w = v.widget_proyecciones
        try:
            cargar_anexo(v, anexo_dane_sintetico(tmp / "agosto.xlsx", "2026-08")); ruta_pavimentos(v)
            assert proyectar(v) == 1; agosto = deepcopy(w.historial[0])
            cargar_anexo(v, anexo_dane_sintetico(tmp / "octubre.xlsx", "2026-10", revision=0.05)); ruta_pavimentos(v)
            assert proyectar(v) == 0  # Fase 4I-J: «Ejecutar» no crea revisiones con otros datos; solo el botón expreso
            with patch.object(v, "_confirmar_reanalisis", return_value=True), patch.object(QMessageBox, "information"):
                v.actualizar_analisis_dane(); esperar(v, lambda: len(w.historial) == 2)
            assert w.historial[0] == {**agosto, "registros_ids": w.historial[0]["registros_ids"], "calculadora": w.historial[0]["calculadora"]}
            assert w.historial[1]["contexto"]["periodo_base"].replace("_", "-").startswith("2026-10")
            assert len(w.grupos_historial()) == 1 and identidad_serie(w.historial[0]["resultado"]) == identidad_serie(w.historial[1]["resultado"])
            assert huella_datos(w.historial[0]["resultado"]) != huella_datos(w.historial[1]["resultado"])
            assert firma_ejecucion(w.historial[0]["resultado"], w.historial[0]["parametros"]) != firma_ejecucion(w.historial[1]["resultado"], w.historial[1]["parametros"])
        finally:
            v.close()
