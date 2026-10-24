"""Fase 4I-C V2: factor económico desde el porcentaje redondeado, clic en puntos de la gráfica y carga minimalista.

Pruebas focalizadas con datos SINTÉTICOS (no oficiales DANE); el motor estadístico no se ejecuta.
"""
import json
import os
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QLabel, QMessageBox, QWidget
from matplotlib.backend_bases import MouseEvent
from openpyxl import load_workbook

from tests.test_slider_comparacion_historica_fase_4d import APP, reales, sintetico
from tests.test_proyecciones_redisenio_visual import ventana_de_prueba
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.widgets.visor_grafica import VisorGraficaDialog, mes_legible
from app_icociv.persistencia.historial_icociv import recuperar_resultado
from app_icociv.servicios.actualizacion_icociv import CRITERIO_FACTOR, calcular_actualizacion_icociv, es_factor_legado
from app_icociv.utilidades.importes_presupuestales import desglose_presupuestal

ROOT = Path(__file__).resolve().parents[1]
MATRIZ = ((102.741, 2.74, 1.0274, 102740.0), (102.745, 2.75, 1.0275, 102750.0), (97.255, -2.75, 0.9725, 97250.0),
          (100.0, 0.0, 1.0, 100000.0), (102.744, 2.74, 1.0274, 102740.0))


def _actividad(w, p, a=0.0, cantidad=1.0, modalidad="unitario", item="Actividad"):
    w._formulario_guardado = None
    w.item.setText(item); w._item_editado = True; w.unidad.setCurrentIndex(1)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modalidad))
    w.cantidad.setValue(cantidad); w.precio_base.setValue(p); w.anticipo.setValue(a)
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


# ---------------------------------------------------------------- C: factor aplicado
def test_C1_matriz_y_falla_si_se_usa_el_cociente_crudo():
    for i, p, f, valor in MATRIZ:
        r = calcular_actualizacion_icociv({"precio_base": 100000, "anticipo_amortizado": 0, "i0_icociv": 100, "i_icociv": i})
        assert (r["diferencia_porcentual"], r["factor_icociv"], r["valor_proyectado"], r["r_total"]) == (p, f, valor, valor - 100000)
        assert Decimal(repr(r["diferencia_porcentual"])) == (Decimal(repr(r["factor_icociv"])) - 1) * 100  # invariante decimal
        assert r["criterio_factor"] == CRITERIO_FACTOR and r["i0_icociv"] == 100 and r["i_icociv"] == i  # índices intactos
        assert r["factor_icociv_indices"] == i / 100
    crudo = 100000 * (102.741 / 100)
    assert calcular_actualizacion_icociv({"precio_base": 100000, "i0_icociv": 100, "i_icociv": 102.741})["valor_proyectado"] != crudo


def test_C2_calculadora_tabla_modalidades_y_edicion_con_el_mismo_factor():
    w = sintetico(); w.valores.clear(); w._refrescar_tabla_valores()
    w.contexto["indice_proyectado"] = 102.741  # I0 = 100 → 2.741 % → 2.74 %
    try:
        u = _actividad(w, 150000, 50000, 3, "unitario", "Unitario")
        t = _actividad(w, 150000, 50000, 3, "total", "Total")
        for r in (u, t):
            assert (r["base_ajustable"], r["factor_icociv"], r["diferencia_porcentual"], r["valor_proyectado"], r["r_total"]) == \
                (100000, 1.0274, 2.74, 102740.0, 2740.0)
        du, dt = desglose_presupuestal(u, u["valor_proyectado"]), desglose_presupuestal(t, t["valor_proyectado"])
        assert (du["base_unitaria"], du["valor_unitario"], du["base_total"], du["valor_total"]) == (100000, 102740, 300000, 308220)
        assert (dt["base_total"], dt["valor_total"]) == (100000, 102740) and dt["valor_unitario"] == 102740 / 3
        m = w.modelo_valores; col = m.claves.index
        assert m.data(m.index(0, col("factor"))) == "1.0274" and m.data(m.index(0, col("diferencia"))) == "2.74"
        assert m.data(m.index(0, col("valor_unitario"))) == "102,740" and m.data(m.index(0, col("valor_total"))) == "308,220"
        w._mostrar_resultado(u); html = w.resultado.toHtml()
        assert "1.0274" in html and "porcentaje" in html.lower()
        # Edición directa: el mismo factor aplicado y Deshacer restaura exactamente.
        previo = w.valores[0]
        assert m.setData(m.index(0, col("precio_base")), "250,000", Qt.EditRole)
        assert w.valores[0]["valor_proyectado"] == 205480.0 and w.valores[0]["factor_icociv"] == 1.0274
        w.pila_deshacer.undo(); assert w.valores[0] is previo
        # Renombrar un ítem no recalcula ni cambia el criterio.
        assert m.setData(m.index(0, col("item")), "Renombrado", Qt.EditRole)
        assert {k: w.valores[0][k] for k in ("factor_icociv", "valor_proyectado", "fecha_calculo")} == \
               {k: previo[k] for k in ("factor_icociv", "valor_proyectado", "fecha_calculo")}
    finally:
        w.close()


def test_C3_slider_expreso_usa_el_porcentaje_redondeado():
    w = sintetico()  # registro con I0 = 100 e I = 110 (vinculado al análisis)
    try:
        r = w.valores[0]; base = r["base_ajustable"]
        w.seleccionar_referencia_historica("2024-04")  # I0 efectivo 120 → −8.3333 % → −8.33 %
        e, estado = w.escenario_registro(r)
        assert "modificada" in estado.lower()
        assert (e["i0_icociv"], e["diferencia_porcentual"], e["factor_icociv"]) == (120.0, -8.33, 0.9167)
        assert e["valor_proyectado"] == float(Decimal(repr(base)) * Decimal("0.9167")) != base * 110 / 120
        assert r["i0_icociv"] == 100 and r["i_icociv"] == 110  # el registro y el pronóstico no cambian
    finally:
        w.close()


def test_C4_excel_pdf_word_y_sesion_con_el_factor_aplicado(tmp_path):
    w = sintetico(); w.valores.clear(); w._refrescar_tabla_valores(); w.contexto["indice_proyectado"] = 102.741
    try:
        _actividad(w, 100000, 0, 2, "unitario", "Excel")
        w._generar_excel(str(tmp_path / "f.xlsx")); libro = load_workbook(tmp_path / "f.xlsx"); ws = libro["Valores proyectados"]
        columnas, _ = w.plan_excel_valores(); n = {k: i for i, (k, _) in enumerate(columnas, 1)}
        from openpyxl.utils import get_column_letter as L
        I, I0, F = L(n["i_icociv"]), L(n["i0_icociv"]), L(n["factor"])
        assert ws.cell(3, n["factor"]).value == f"=1+ROUND(({I}3/{I0}3-1)*100,2)/100" and ws.cell(3, n["factor"]).number_format == "0.0000"
        assert ws.cell(3, n["diferencia"]).value == f"=ROUND(({F}3-1)*100,2)"
        assert f"{F}3" in ws.cell(3, n["valor_unitario"]).value  # R/Valor usan el factor APLICADO de la fila
        assert len(libro.sheetnames) == 5 and "ROUND" in str([c.value for c in libro["Metodología y fórmulas"]["B"]])
        libro.close()
        from app_icociv.reportes.sesion_modulos import construir_proyecciones
        texto = json.dumps([[getattr(b, "texto", None) or getattr(b, "filas", None) for b in s.bloques] for s in construir_proyecciones(w).secciones],
                           ensure_ascii=False, default=str)
        assert "1.0274" in texto and "102,740" in texto and "ROUND" in texto  # PDF y Word comparten este modelo
        datos = deepcopy(w.datos_sesion()); otro = WidgetProyeccionesICOCIV(); otro.restaurar_datos_sesion(datos)
        try:
            assert otro.valores == w.valores and otro.valores[0]["criterio_factor"] == CRITERIO_FACTOR
        finally:
            otro.close()
    finally:
        w.close()


def test_C5_legados_se_conservan_y_solo_se_actualizan_expresamente(tmp_path):
    fixture = json.loads((ROOT / "docs/fase_4a_d/FIXTURE_DATOS.json").read_text(encoding="utf-8"))
    w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(deepcopy(fixture["proyecciones"]))
    try:
        antes = deepcopy(w.valores); r0 = w.valores[0]
        assert es_factor_legado(r0) and r0["factor_icociv"] == r0["i_icociv"] / r0["i0_icociv"]
        m = w.modelo_valores
        # Fase 4I-D: el factor original se muestra solo como número; el estado va en el aviso sobre la tabla.
        assert m.data(m.index(0, m.claves.index("factor"))) == f"{r0['factor_icociv']:.6f}" and w.aviso_legados.isVisibleTo(w)
        w._generar_excel(str(tmp_path / "l.xlsx")); ws = load_workbook(tmp_path / "l.xlsx")["Valores proyectados"]
        columnas, _ = w.plan_excel_valores(); n = {k: i for i, (k, _) in enumerate(columnas, 1)}
        assert "ROUND" not in str(ws.cell(3, n["factor"]).value)  # el legado conserva I/I0 en Excel
        assert w.valores == antes  # abrir, mostrar y exportar no cambian nada
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=False):
            assert not w.actualizar_factor_legado(r0["calculo_id"]) and w.valores == antes
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=True) as pregunta, \
             patch("app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis", side_effect=AssertionError):
            assert w.actualizar_factor_legado(r0["calculo_id"])
        assert len(pregunta.call_args[0][0]) == 1  # vista previa tabular antes → después de esa actividad
        nuevo = w.valores[0]
        p = (Decimal(repr(r0["i_icociv"])) / Decimal(repr(r0["i0_icociv"])) - 1) * 100
        p = p.quantize(Decimal("0.01"), ROUND_HALF_UP)
        assert nuevo["criterio_factor"] == CRITERIO_FACTOR and nuevo["diferencia_porcentual"] == float(p)
        assert nuevo["i0_icociv"] == r0["i0_icociv"] and nuevo["calculo_id"] == r0["calculo_id"]
        assert nuevo["historial_ediciones"][-1]["antes"]["factor_icociv"] == r0["factor_icociv"]
        w.pila_deshacer.undo(); assert w.valores[0] is r0  # reversible
    finally:
        w.close()


# ---------------------------------------------------------------- A: clic en puntos de la gráfica
class _Visor(VisorGraficaDialog):
    def exec(self):
        return 0


def _visor(tema="claro"):
    v = ventana_de_prueba(); v._aplicar_tema(tema, persistir=False)
    analisis = reales()["historial_analisis"][0]; resultado = recuperar_resultado(analisis["resultado"])
    from matplotlib.figure import Figure
    from app_icociv.interfaz.ventana_principal import paleta_tema
    figura = Figure(figsize=(6, 4), dpi=100, facecolor=paleta_tema(tema)["grafica"])
    v._dibujar_grafica(figura, resultado["serie_df"], resultado["proyeccion"])
    visor = _Visor(v, figura, "Gráfica"); visor.resize(900, 640); visor.show(); APP.processEvents(); visor.canvas.draw()
    return v, visor, resultado


def _pixel(visor, nombre, k):
    linea = next(l for l in visor.eje.get_lines() if l.get_label() == nombre)
    x, y = float(linea.get_xdata()[k]), float(linea.get_ydata()[k])
    return visor.eje.transData.transform((x, y)), x, y


def _clic(visor, px, py):
    for tipo in ("button_press_event", "button_release_event"):
        MouseEvent(tipo, visor.canvas, px, py, button=1)._process()


def test_A_clic_muestra_mes_indice_y_tipo_y_reutiliza_una_etiqueta():
    for tema in ("claro", "oscuro"):
        v, visor, resultado = _visor(tema)
        try:
            periodos = visor.eje.xaxis.get_major_locator().periodos
            (px, py), x, y = _pixel(visor, "Serie observada", 5)
            _clic(visor, px + 3, py - 2)  # clic real (press + release) cerca del punto
            a = visor._anotacion
            assert a.get_visible() and mes_legible(periodos[round(x)]) in a.get_text() and "Histórico observado" in a.get_text()
            assert f"{y:,.3f}" in a.get_text() and y == resultado["serie_df"]["Indice"].iloc[5]
            (px, py), x, y = _pixel(visor, "Proyección", -1)
            _clic(visor, px, py)
            assert visor._anotacion is a and "Proyección" in a.get_text() and mes_legible(periodos[round(x)]) in a.get_text()
            assert a.get_ha() == "right"  # cerca del borde derecho: se orienta hacia dentro
            caja = visor.eje.get_window_extent(); _clic(visor, caja.x0 + 3, caja.y1 - 3)  # zona vacía
            assert not a.get_visible()
            _clic(visor, *_pixel(visor, "Serie observada", 0)[0]); assert a.get_visible()
            QTest.keyClick(visor, Qt.Key.Key_Escape); assert not a.get_visible() and visor.isVisible()  # Esc retira la etiqueta
            assert len([t for t in visor.eje.texts if t.get_visible()]) == 0 and len(visor.eje.texts) <= 1  # nunca acumula
            # Arrastrar (desplazar) no abre etiquetas.
            px, py = _pixel(visor, "Serie observada", 3)[0]
            MouseEvent("button_press_event", visor.canvas, px, py, button=1)._process()
            MouseEvent("motion_notify_event", visor.canvas, px + 40, py)._process()
            MouseEvent("button_release_event", visor.canvas, px + 40, py, button=1)._process()
            assert not a.get_visible()
            visor.grab().save(str(ROOT / f"docs/fase_4i_c/capturas/_prueba_anotacion_{tema}.png")) if (ROOT / "docs/fase_4i_c/capturas").exists() else None
        finally:
            visor.close(); v.close()
    assert mes_legible("2026-12") == "Diciembre de 2026" and mes_legible("2027_1") == "Enero de 2027" and mes_legible("x") == "x"


# ---------------------------------------------------------------- B: carga minimalista
def test_B_cuadro_de_carga_solo_titulo_y_barra():
    from app_icociv.interfaz.componentes import VeloCarga
    from app_icociv.interfaz.animaciones.transiciones import establecer_movimiento_reducido
    establecer_movimiento_reducido(False)
    base = QWidget(); base.resize(500, 300); base.show(); velo = VeloCarga(base)
    try:
        velo.mostrar("Ejecutando proyección…", "La aplicación sigue respondiendo."); QTest.qWait(3500)
        visibles = [l.text() for l in velo.findChildren(QLabel) if l.isVisible()]
        # Fase V1.1.2-01: la barra es la franja indeterminada (sin porcentaje inventado).
        assert visibles == ["Ejecutando proyección…"] and velo.franja.isVisible()
        assert velo.franja.porcentaje() is None and velo.panel.sizeHint().height() < 140
        velo.ocultar(exito=False); assert velo.franja.porcentaje() is None  # error: sin 100 %
        velo.mostrar("Cargando archivo…"); QTest.qWait(100); velo.ocultar(); QTest.qWait(200)
        assert not velo.isVisible()  # tarea breve: nunca apareció
    finally:
        base.close()
