"""Verificación 4I-D en la VENTANA real: abrir sesión con factores anteriores → aviso arriba → convertir por lote
→ guardar → reabrir → Excel. Datos SINTÉTICOS con los factores de la captura; el motor queda bloqueado."""
import json
import os
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtCore import QCoreApplication  # Fase 4I-G: la apertura de sesiones termina en el bucle de eventos
from PySide6.QtWidgets import QFileDialog, QMessageBox
from openpyxl import load_workbook

from tests.test_fase_4i_d_legados_factor import APP, sesion_legada
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from tests.test_proyecciones_redisenio_visual import ventana_de_prueba
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.persistencia.gestor_sesiones import guardar_sesion
from app_icociv.servicios.actualizacion_icociv import CRITERIO_FACTOR, es_factor_legado

MOTOR = "app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis"


def _abrir(ruta):
    v = ventana_de_prueba(); v.resize(1366, 768); v.show()
    with patch.object(VentanaPrincipal, '_corte_dane_cargado', return_value='9999-12'), patch.object(QFileDialog, "getOpenFileName", return_value=(str(ruta), "")), patch(MOTOR, side_effect=AssertionError("motor")):
        v.abrir_sesion()
        while v._operacion_sesion: QCoreApplication.processEvents()
    v._cambiar_seccion(1); APP.processEvents()
    return v, v.widget_proyecciones


def _factores(w):
    m = w.modelo_valores
    return [m.data(m.index(i, m.claves.index("factor"))) for i in range(m.rowCount())]


def test_ventana_abrir_convertir_guardar_reabrir_y_excel(tmp_path):
    w0, datos = sesion_legada(); w0.close()
    legada = guardar_sesion(tmp_path / "legada.json", {"usuario": "prueba", "archivo_excel": "", "proyecciones_icociv": datos})
    v, w = _abrir(legada)
    try:
        assert _factores(w)[:3] == ["1.027413", "1.023369", "1.030601"]
        # El aviso está ARRIBA del módulo (visible al abrir), además de junto a la tabla.
        assert w.aviso_legados_superior.isVisibleTo(w) and w.aviso_legados.isVisibleTo(w)
        assert w.aviso_legados_superior.mapTo(w, w.aviso_legados_superior.rect().topLeft()).y() < 250
        assert w.boton_actualizar_legados_superior.text() == "Actualizar factores de 6 actividades"
        # Filtro activo con un subconjunto: la conversión es de TODA la sesión y no pierde filas ocultas.
        w.solo_analisis_activo.setChecked(True); APP.processEvents()
        ocultas = len(w.valores) - len(w._valores_visibles)
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=False):
            w.boton_actualizar_legados_superior.click()
        assert all(es_factor_legado(r) for r in w.valores)  # cancelar: nada cambia
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_actualizacion", return_value=True) as pregunta, patch(MOTOR, side_effect=AssertionError("motor")):
            w.boton_actualizar_legados_superior.click(); APP.processEvents()
        assert len(pregunta.call_args[0][0]) == 6 and all("valor_proyectado" in n for _, n in pregunta.call_args[0][0])
        assert len(w.valores) == 6 and len(w.valores) - len(w._valores_visibles) == ocultas
        assert not any(es_factor_legado(r) for r in w.valores) and not w.aviso_legados_superior.isVisibleTo(w)
        w.solo_analisis_activo.setChecked(False); APP.processEvents()
        assert _factores(w) == ["1.0274", "1.0234", "1.0306", "1.0236", "1.0200", "0.9725"]
        convertida = tmp_path / "convertida.json"
        with patch.object(v, "_solicitar_usuario", return_value="prueba"), \
             patch.object(QFileDialog, "getSaveFileName", return_value=(str(convertida), "")):
            v.guardar_sesion_actual(como=True)
        valores = json.loads(convertida.read_text(encoding="utf-8"))["proyecciones_icociv"]["valores"]
        assert [r["criterio_factor"] for r in valores] == [CRITERIO_FACTOR] * 6
        assert [r["factor_icociv"] for r in valores] == [1.0274, 1.0234, 1.0306, 1.0236, 1.02, 0.9725]
        assert valores[0]["historial_ediciones"][-1]["antes"]["factor_icociv"] == datos["valores"][0]["factor_icociv"]
    finally:
        v.close()
    v2, w2 = _abrir(convertida)
    try:
        assert _factores(w2) == ["1.0274", "1.0234", "1.0306", "1.0236", "1.0200", "0.9725"]
        assert not w2.aviso_legados_superior.isVisibleTo(w2) and not w2.actualizar_legados_lote()  # sin conversión repetida
        assert [r["valor_proyectado"] for r in w2.valores] == [102740.0, 102340.0, 103060.0, 81888.0, 1020000.0, 97250.0]
        w2._generar_excel(str(tmp_path / "reabierta.xlsx")); libro = load_workbook(tmp_path / "reabierta.xlsx")
        ws = libro["Valores proyectados"]; columnas, _ = w2.plan_excel_valores(); n = {k: i for i, (k, _) in enumerate(columnas, 1)}
        assert len(libro.sheetnames) == 5
        assert all(ws.cell(f, n["factor"]).value.startswith("=1+ROUND(") for f in range(3, 9))
        libro.close()
    finally:
        v2.close()
