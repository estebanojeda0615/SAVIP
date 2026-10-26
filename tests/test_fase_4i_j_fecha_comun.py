"""Fase 4I-J V2: trayectoria de 24 meses consultable sin motor, fecha objetivo común del presupuesto,
excepción individual avisada, advertencia antes de exportar y reanálisis SOLO por botón con DANE nuevo.

Motor real sobre anexos sintéticos: «antiguo» con corte 2026-10 y «nuevo» con corte 2026-12 (mismos datos
hasta 2026-10, como un anexo posterior del DANE)."""

import json
import math
import os
import time
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtGui import QFontDatabase
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QFileDialog, QMessageBox, QScrollArea

from tests.test_slider_comparacion_historica_fase_4d import APP, reales
from app_icociv.interfaz.controladores.controlador_principal import ControladorPrincipal
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.servicios import fecha_objetivo as fo
from app_icociv.utilidades.utilidades import COLUMNAS_FIJAS_TABLAS
from tests.fixtures.salidas import carpeta_salida

for _fuente in ("segoeui.ttf", "segoeuib.ttf", "seguisb.ttf", "seguisym.ttf"):
    QFontDatabase.addApplicationFont(str(Path("C:/Windows/Fonts") / _fuente))
CAPTURAS = carpeta_salida("fase_4i_j_capturas")  # Fase 1A (AS-007): antes en docs/, no publicado
_MOTOR_REAL = ControladorPrincipal.ejecutar_analisis


def _anexo(ruta, meses):
    """Anexo sintético desde 2023-01 con `meses` periodos (46 → corte 2026-10; 48 → 2026-12)."""
    libro = Workbook(); hoja = libro.active; hoja.title = "Anexo 16"
    periodos = [f"{2023 + i // 12}_{i % 12 + 1}" for i in range(meses)]
    for tag, columnas in COLUMNAS_FIJAS_TABLAS.items():
        hoja.append([tag + ". PRUEBA - datos sintéticos"]); hoja.append(columnas + periodos)
        for grupo in (1, 2):
            hoja.append([100.0 if "Ponderacion" in c else grupo if "cod" in c.lower() else f"{c} sintético {grupo}" for c in columnas]
                        + [100 + grupo + i * .2 + math.sin(i / 3) * .1 for i in range(meses)])
        for _ in range(3):
            hoja.append([None])
    libro.save(ruta)
    return ruta


def _esperar(v, condicion, limite=120):
    fin = time.monotonic() + limite
    while time.monotonic() < fin:
        APP.processEvents()
        if condicion() and (v.trabajador is None or not v.trabajador.isRunning()):
            return
        QTest.qWait(20)
    raise AssertionError("No terminó: " + v.etiqueta_estado.text())


def _cargar(v, ruta):
    v._archivo_cargado(v.controlador.cargar_archivo(str(ruta))); APP.processEvents()


def _proyectar(v, grupo, horizonte):
    w = v.widget_proyecciones; n = len(w.historial)
    v.combo_grupo.setCurrentIndex(grupo)
    v.combo_horizonte.setCurrentIndex(v.combo_horizonte.findData(None))
    v.spin_horizonte_personalizado.setValue(horizonte)
    v.ejecutar_proyeccion()
    _esperar(v, lambda: len(w.historial) == n + 1)
    return w.historial[-1]


def _calcular(w, nombre):
    w.tipo_componente.setCurrentText("Actividad"); w.nombre_componente.setText(nombre); w.item.setText(nombre)
    w._item_editado = True; w.unidad.setCurrentIndex(1); w.cantidad.setValue(2); w.precio_base.setValue(1000)
    w.anticipo.setValue(100); w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData("unitario"))
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


_APLICAR = WidgetProyeccionesICOCIV.aplicar_fecha_objetivo


def _fecha(v, anio, mes, alcance="todas"):
    """Escenario de SOLO consulta (fecha común con actividades aún en otra fecha, como tras reutilizar una
    trayectoria); la actualización confirmada de actividades se prueba en test_fase_4i_j_k_actividades."""
    v.spin_anio.setValue(anio); v.spin_mes.setValue(mes)
    v.combo_alcance_fecha.setCurrentIndex(v.combo_alcance_fecha.findData(alcance))
    solo_consulta = lambda self, periodo, solo_id=None, confirmar=True, actualizar_actividades=True: _APLICAR(self, periodo, solo_id, confirmar, False)
    with patch.object(WidgetProyeccionesICOCIV, "aplicar_fecha_objetivo", solo_consulta):
        return v.aplicar_fecha_objetivo()


def test_presupuesto_fecha_comun_excepcion_exportacion_y_reanalisis_dane(tmp_path):
    antiguo, nuevo = _anexo(tmp_path / "anex-ICOCIV-oct2026.xlsx", 46), _anexo(tmp_path / "anex-ICOCIV-dic2026.xlsx", 48)
    v = VentanaPrincipal(); v.resize(1366, 900); v.show(); w = v.widget_proyecciones
    try:
        # Serie A con DANE de corte 2026-10 (H9 = 2027-07) y serie B con DANE de corte 2026-12 (H7 = 2027-07).
        _cargar(v, antiguo)
        a = _proyectar(v, 1, 9)
        assert fo.origen_analisis(a) == "2026-10" and a["contexto"]["periodo_proyectado"] == "2027-07"
        tray_a = fo.trayectoria(a)
        assert len(tray_a) == 24 and a["contexto"]["indice_proyectado"] == tray_a[8]  # 1) H9 = punto 9 de 24
        for i in range(10):
            _calcular(w, f"Actividad A{i + 1}")
        historial_a = deepcopy(a)
        # 7) Cargar un DANE más reciente NO recalcula nada.
        with patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")):
            _cargar(v, nuevo)
        assert w.historial == [historial_a] and len(w.valores) == 10
        b = _proyectar(v, 2, 7)
        assert fo.origen_analisis(b) == "2026-12" and b["contexto"]["periodo_proyectado"] == "2027-07"
        _calcular(w, "Actividad B1")
        valores = deepcopy(w.valores); n_hist = len(w.historial)

        # 2) y 3) Fecha común 2027-09 en todas las series: sin motor, sin historial nuevo, actividades intactas.
        with patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")) as motor:
            assert _fecha(v, 2027, 9)
        assert motor.call_count == 0 and len(w.historial) == n_hist and w.valores == valores
        assert w.fecha_objetivo_presupuesto == "2027-09"
        ca, cb = w.historial[0]["consulta_horizonte"], w.historial[1]["consulta_horizonte"]
        assert (ca["horizonte"], cb["horizonte"]) == (11, 9) and ca["indice"] == tray_a[10] and cb["indice"] == fo.trayectoria(b)[8]
        assert w.historial[0]["contexto"] == historial_a["contexto"]  # el análisis original no se reescribe
        # 5) Las 11 actividades siguen en 2027-07: se informa la discrepancia y un cálculo nuevo usa la fecha consultada.
        assert w.actividades_en_otra_fecha() == 11
        nueva = _calcular(w, "Actividad B2 (nueva fecha)")
        assert (nueva["periodo_proyectado"], nueva["i_icociv"], nueva["horizonte"]) == ("2027-09", cb["indice"], 9)
        w._registrar_cambio("quitar prueba", w.valores[:-1], None); assert w.valores == valores

        # 4) Excepción solo para B: Cancelar no cambia nada; confirmar la registra y la muestra.
        v._activar_analisis(b["id"], conservar_formulario=True)
        antes = deepcopy(w.historial)
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_fecha_individual", return_value=False) as aviso:
            assert not _fecha(v, 2027, 11, "individual")
        assert aviso.call_count == 1 and w.historial == antes and not w.excepciones_fecha
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_fecha_individual", return_value=True) as aviso:
            assert _fecha(v, 2027, 11, "individual")
        assert aviso.call_args[0][1:] == ("2027-11", "2027-09") and w.historial[1]["consulta_horizonte"]["horizonte"] == 11
        assert w.excepciones_fecha[-1]["analisis_id"] == b["id"] and "Fuera de la fecha común" in v.lbl_fecha_presupuesto.text()
        assert w.historial[0]["consulta_horizonte"]["periodo"] == "2027-09"  # A no cambia
        with patch.object(WidgetProyeccionesICOCIV, "_confirmar_fecha_individual", side_effect=AssertionError("falsa alerta")):
            assert _fecha(v, 2027, 9, "individual")  # vuelve a la fecha común: sin advertencia
        # Fecha fuera de la trayectoria de alguna serie: no se aplica en ninguna.
        with patch.object(QMessageBox, "warning") as aviso:
            assert not _fecha(v, 2028, 12)
        assert "máximo es 24" in aviso.call_args[0][2] and w.fecha_objetivo_presupuesto == "2027-09"

        # 6) Exportación: discrepancia (11 actividades en 2027-07 vs presupuesto 2027-09) → aviso; cancelar no exporta.
        with patch.object(w, "_elegir_alcance_exportacion", return_value=w.alcance_exportacion("sesion")), \
             patch.object(w, "_dialogo_exportacion", return_value=False) as dialogo, \
             patch.object(QFileDialog, "getSaveFileName", side_effect=AssertionError("no debía pedir archivo")):
            for formato in ("xlsx", "pdf", "docx"):
                w._exportar_formato(formato)
        assert dialogo.call_count == 3 and "requieren actualización expresa" in dialogo.call_args[0][0]
        # Mixto 10 + 1: el aviso enumera ambas fechas; continuar exporta el Excel con sus 5 hojas intactas.
        w.valores[-1]["periodo_proyectado"] = "2027-09"
        destino = tmp_path / "mixto.xlsx"
        with patch.object(w, "_elegir_alcance_exportacion", return_value=w.alcance_exportacion("sesion")), \
             patch.object(w, "_dialogo_exportacion", return_value=True) as dialogo, \
             patch.object(QFileDialog, "getSaveFileName", return_value=(str(destino), "")), patch.object(QMessageBox, "information"):
            w._exportar_formato("xlsx")
        assert "jul-2027 (10 actividades) y sep-2027 (1 actividad)" in dialogo.call_args[0][0]
        assert load_workbook(destino).sheetnames == ['Resumen', 'Valores proyectados', 'Proyecciones ICOCIV', 'Trazabilidad', 'Metodología y fórmulas']
        w.valores[-1]["periodo_proyectado"] = "2027-07"
        # Alcance restringido homogéneo con fecha de presupuesto igual: sin falsa alarma.
        w.fecha_objetivo_presupuesto = "2027-07"
        uno = w.alcance_exportacion("proyeccion", analisis_id=a["id"])  # las 10 actividades de A, todas en 2027-07
        assert len(uno["calculos"]) == 10
        with patch.object(w, "_elegir_alcance_exportacion", return_value=uno), patch.object(w, "_confirmar_avisos_alcance", return_value=True), \
             patch.object(w, "_dialogo_exportacion", side_effect=AssertionError("falsa alarma")), \
             patch.object(QFileDialog, "getSaveFileName", return_value=("", "")):
            w._exportar_formato("xlsx")
        w.fecha_objetivo_presupuesto = "2027-09"

        # 10) Persistencia: guardar y reabrir conserva fecha común, consultas, excepciones y actividades.
        ruta = tmp_path / "presupuesto.json"
        with patch.object(v, "_solicitar_usuario", return_value="QA"), patch.object(QFileDialog, "getSaveFileName", return_value=(str(ruta), "")):
            assert v.guardar_sesion_actual(como=True)
        otra = VentanaPrincipal()
        try:
            _cargar(otra, nuevo)
            with patch.object(QFileDialog, "getOpenFileName", return_value=(str(ruta), "")), \
                 patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")):
                otra.abrir_sesion()
                while otra._operacion_sesion:
                    APP.processEvents()
            o = otra.widget_proyecciones
            assert o.fecha_objetivo_presupuesto == "2027-09" and o.excepciones_fecha == w.excepciones_fecha
            assert [x.get("consulta_horizonte") for x in o.historial] == [x.get("consulta_horizonte") for x in w.historial]
            assert o.valores == w.valores and o.contexto["periodo_proyectado"] == o.analisis_activo()["consulta_horizonte"]["periodo"]
            assert o.datos_sesion()["comparacion_historica"] == w.datos_sesion()["comparacion_historica"]  # slider I0 intacto
        finally:
            otra.close()

        # 9) Reanálisis cancelado o fallido: nada cambia.
        estado = deepcopy(w.datos_sesion())
        with patch.object(v, "_confirmar_reanalisis", return_value=False), \
             patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=AssertionError("motor")):
            v.actualizar_analisis_dane()
        with patch.object(v, "_confirmar_reanalisis", return_value=True), \
             patch.object(ControladorPrincipal, "ejecutar_analisis", side_effect=ValueError("fallo simulado")), \
             patch.object(QMessageBox, "warning") as aviso:
            v.actualizar_analisis_dane(); _esperar(v, lambda: aviso.called)
        assert "No se creó ninguna revisión" in aviso.call_args[0][2] and w.datos_sesion() == estado

        # 8) Reanálisis explícito: solo A (corte 2026-10 < 2026-12) → revisión nueva desde 2026-12; A y actividades intactas.
        llamadas = []
        espia = lambda self, *a, **k: (llamadas.append(a[0]), _MOTOR_REAL(self, *a, **k))[1]
        previa = deepcopy(w.historial[0])
        with patch.object(v, "_confirmar_reanalisis", return_value=True) as confirmar, \
             patch.object(ControladorPrincipal, "ejecutar_analisis", espia), patch.object(QMessageBox, "information"):
            v.actualizar_analisis_dane(); _esperar(v, lambda: len(w.historial) == n_hist + 1)
        assert "elegibles: 1; no elegibles: 1" in confirmar.call_args[0][0] and len(llamadas) == 1
        r = w.historial[-1]
        assert (r["revision_de"], r["revision"], r["corte_anterior"], r["corte_nuevo"]) == (a["id"], 2, "2026-10", "2026-12")
        assert len(fo.trayectoria(r)) == 24 and fo.origen_analisis(r) == "2026-12"  # nueva ventana H1=2027-01..H24=2028-12
        assert r["consulta_horizonte"]["periodo"] == "2027-09" and r["consulta_horizonte"]["horizonte"] == 9  # antes H11
        assert w.historial[0] == previa and w.valores == valores  # revisión anterior y economía intactas
        assert [x["id"] for x in fo.vigentes(w.historial)] == [b["id"], r["id"]]
        assert r["contexto"]["serie_final"] == previa["contexto"]["serie_final"]

        # 11) Bloque de parámetros en pantalla dividida: sin barra horizontal, claro y oscuro.
        CAPTURAS.mkdir(parents=True, exist_ok=True)
        v._cambiar_seccion(1)
        for tema in ("claro", "oscuro"):
            v._aplicar_tema(tema, persistir=False)
            for ancho in (1920, 1366, 960, 683):
                v.resize(ancho, 900); QTest.qWait(450)
                area = v.boton_aplicar_fecha.parent()
                while not isinstance(area, QScrollArea):
                    area = area.parent()
                area.ensureWidgetVisible(v.boton_reanalizar_dane, 10, 240); QTest.qWait(60)
                assert area.horizontalScrollBar().maximum() == 0
                for control in (v.combo_alcance_fecha, v.boton_aplicar_fecha, v.boton_reanalizar_dane):
                    assert control.mapTo(area.viewport(), control.rect().topRight()).x() <= area.viewport().width()
                v.grab().save(str(CAPTURAS / f"parametros_{ancho}_{tema}.png"))
    finally:
        v.close()


def test_sesion_antigua_sin_trayectoria_no_inventa_puntos():
    """12) Sesiones guardadas: no se migran ni se fabrican H; una trayectoria ausente bloquea la consulta."""
    datos = reales()
    sin_trayectoria = deepcopy(datos["historial_analisis"][0])
    sin_trayectoria["resultado"]["proyeccion"]["trayectoria_24_meses"] = []
    try:
        fo.consultar(sin_trayectoria, "2027-05")
    except ValueError as error:
        assert "guarda 0 meses" in str(error)
    else:
        raise AssertionError("se fabricó un punto")
    a = datos["historial_analisis"][0]
    assert fo.consultar(a, "2027-05")["indice"] == a["resultado"]["proyeccion"]["trayectoria_24_meses"][4]
    assert "consulta_horizonte" not in a and fo.vigentes(datos["historial_analisis"]) == datos["historial_analisis"]
