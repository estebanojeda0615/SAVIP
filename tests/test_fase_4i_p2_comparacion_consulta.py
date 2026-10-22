"""Fase 4I-P2: el panel de comparación histórica de I0 muestra la consulta vigente y calcula escenarios contra
ella. Una actividad es demostrable si su I es EXACTAMENTE el punto guardado de la trayectoria para su fecha
(objetivo original o consultado). Sin slider movido los importes no cambian; con slider movido el escenario se
aplica igual en fecha original y consultada. Sin motor; snapshots y cifras registradas intactos."""

import os
from copy import deepcopy
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QMessageBox

from tests.test_fase_4i_j_fecha_comun import APP, _calcular, _cargar
from tests.test_fase_4i_j_k_actividades import _aplicar
from tests.test_fase_4i_j_reutilizacion import _Contador, _anexo, _ejecutar
from app_icociv.interfaz import ventana_principal as vp
from app_icociv.interfaz.controladores.controlador_principal import ControladorPrincipal
from app_icociv.interfaz.widgets.comparacion_historica import escenario_registro
from app_icociv.reportes.sesion_modulos import _numero, construir_proyecciones
from app_icociv.servicios import fecha_objetivo as fo
from app_icociv.servicios.actualizacion_icociv import aplicar_factor

NO_DEMOSTRABLE = "vínculo no demostrable"


def _xlsx(par):
    """openpyxl serializa los float con '%.16g' (preexistente): se compara con esa misma precisión."""
    return tuple(float("%.16g" % x) for x in par)


def _salidas(v, w, ruta_xlsx):
    """Tabla, Excel (I0/I de «Valores proyectados»), informe PDF/DOCX (tablas I0 y Estado) y panel."""
    m = w.modelo_valores
    ci0, ci = m.claves.index("i0_icociv"), m.claves.index("i_icociv")
    tabla = [(m.index(i, ci0).data(Qt.UserRole), m.index(i, ci).data(Qt.UserRole)) for i in range(m.rowCount())]
    with patch.object(QFileDialog, "getSaveFileName", return_value=(str(ruta_xlsx), "")), \
         patch.object(QMessageBox, "information", return_value=QMessageBox.StandardButton.Ok):
        w._exportar_excel()
    ws = load_workbook(ruta_xlsx)["Valores proyectados"]
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    cab = next(i for i, f in enumerate(filas) if "I0 ICOCIV" in [str(c).split("\n")[0] for c in f if c])
    nombres = [str(c).split("\n")[0] if c else "" for c in filas[cab]]
    k0, k1 = nombres.index("I0 ICOCIV"), nombres.index("I ICOCIV proyectado")
    excel = [(f[k0], f[k1]) for f in filas[cab + 1:cab + 1 + len(w.valores)]]
    tablas = [b for s in construir_proyecciones(w).secciones for b in s.bloques if hasattr(b, "encabezados")]
    t_i0 = next(t for t in tablas if "I0 ICOCIV" in t.encabezados and "Valor (COP)" in t.encabezados)
    t_estado = next(t for t in tablas if t.encabezados == ["N.º", "Ajuste R (COP)", "Estado"])
    informe = [(f[2], f[3], f[5]) for f in t_i0.filas]
    return {"tabla": tabla, "excel": excel, "informe": informe, "estados": [f[2] for f in t_estado.filas],
            "escenarios": [w.escenario_registro(r) for r in w.valores],
            "panel": (w.lbl_objetivo_comparacion.text(), w.lbl_estado_comparacion.text()),
            "resultados": v.lbl_comparacion_resultados.text()}


def _sin_cambio(s, w):
    """Sin slider movido: tabla, Excel e informe muestran exactamente las cifras registradas."""
    registrados = [(r["i0_icociv"], r["i_icociv"]) for r in w.valores]
    assert s["tabla"] == registrados and s["excel"] == [_xlsx(x) for x in registrados], (s["tabla"], s["excel"], registrados)
    assert s["informe"] == [(_numero(i0), _numero(i), _numero(r["valor_proyectado"], 0))
                            for (i0, i), r in zip(registrados, w.valores)]
    assert all(e == r for (e, _), r in zip(s["escenarios"], w.valores))
    assert not any(NO_DEMOSTRABLE in x for x in s["estados"]), s["estados"]


def _con_escenario(s, w, referencia):
    """Slider movido: I0 = referencia y Valor = Base × F con la regla vigente (ROUND_HALF_UP 2 decimales)."""
    for (e, estado), r, fila_t, fila_x, fila_i in zip(s["escenarios"], w.valores, s["tabla"], s["excel"], s["informe"]):
        esperado = aplicar_factor(r["precio_base"] - r["anticipo_amortizado"], r["i_icociv"], referencia)
        assert estado == "Referencia modificada (escenario)", estado
        assert e["i0_icociv"] == referencia and e["valor_proyectado"] == esperado["valor"] and e["factor_icociv"] == esperado["factor"]
        assert fila_t == (referencia, r["i_icociv"]) and fila_x == _xlsx((referencia, r["i_icociv"]))
        assert fila_i == (_numero(referencia), _numero(r["i_icociv"]), _numero(esperado["valor"], 0))
    assert all(x == "Referencia modificada (escenario)" for x in s["estados"])


def test_panel_consulta_vigente_escenarios_y_persistencia(tmp_path):
    v = vp.VentanaPrincipal(); v.resize(1366, 900); v.show(); w = v.widget_proyecciones
    motor = _Contador()
    try:
        with patch.object(ControladorPrincipal, "ejecutar_analisis", motor.metodo):
            _cargar(v, _anexo(tmp_path / "anex-ICOCIV-ago2026.xlsx", meses=44))   # sintético, último 2026-08
            assert _ejecutar(v, 1, 11) == 1
            a = w.historial[0]; tray = fo.trayectoria(a)
            snapshot = deepcopy([a["nombre"], a["resultado"], a["contexto"], a["parametros"]])
            for nombre in ("A1", "A2"):
                _calcular(w, nombre)
            registrados_h11 = deepcopy(w.valores)

            # H11, slider en la referencia predeterminada: panel original, sin cambios de importes.
            s = _salidas(v, w, tmp_path / "h11.xlsx")
            assert s["panel"][0].startswith("Objetivo guardado: 2027-07") and "2027-07 · 11 meses" in s["panel"][1]
            _sin_cambio(s, w)

            # H23 aplicado (actividades actualizadas): el panel muestra la consulta y el vínculo es demostrable.
            assert _aplicar(v, 2028, 7)[0]
            s = _salidas(v, w, tmp_path / "h23.xlsx")
            assert s["panel"][0].startswith(f"Objetivo consultado: 2028-07 · H23 · Índice proyectado: ") and "ejecución original: 2027-07 · H11" in s["panel"][0]
            assert "2028-07 · 23 meses" in s["panel"][1] and "Objetivo consultado: 2028-07" in s["resultados"]
            assert all(r["i_icociv"] == tray[22] for r in w.valores)
            _sin_cambio(s, w)

            # Slider movido en H23 (consultada) y luego en H11 (original): mismo cálculo.
            h = w._historicos_comparacion[a["id"]]; mes = "2025-06"; ref = h[mes]
            assert w.seleccionar_referencia_historica(mes)
            _con_escenario(_salidas(v, w, tmp_path / "h23_slider.xlsx"), w, ref)
            assert _aplicar(v, 2027, 7)[0]
            _con_escenario(_salidas(v, w, tmp_path / "h11_slider.xlsx"), w, ref)
            assert _aplicar(v, 2028, 7)[0]
            _con_escenario(_salidas(v, w, tmp_path / "h23_slider_2.xlsx"), w, ref)

            # Deshacer / rehacer la fecha: el escenario sigue a las actividades (H11 y vuelta a H23).
            w.pila_deshacer.undo(); APP.processEvents()
            assert all(r["i_icociv"] == tray[10] for r in w.valores)
            _con_escenario(_salidas(v, w, tmp_path / "undo.xlsx"), w, ref)
            assert w.lbl_objetivo_comparacion.text().startswith("Objetivo consultado: 2027-07 · H11")
            w.pila_deshacer.redo(); APP.processEvents()
            _con_escenario(_salidas(v, w, tmp_path / "redo.xlsx"), w, ref)

            # Guardar y reabrir con el escenario guardado: mismas salidas en la ventana reabierta.
            antes = _salidas(v, w, tmp_path / "antes.xlsx")
            ruta = tmp_path / "presupuesto.json"
            with patch.object(v, "_solicitar_usuario", return_value="QA"), \
                 patch.object(QFileDialog, "getSaveFileName", return_value=(str(ruta), "")):
                assert v.guardar_sesion_actual(como=True)
            otra = vp.VentanaPrincipal()
            try:
                _cargar(otra, tmp_path / "anex-ICOCIV-ago2026.xlsx")
                with patch.object(QFileDialog, "getOpenFileName", return_value=(str(ruta), "")):
                    otra.abrir_sesion()
                    while otra._operacion_sesion:
                        APP.processEvents()
                w2 = otra.widget_proyecciones; otra._activar_analisis(w2.historial[0]["id"])
                despues = _salidas(otra, w2, tmp_path / "reabierta.xlsx")
                assert {k: despues[k] for k in ("tabla", "excel", "informe", "estados", "panel")} == \
                       {k: antes[k] for k in ("tabla", "excel", "informe", "estados", "panel")}
                _con_escenario(despues, w2, ref)
                assert w2.valores == w.valores
            finally:
                otra.close()

            # Volver a la referencia predeterminada: las cifras registradas reaparecen tal cual.
            assert w.seleccionar_referencia_historica(next(reversed(h)))
            _sin_cambio(_salidas(v, w, tmp_path / "default.xlsx"), w)

            # Cifras registradas, snapshots y motor intactos: el slider nunca modifica los registros.
            assert deepcopy([a["nombre"], a["resultado"], a["contexto"], a["parametros"]]) == snapshot
            assert [(r["i0_icociv"], r["precio_base"], r["anticipo_amortizado"], r["cantidad"]) for r in w.valores] == \
                   [(r["i0_icociv"], r["precio_base"], r["anticipo_amortizado"], r["cantidad"]) for r in registrados_h11]
            assert motor.n == 1
    finally:
        v.close()


def test_registro_con_indice_ajeno_a_la_trayectoria_sigue_no_demostrable(tmp_path):
    """La regla no se relaja: un I que no es el punto guardado de su fecha sigue sin vínculo demostrable."""
    v = vp.VentanaPrincipal(); w = v.widget_proyecciones
    try:
        _cargar(v, _anexo(tmp_path / "a.xlsx", meses=44))
        assert _ejecutar(v, 1, 11) == 1
        r = _calcular(w, "A1")
        alterado = {**r, "i_icociv": r["i_icociv"] + 0.001}
        assert NO_DEMOSTRABLE in w.escenario_registro(alterado)[1]
        otra_fecha = {**r, "periodo_proyectado": "2028-07"}                    # fecha consultable, I de otra fecha
        assert NO_DEMOSTRABLE in w.escenario_registro(otra_fecha)[1]
        fuera = {**r, "periodo_proyectado": "2029-12", "i_icociv": r["i_icociv"]}  # fuera de la trayectoria (H40)
        assert NO_DEMOSTRABLE in w.escenario_registro(fuera)[1]
        # Referencia (p. ej. DANE posterior) no anterior al objetivo de la actividad: no se compara ni falla.
        a = w.historial[0]; h = dict(w._historicos_comparacion[a["id"]]); h["2027-09"] = next(reversed(h.values()))
        e, estado = escenario_registro(r, a, h, "2027-09", True)
        assert e == r and "no es anterior al objetivo 2027-07" in estado
    finally:
        v.close()
