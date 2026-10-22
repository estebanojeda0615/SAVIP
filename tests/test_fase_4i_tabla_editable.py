"""Fase 4I: tabla ICOCIV editable, orden presupuestal, inserción, formatos, pegado, versión y logo.

Pruebas focalizadas con Qt real y fixtures de sesión anteriores (sin motor estadístico).
"""
import json
import os
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox
from openpyxl import load_workbook

from tests.test_slider_comparacion_historica_fase_4d import APP, reales, sintetico
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV

ROOT = Path(__file__).resolve().parents[1]


def _actividad(w, item, p, cantidad=2, modalidad="unitario"):
    w._formulario_guardado = None
    w.item.setText(item); w._item_editado = True; w.unidad.setCurrentIndex(1)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modalidad))
    w.cantidad.setValue(cantidad); w.precio_base.setValue(p); w.anticipo.setValue(0)
    with patch("app_icociv.interfaz.widgets.controles.confirmar_ingreso_presupuestal", return_value=True):
        w.boton_calcular.click()
    return w.valores[-1]


def _ids(w):
    return [r["calculo_id"] for r in w.valores]


def _widget_con_actividades(n=4):
    w = sintetico(); w.valores.clear(); w._refrescar_tabla_valores()
    for i in range(1, n + 1):
        _actividad(w, f"Actividad {i}", 1000 * i)
    return w


def test_A_insertar_en_posicion_y_persistir_orden(tmp_path):
    w = _widget_con_actividades(4)
    try:
        ids = _ids(w)
        assert w.preparar_insercion(ids[0], despues=True) and w.lbl_insercion.isVisibleTo(w)
        _actividad(w, "Actividad 1b", 1500)
        nueva = next(i for i in _ids(w) if i not in ids)
        assert _ids(w) == [ids[0], nueva, *ids[1:]] and w._insercion is None
        datos = deepcopy(w.datos_sesion()); otro = WidgetProyeccionesICOCIV(); otro.restaurar_datos_sesion(datos)
        try:
            assert _ids(otro) == _ids(w) and [r["numero"] for _, r in []] == []
            p = tmp_path / "orden.xlsx"; otro._generar_excel(str(p)); libro = load_workbook(p)
            t = libro["Trazabilidad"]; assert [t.cell(f, 2).value for f in range(3, 3 + len(ids) + 1)] == _ids(w)
            libro.close()
        finally:
            otro.close()
    finally:
        w.close()


def test_B_mover_por_arrastre_y_alternativas_sin_recalcular():
    w = _widget_con_actividades(4)
    try:
        sin_n = lambda r: {k: v for k, v in r.items() if k != "numero_fila"}  # número visible de fila, no dato
        antes = {r["calculo_id"]: sin_n(r) for r in w.valores}; ids = _ids(w); borrador = w._datos_calculadora()
        with patch("app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv", side_effect=AssertionError("No recalcular")):
            w.tabla_valores.verticalHeader().moveSection(3, 0)  # arrastre del número de fila
            assert _ids(w) == [ids[3], ids[0], ids[1], ids[2]]
            assert w.tabla_valores.verticalHeader().visualIndex(0) == 0  # la cabecera vuelve a su orden natural
            assert w.mover_actividad(ids[3], 3) and _ids(w) == ids
            assert w.mover_actividad(ids[1], 0) and _ids(w)[0] == ids[1]
            with patch("app_icociv.interfaz.widgets.proyecciones_icociv.QInputDialog.getInt", return_value=(4, True)):
                w._pedir_posicion(ids[1])
            assert _ids(w)[-1] == ids[1]
        assert {r["calculo_id"]: sin_n(r) for r in w.valores} == antes and w._datos_calculadora() == borrador
    finally:
        w.close()


def test_C_reordenar_con_filtro_respeta_filas_ocultas():
    w = WidgetProyeccionesICOCIV(); w.restaurar_datos_sesion(reales())
    try:
        extra = deepcopy(w.valores[0]); extra["calculo_id"] = "extra001"; w.valores.insert(1, extra)
        for a in w.historial: a["registros_ids"] = [r["calculo_id"] for r in w.valores if r.get("analisis_id") == a["id"]]
        w._preparar_comparacion_historica(); w._refrescar_historial(); w._refrescar_tabla_valores()
        antes = _ids(w); objetivo = w.valores[0]["analisis_id"]
        w.selector_analisis_valores.setCurrentIndex(w.selector_analisis_valores.findData(objetivo)); w.solo_analisis_activo.setChecked(True)
        visibles = [r["calculo_id"] for r in w._valores_visibles]; assert len(visibles) == 2
        w.mover_actividad(visibles[1], 0)
        assert _ids(w) == [antes[1], antes[0], *antes[2:]]  # solo cambian los huecos de las filas visibles
        w.solo_analisis_activo.setChecked(False); assert sorted(_ids(w)) == sorted(antes)
    finally:
        w.close()


def test_D_insercion_cancelada_o_sin_analisis_no_crea_nada():
    w = _widget_con_actividades(2)
    try:
        ids = _ids(w)
        assert w.preparar_insercion(ids[0]); w.cancelar_insercion()
        assert _ids(w) == ids and w._insercion is None and not w.lbl_insercion.isVisibleTo(w)
        w.analisis_activo_id = None
        with patch.object(QMessageBox, "information") as aviso:
            assert not w.preparar_insercion(ids[0])
        assert aviso.called and w._insercion is None and _ids(w) == ids
    finally:
        w.close()


def test_E_edicion_directa_valida_invalida_y_solo_lectura():
    w = _widget_con_actividades(2)
    try:
        m = w.modelo_valores; col = m.claves.index; r0 = deepcopy(w.valores[0]); borrador = w._datos_calculadora()
        for clave in ("i0_icociv", "i_icociv", "factor", "valor_unitario", "diferencia"):
            assert not (m.flags(m.index(0, col(clave))) & Qt.ItemIsEditable)
        assert m.setData(m.index(0, col("item")), "Ítem renombrado", Qt.EditRole)
        assert m.setData(m.index(0, col("precio_base")), "1.250.000", Qt.EditRole)
        r = w.valores[0]
        assert (r["calculo_id"], r["analisis_id"], r["item"], r["precio_base"]) == (r0["calculo_id"], r0["analisis_id"], "Ítem renombrado", 1250000.0)
        assert r["base_ajustable"] == 1250000.0 and r["valor_proyectado"] == 1250000.0 * (r0["i_icociv"] / r0["i0_icociv"])
        assert len(r["historial_ediciones"]) == 2 and r["historial_ediciones"][1]["antes"] == {"precio_base": r0["precio_base"]}
        antes = deepcopy(w.valores[0])
        with patch.object(w, "_avisar_edicion") as aviso:
            assert not m.setData(m.index(0, col("anticipo")), "2.000.000", Qt.EditRole)  # A >= P
            assert not m.setData(m.index(0, col("cantidad")), "abc", Qt.EditRole)
            assert not m.setData(m.index(0, col("unidad")), "galones imperiales", Qt.EditRole)
        assert aviso.call_count == 3 and w.valores[0] == antes and w._datos_calculadora() == borrador
    finally:
        w.close()


def test_F_pegar_en_tabla_y_formatos_de_pegado():
    from app_icociv.interfaz.widgets.numerico_icociv import _pegado
    assert [float(_pegado(t)) for t in ("1250000", "1.250.000", "1.250.000,50", "1,250,000.50", "$ 1.250.000", "1 250 000")] == \
        [1250000, 1250000, 1250000.5, 1250000.5, 1250000, 1250000]
    # Fase 4I-A V4: «1.250» = 1.25 y «1,234» = 1234 son ahora formatos aceptados.
    assert [float(_pegado(t)) for t in ("1.250", "1,234")] == [1.25, 1234]
    for malo in ("12 a", "$$"):
        try:
            _pegado(malo)
        except ValueError:
            pass
        else:
            raise AssertionError(malo)
    w = _widget_con_actividades(1)
    try:
        m = w.modelo_valores; index = m.index(0, m.claves.index("precio_base")); w.tabla_valores.setCurrentIndex(index)
        QApplication.clipboard().setText("1.250.000,50"); assert w.tabla_valores.pegar_en_celda()
        assert w.valores[0]["precio_base"] == 1250000.5
        w.tabla_valores.setCurrentIndex(m.index(0, m.claves.index("precio_base"))); QApplication.clipboard().setText("1\t2\n3\t4")
        from PySide6.QtWidgets import QToolTip  # Fase 4I-A: aviso no modal junto a la celda
        with patch.object(QToolTip, "showText") as aviso:
            assert not w.tabla_valores.pegar_en_celda()
        assert aviso.called and w.valores[0]["precio_base"] == 1250000.5
    finally:
        w.close()


def test_G_formatos_pesos_y_porcentaje_en_tabla_resultado_y_excel(tmp_path):
    from app_icociv.interfaz.formato_numerico import pesos_ui, porcentaje_ui
    assert pesos_ui(1234567.5) == "1,234,568" and porcentaje_ui(3.094782) == "3.09"
    w = _widget_con_actividades(1)
    try:
        m = w.modelo_valores; col = m.claves.index
        m.setData(m.index(0, col("precio_base")), "1234567,5", Qt.EditRole)
        assert m.index(0, col("valor_unitario")).data() == pesos_ui(w.valores[0]["valor_proyectado"])
        assert "." not in m.index(0, col("base_unitaria")).data()
        assert len(m.index(0, col("diferencia")).data().split(".")[-1]) <= 2
        w._mostrar_resultado(w.valores[0]); assert "1,234,568" in w.resultado.toPlainText()
        p = tmp_path / "formato.xlsx"; w._generar_excel(str(p)); libro = load_workbook(p)
        try:
            s = libro["Valores proyectados"]; C = {c.value.split("\n")[0]: c.column for c in s[2] if c.value}
            assert libro.sheetnames == ["Resumen", "Valores proyectados", "Proyecciones ICOCIV", "Trazabilidad", "Metodología y fórmulas"]
            assert s.cell(3, C["P"]).number_format == "#,##0" and s.cell(3, C["P"]).data_type == "n"
            assert s.cell(3, C["Valor proyectado unitario (COP)"]).data_type == "f" and s.cell(3, C["Diferencia %"]).number_format == "#,##0.00"
        finally:
            libro.close()
    finally:
        w.close()


def test_J_version_titulo_y_logo_oficial():
    from app_icociv.config.rutas import VERSION, TITULO_APLICACION, ruta_recurso
    from app_icociv.interfaz.ventana_principal import VentanaPrincipal
    from PIL import Image
    assert (ROOT / "VERSION").read_text(encoding="utf-8").strip() == "1.1.1" == VERSION
    assert TITULO_APLICACION.startswith("SAVIP v1.1.1") and TITULO_APLICACION.count("SAVIP") == 1
    v = VentanaPrincipal()
    try:
        assert v.windowTitle() == TITULO_APLICACION and not v.windowIcon().isNull()
        assert v.pantalla_inicio.version.text() == "SAVIP v1.1.1"
    finally:
        v.close()
    oficial = Image.open(ruta_recurso("app_icociv/interfaz/recursos/savip_logo.png")).convert("RGBA").resize((64, 64))
    ico = Image.open(ROOT / "packaging/savip.ico"); ico.size = (64, 64); ico = ico.convert("RGBA").resize((64, 64))
    diferencia = sum(abs(a - b) for pa, pb in zip(oficial.getdata(), ico.getdata()) for a, b in zip(pa, pb)) / (64 * 64 * 4)
    assert diferencia < 8, diferencia  # el .ico maestro es el logo oficial, no la moneda dorada anterior
