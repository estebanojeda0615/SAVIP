import json
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch
from PySide6.QtCore import QCoreApplication  # Fase 4I-G: la apertura de sesiones termina en el bucle de eventos
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog
from PySide6.QtCore import Qt
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.servicios.empalme_iccp_icociv import calcular_empalme_iccp_icociv
from app_icociv.persistencia.gestor_sesiones import cargar_sesion,guardar_sesion
ROOT=Path(__file__).resolve().parents[1]
APP=QApplication.instance() or QApplication([])
ICCP={'Total ICCP':{'2021_1':100.,'2021_6':105.,'2021_12':110.}}
from app_icociv.persistencia.historial_icociv import recuperar_resultado
from app_icociv.interfaz.widgets.proyecciones_icociv import _extraer_contexto
from tests.fixtures.sesiones import cargar as cargar_sesion_sintetica, ruta_temporal as ruta_sesion_sintetica  # Fase 1A (AS-007)
# Fase 1A (AS-007): fixture sintético trasladado de docs/fase_4a_excel_formulas_simples/ a tests/fixtures; el
# centinela de la Fase 0 ya no hace falta porque el fixture siempre está disponible.
PREVIO=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
FUTURO=_extraer_contexto(recuperar_resultado(PREVIO['proyecciones']['historial_analisis'][-1]['resultado']))
INDICES={'2021_12':100.,'2022_1':120.,'2022_6':150.,'2027_3':FUTURO['indice_proyectado']}

def registrar_casos(w):
    entradas=[]
    for inicio,fin,acero in [('2021_1','2021_6',False),('2022_1','2022_6',False),('2021_1','2022_1',False),('2021_1','2027_3',True),('2021_1','2022_1',False)]:
        entrada={'item':'Insumo repetido' if inicio=='2021_1' and not acero else 'Acero' if acero else 'Solo ICOCIV',
                 'unidad':'kg' if acero else 'm3','precio_base':1000.,'anticipo_amortizado':100.,'cantidad':2.,
                 'fecha_inicial':inicio,'fecha_final':fin,'tipo_serie_iccp':'total_iccp','serie_iccp':'Total ICCP','ruta_icociv':'Ruta > Sintética',
                 'calculo_acero':acero,'p0':1000.,'ix':5.5,'q':200.,'observacion_tecnica':'Observación real de prueba'}
        opcion={'indices':INDICES,'seleccion':{},'ruta_estructurada':[],
                'metadata_proyeccion':{'icociv_final_es_proyectado':fin=='2027_3','modelo_proyeccion':FUTURO['modelo'] if fin=='2027_3' else None,
                                       'horizonte_usado':3 if fin=='2027_3' else None,'estado_horizonte':FUTURO['estado'] if fin=='2027_3' else None}}
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)) as servicio:
            w._finalizar_calculo(entrada,opcion);assert servicio.call_count==1
        entradas.append((entrada,opcion))
    assert len(w.historial_empalme)==len(w.calculos)==5
    return entradas

def guardar(v,path):
    with patch.object(v,'_solicitar_usuario',return_value='QA4B'),patch('app_icociv.interfaz.ventana_principal.CARPETA_SESIONES',path.parent),patch.object(QFileDialog,'getSaveFileName',return_value=(str(path),'')):v.guardar_sesion_actual(como=True)
    assert path.exists()

def abrir(v,path):
    with patch.object(VentanaPrincipal, '_corte_dane_cargado', return_value='9999-12'),patch('app_icociv.interfaz.ventana_principal.CARPETA_SESIONES',path.parent),patch.object(QFileDialog,'getOpenFileName',return_value=(str(path),'')),patch.object(v,'_decidir_cierre',return_value='descartar'):
        v.abrir_sesion()
        while v._operacion_sesion: QCoreApplication.processEvents()

def test_cinco_ejecuciones_selectores_independientes_sin_recalculo():
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
        v.widget_proyecciones.restaurar_datos_sesion(deepcopy(fixture['proyecciones']));v._sincronizar_selector_resultados()
        ico=deepcopy(v.widget_proyecciones.datos_sesion());registrar_casos(w)
        assert 'revisión 2' in w.historial_empalme[-1]['nombre']
        original=deepcopy(w.datos_sesion());ids=[a['id'] for a in w.historial_empalme]
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Recalcular al consultar')) as motor,patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor ICOCIV')):
            for i in (0,1,2,3,0,4,2):
                w.selector_analisis_empalme.setCurrentIndex(w.selector_analisis_empalme.findData(ids[i]))
                assert w.analisis_empalme_activo_id==ids[i]==v.selector_resultados_empalme.currentData()
                assert w.item.text()==w.calculos[i]['item'] and w._periodo_inicial()==w.calculos[i]['fecha_inicial']
                assert w.resultado_analisis_empalme(w.analisis_empalme_activo())==original['calculos'][i]
                assert 'Resultado de actualización' in v.texto_resultados_empalme.toPlainText()
                assert v.widget_proyecciones.datos_sesion()==ico
            v.selector_resultados_empalme.setCurrentIndex(v.selector_resultados_empalme.findData(ids[1]))
            assert w.selector_analisis_empalme.currentData()==ids[1]
            assert motor.call_count==0
        assert w.calculos==original['calculos'] and w.historial_empalme==original['historial_empalme']
        assert v.selector_resultados_icociv.count()==4 and v.selector_resultados_empalme.count()==6
    finally:v.close()

def test_edicion_borrado_snapshot_y_cambios_no_guardados():
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        entradas=registrar_casos(w);ids=[a['id'] for a in w.historial_empalme];original=deepcopy(w.calculos[0])
        w.item.setText('Borrador sin guardar')
        with patch.object(QMessageBox,'question',return_value=QMessageBox.No):assert not w.activar_analisis_empalme(ids[0])
        assert w.item.text()=='Borrador sin guardar'
        with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):assert w.activar_analisis_empalme(ids[0])
        w.tabla.selectRow(0);w.editar_seleccionado()
        entrada,opcion=deepcopy(entradas[0]);entrada['precio_base']=1200.
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):w._finalizar_calculo(entrada,opcion)
        assert len(w.historial_empalme)==5 and w.calculos[0]['calculo_id']==original['calculo_id'] and w.calculos[0]['valor_actualizado']!=original['valor_actualizado']
        assert w.resultado_analisis_empalme(w.historial_empalme[0])==original
        w.tabla.selectRow(0);w.eliminar_seleccionado();assert len(w.calculos)==4 and len(w.historial_empalme)==5
        w.validar_datos_sesion(w.datos_sesion())
        assert w.activar_analisis_empalme(ids[0],confirmar=False)
        assert w.resultado_analisis_empalme(w.analisis_empalme_activo())==original
    finally:v.close()

def test_sesion_nueva_sin_fuente_y_rechazo_no_destructivo(tmp_path):
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        registrar_casos(w);w.contrato.setText('Proyecto común');path=tmp_path/'sesion.json';guardar(v,path)
        original=cargar_sesion(path);otro=VentanaPrincipal()
        try:
            with patch.object(otro,'ejecutar_proyeccion',side_effect=AssertionError('Motor al abrir')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Economía al abrir')):
                abrir(otro,path)
                assert otro.widget_empalme.datos_sesion()==original['empalme_iccp_icociv']
                assert not otro.widget_empalme.boton_calcular.isEnabled()
                assert otro.widget_empalme.contrato.text()=='Proyecto común'
                guardar(otro,tmp_path/'copia.json');assert cargar_sesion(tmp_path/'copia.json')['empalme_iccp_icociv']==original['empalme_iccp_icociv']
            before=deepcopy(otro.widget_empalme.datos_sesion())
            for mutacion in ('duplicado','huerfano','periodo','seleccion','activo','valores','json'):
                datos=deepcopy(original);emp=datos['empalme_iccp_icociv']
                if mutacion=='duplicado':emp['historial_empalme'].append(deepcopy(emp['historial_empalme'][0]))
                if mutacion=='huerfano':emp['calculos'].pop(0)
                if mutacion=='periodo':emp['calculos'][0]['fecha_final']='2021_99'
                if mutacion=='seleccion':emp['calculos'][0]['seleccion_icociv']={'idx_g':-1}
                if mutacion=='activo':emp['analisis_empalme_activo_id']='inexistente'
                if mutacion=='valores':emp['calculos'][0]['r_total']+=1
                bad=tmp_path/'bad.json';bad.write_text('{no json' if mutacion=='json' else json.dumps(datos),encoding='utf-8')
                with patch.object(QMessageBox,'warning') as warning:abrir(otro,bad);assert warning.call_count==1
                assert otro.widget_empalme.datos_sesion()==before
        finally:otro.close()
    finally:v.close()

def test_sesion_ambos_historiales_y_compatibilidad_anteriores(tmp_path):
    fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
    v=VentanaPrincipal()
    try:
        v.sesion_pendiente={'proyecciones_icociv':deepcopy(fixture['proyecciones'])};v._restaurar_sesion_pendiente()
        ico=deepcopy(v.widget_proyecciones.datos_sesion());registrar_casos(v.widget_empalme)
        v.selector_tipo_resultado.setCurrentIndex(1);path=tmp_path/'ambos.json';guardar(v,path);data=cargar_sesion(path)
        otro=VentanaPrincipal()
        try:
            with patch.object(otro,'ejecutar_proyeccion',side_effect=AssertionError('Motor')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Recalculo')):
                abrir(otro,path)
                assert otro.widget_proyecciones.datos_sesion()==data['proyecciones_icociv']
                assert otro.widget_empalme.datos_sesion()==data['empalme_iccp_icociv']
                assert otro.selector_tipo_resultado.currentIndex()==1
                guardar(otro,tmp_path/'ambos2.json');otra=cargar_sesion(tmp_path/'ambos2.json')
                for key in ('proyecciones_icociv','empalme_iccp_icociv','tipo_resultado_visible'):assert otra[key]==data[key]
                solo=tmp_path/'anterior.json';guardar_sesion(solo,{'proyecciones_icociv':ico})
                abrir(otro,solo);assert otro.widget_empalme.calculos==otro.widget_empalme.historial_empalme==[]
                assert otro.widget_proyecciones.datos_sesion()==ico
            heredado={'generales':{'contrato':'Anterior'},'calculos':deepcopy(fixture['empalme'])}
            otro.widget_empalme.restaurar_datos_sesion(heredado)
            assert otro.widget_empalme.calculos==heredado['calculos'] and otro.widget_empalme.historial_empalme==[]
            assert otro.widget_empalme.datos_sesion()['generales']['contrato']=='Anterior'
        finally:otro.close()
    finally:v.close()

def test_proyeccion_empalme_async_no_invade_icociv():
    from PySide6.QtTest import QTest
    v=VentanaPrincipal();source={'proyeccion':{'resultado_horizonte_solicitado':{}},'serie_df':None}
    try:
        fixture=cargar_sesion_sintetica('FASE_4A_FIXTURE_DATOS')
        v.sesion_pendiente={'proyecciones_icociv':deepcopy(fixture['proyecciones'])};v._restaurar_sesion_pendiente()
        ico=deepcopy(v.widget_proyecciones.datos_sesion());visible=deepcopy(v.resultado_ui_actual)
        fechas=(v.spin_anio.value(),v.spin_mes.value());calls=[]
        with patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',return_value=source) as motor:
            v._ejecutar_proyeccion_para_empalme_async({},2027,3,lambda r,e:calls.append((r,e)))
            for _ in range(50):
                QTest.qWait(50)
                if calls:break
            assert calls==[(source,None)] and motor.call_count==1
        assert v.widget_proyecciones.datos_sesion()==ico and fechas==(v.spin_anio.value(),v.spin_mes.value())
        from app_icociv.persistencia.historial_icociv import serializar_resultado
        assert serializar_resultado(v.resultado_ui_actual)==serializar_resultado(visible)
    finally:v.close()

def test_fuente_disponible_y_historial_archivado_exportable(tmp_path):
    from tests.test_proyecciones_fase_3 import anexo_sintetico
    from openpyxl import load_workbook
    path=tmp_path/'anexo_sintetico.xlsx';anexo_sintetico(path)
    v=VentanaPrincipal();w=v.widget_empalme
    try:
        v.controlador.cargar_archivo(path);w.actualizar_icociv(v.controlador.tablas,v.controlador.periodos,path)
        registrar_casos(w);w.contrato.setText('Proyecto actual distinto')
        with patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.cargar_archivo',side_effect=AssertionError('Lectura al consultar')) as lector,patch('app_icociv.interfaz.controladores.controlador_principal.ControladorPrincipal.ejecutar_analisis',side_effect=AssertionError('Motor al consultar')) as motor:
            for a in w.historial_empalme:
                assert w.activar_analisis_empalme(a['id'],confirmar=False)
                assert w.boton_calcular.isEnabled() and w.contrato.text()=='Proyecto actual distinto'
            assert lector.call_count==motor.call_count==0
        before=deepcopy(w.datos_sesion())
        with patch.object(QMessageBox,'exec',lambda box:0),patch.object(QMessageBox,'clickedButton',lambda box:next(b for b in box.buttons() if b.text()=='Cancelar')):w.eliminar_todo()
        assert w.datos_sesion()==before
        with patch.object(QMessageBox,'exec',lambda box:0),patch.object(QMessageBox,'clickedButton',lambda box:next(b for b in box.buttons() if b.text()=='Eliminar')):w.eliminar_todo()
        assert not w.calculos and len(w.historial_empalme)==5
        w.validar_datos_sesion(w.datos_sesion());assert w.activar_analisis_empalme(w.historial_empalme[0]['id'],confirmar=False)
        w.barra.actualizar();assert w.barra.boton_exportar.isEnabled()
        out=tmp_path/'archivados.xlsx';w._generar_excel_empalme(str(out));b=load_workbook(out)
        assert b['Historial de Empalme'].max_row>=7 and all(b['Historial de Empalme'].cell(r,12).value for r in range(3,8));b.close()  # 4E añade un bloque separado; las cinco filas originales no cambian.
        original_largo=deepcopy(w.historial_empalme[0]['resultado_archivado'])
        original_largo['observacion_tecnica']='Observación extensa de prueba. '*2000
        w.historial_empalme[0]['resultado_archivado']=original_largo
        largo=tmp_path/'archivado_largo.xlsx';w._generar_excel_empalme(str(largo));b=load_workbook(largo)
        partes=[(r[12].value,r[11].value) for r in b['Historial de Empalme'].iter_rows(min_row=3) if r[0].value==w.historial_empalme[0]['id']]
        assert len(partes)>1 and json.loads(''.join(x[1] for x in sorted(partes)))==original_largo
        b.close()
    finally:v.close()
