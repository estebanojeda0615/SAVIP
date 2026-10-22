"""4E: declaración explícita, deriva presupuesto y conserva la economía original."""
import json,math
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from dataclasses import asdict
from PySide6.QtCore import Qt,QTimer,QUrl,QPoint,QPointF
from PySide6.QtGui import QFontDatabase,QFont,QWheelEvent
from PySide6.QtWidgets import QDialog,QDialogButtonBox,QComboBox,QApplication,QScrollArea,QMessageBox
from PySide6.QtTest import QTest,QSignalSpy
from openpyxl import load_workbook
from app_icociv.utilidades.importes_presupuestales import MODALIDADES,desglose_importe,validar_modalidad,texto_desglose
from app_icociv.servicios.actualizacion_icociv import calcular_actualizacion_icociv
from app_icociv.servicios.empalme_iccp_icociv import calcular_empalme_iccp_icociv
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.widgets.empalme_iccp_icociv import WidgetEmpalmeICCPICOCIV
from app_icociv.reportes.sesion_modulos import construir_proyecciones,construir_empalme
from tests.test_slider_comparacion_historica_fase_4d import APP,reales,sintetico,cargar
from tests.test_empalme_fase_4c import registrar_seis
from tests.test_historial_empalme_fase_4b import ICCP,guardar,abrir
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
for fuente in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):assert QFontDatabase.addApplicationFont('C:/Windows/Fonts/'+fuente)>=0
APP.setFont(QFont('Segoe UI',10))
NUMEROS=('cantidad','precio_base','anticipo_amortizado','base_ajustable','i0_icociv','i_icociv','factor_icociv','r_total','valor_proyectado')

def test_01_servicio_real_990_y_matriz_cantidades_tres_modos():
    assert tuple(MODALIDADES)==('heredado','unitario','total')
    for q in (2.,2.5,1.,None,0.,-2.,1e308):
        r=calcular_actualizacion_icociv({'precio_base':1000.,'anticipo_amortizado':100.,'i0_icociv':100.,'i_icociv':110.,'cantidad':q});original=deepcopy(r);valor=r['valor_proyectado']
        assert math.isclose(valor,990.,abs_tol=1e-10)
        for modo in MODALIDADES:
            copia={**r,'modalidad_importe':modo};antes=deepcopy(copia);d=desglose_importe(copia,valor);assert copia==antes and r==original
            if modo=='heredado':assert d['unitario'] is d['total'] is None
            elif modo=='unitario':
                assert d['unitario']==valor
                if q in (2.,2.5,1.):assert d['total']==valor*q
                else:assert d['total'] is None
            else:
                assert d['total']==valor
                if q in (2.,2.5,1.,1e308):assert d['unitario']==valor/q
                else:assert d['unitario'] is None
        if q==2.:
            assert math.isclose(desglose_importe({**r,'modalidad_importe':'unitario'},valor)['total'],1980.,abs_tol=1e-9)
            assert math.isclose(desglose_importe({**r,'modalidad_importe':'total'},valor)['unitario'],495.,abs_tol=1e-9)
    r=calcular_actualizacion_icociv({'precio_base':1000.,'anticipo_amortizado':100.,'i0_icociv':100.,'i_icociv':-10.,'cantidad':2.5})
    assert r['valor_proyectado']<0 and desglose_importe({**r,'modalidad_importe':'unitario'},r['valor_proyectado'])['total']<0

def test_02_extremos_y_metadata_invalida_sin_ceros_inventados():
    for malo in ('parcial',None,False,1,float('nan'),float('inf'),[],{}):
        try:validar_modalidad({'modalidad_importe':malo})
        except ValueError:pass
        else:raise AssertionError('Enum inválida aceptada')
    for q in (None,0,-1,float('inf'),float('nan'),'2',True):
        assert desglose_importe({'cantidad':q,'modalidad_importe':'unitario'},10)['total'] is None
        assert desglose_importe({'cantidad':q,'modalidad_importe':'total'},10)['unitario'] is None
    assert desglose_importe({'cantidad':1e308,'modalidad_importe':'unitario'},1e308)['total'] is None
    assert desglose_importe({'cantidad':1e-308,'modalidad_importe':'total'},1e308)['unitario'] is None
    assert desglose_importe({'cantidad':1e-308,'modalidad_importe':'unitario'},1e-308)['total'] is None
    assert desglose_importe({'cantidad':2.,'modalidad_importe':'unitario'},0.)['total']==0.
    assert '1e-308' in texto_desglose({'cantidad':1.,'modalidad_importe':'unitario'},1e-308,'Original')

def test_03_formulario_icociv_metadata_y_precision_original():
    w=sintetico()
    try:
        assert w.modalidad_importe.currentData()=='heredado' and 'modalidad_importe' not in w.valores[0]
        original=deepcopy(w.valores[0])
        for modo in ('unitario','total'):
            w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modo));antes=deepcopy(w.valores);w.boton_calcular.click();r=w.valores[-1]
            assert all(r[k]==original[k] for k in NUMEROS) and w.valores[:-1]==antes
            assert validar_modalidad(r)==modo and 'unitario' in w.resultado.toPlainText() and 'total' in w.resultado.toPlainText()
        assert w.modelo_valores.columnCount()==17
    finally:w.close()

def test_04_edicion_icociv_dialogo_real_guardar_cancelar_sin_servicio():
    w=sintetico()
    try:
        r=w.valores[0];original=deepcopy(r);historial=deepcopy(w.historial);referencias=deepcopy(w._periodo_por_analisis)
        def respuesta(modo,guardar_):
            def actuar():
                dialogo=QApplication.activeModalWidget();assert isinstance(dialogo,QDialog)
                c=dialogo.findChild(QComboBox);c.setCurrentIndex(c.findData(modo))
                b=dialogo.findChild(QDialogButtonBox).button(QDialogButtonBox.Save if guardar_ else QDialogButtonBox.Cancel);QTest.mouseClick(b,Qt.LeftButton)
            QTimer.singleShot(0,actuar);w._editar_interpretacion(QUrl('interpretacion'))
        with patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')):
            respuesta('unitario',False);assert r==original
            respuesta('unitario',True);assert r=={**original,'modalidad_importe':'unitario'}
            respuesta('total',False);assert r['modalidad_importe']=='unitario'
            respuesta('total',True);assert r=={**original,'modalidad_importe':'total'}
            respuesta('heredado',True);assert r=={**original,'modalidad_importe':'heredado'}
        assert w.historial==historial and w._periodo_por_analisis==referencias
    finally:w.close()

def test_05_slider_individual_global_rueda_y_resumen_por_id():
    v=VentanaPrincipal();w=cargar(v)
    try:
        for r,modo in zip(w.valores,('unitario','total','heredado')):r['modalidad_importe']=modo
        v._activar_analisis(w.historial[0]['id']);original=deepcopy(w.valores);snapshots=deepcopy([a['resultado'] for a in w.historial])
        w.tabla_valores.selectRow(0);assert original[0]['calculo_id'] in w.resumen_presupuestal.text()
        with patch.object(v.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor')),patch.object(v.controlador,'cargar_archivo',side_effect=AssertionError('DANE')),patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')):
            assert w.seleccionar_referencia_historica('2024-02');e=w.escenario_registro(w.valores[0])[0];assert desglose_importe(w.valores[0],e['valor_proyectado'])['total']==e['valor_proyectado']*w.valores[0]['cantidad']
            otros=[w.escenario_registro(r)[0] for r in w.valores][1:];assert otros==original[1:]
            v._cambiar_seccion(1);v.resize(1366,768);v.show();APP.processEvents();s=w.slider_comparacion;scroll=w.findChild(QScrollArea);bar=scroll.verticalScrollBar();bar.setValue(bar.maximum()//2)
            previo=deepcopy(w.datos_sesion());pos=s.value();valor=bar.value();spy=QSignalSpy(w.comparacion_cambiada)
            APP.sendEvent(s,QWheelEvent(QPointF(s.rect().center()),QPointF(s.mapToGlobal(s.rect().center())),QPoint(),QPoint(0,-120),Qt.NoButton,Qt.NoModifier,Qt.NoScrollPhase,False));APP.processEvents()
            assert s.value()==pos and w.datos_sesion()==previo and spy.count()==0 and bar.value()>valor
            w.aplicar_global.setChecked(True);assert w.seleccionar_referencia_historica('2024-05')
            for r in w.valores:assert validar_modalidad(w.escenario_registro(r)[0])==validar_modalidad(r)
            assert w.seleccionar_referencia_historica('2026-12') and [w.escenario_registro(r)[0] for r in w.valores]==original
        assert w.valores==original and [a['resultado'] for a in w.historial]==snapshots
    finally:v.close()

def test_06_empalme_real_acero_edicion_cancelada_y_archivo():
    w=WidgetEmpalmeICCPICOCIV()
    try:
        entradas=registrar_seis(w);original=deepcopy(w.calculos)
        for r in w.calculos:
            for modo in MODALIDADES:
                d=desglose_importe({**r,'modalidad_importe':modo},r['valor_actualizado'])
                if modo=='unitario':assert d['unitario']==r['valor_actualizado'] and d['total']==r['valor_actualizado']*r['cantidad']
                if modo=='total':assert d['total']==r['valor_actualizado'] and d['unitario']==r['valor_actualizado']/r['cantidad']
        acero=w.calculos[3];assert acero['q']!=acero['cantidad'] and acero['anticipo_amortizado']==0
        assert desglose_importe({**acero,'p0':None,'modalidad_importe':'unitario'},acero['valor_actualizado'])['unitario'] is None
        w.tabla_operativa.selectRow(0);w.editar_seleccionado();w.modalidad_importe.setCurrentIndex(1);w.precio_base.setValue(1234);w.boton_limpiar.click()
        assert w.indice_edicion is None and w.modalidad_importe.currentData()=='heredado' and w.precio_base.value()==original[0]['precio_base'] and w.calculos==original
        for modo in ('unitario','total'):
            w.tabla_operativa.selectRow(0);w.editar_seleccionado();w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData(modo));e,o=deepcopy(entradas[0]);e['modalidad_importe']=modo
            with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):w._finalizar_calculo(e,o)
            assert all(w.calculos[0].get(k)==original[0].get(k) for k in original[0] if type(original[0][k]) in (int,float))
        assert w.historial_empalme[0]['resultado_archivado']==original[0]
        w.tabla_operativa.selectRow(0);w.eliminar_seleccionado();cantidad=len(w.calculos)
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Consulta')):
            assert w.activar_analisis_empalme(w.historial_empalme[0]['id'],confirmar=False) and len(w.calculos)==cantidad
            assert w.modalidad_importe.currentData()=='heredado'
        assert w.modelo_operativo.columnCount()==18
        e,o=deepcopy(entradas[1]);e['modalidad_importe']='unitario'
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=lambda e,i:calcular_empalme_iccp_icociv(e,i,ICCP)):w._finalizar_calculo(e,o)
        original_unitario=deepcopy(w.calculos[-1]);w.tabla_operativa.selectRow(len(w.calculos)-1);w.editar_seleccionado();w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('total'))
        with patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Servicio por modalidad')),patch.object(w,'_opcion_icociv_o_vacia',side_effect=AssertionError('Fuente por modalidad')):
            w.boton_calcular.click()
        assert w.calculos[-1]=={**original_unitario,'modalidad_importe':'total'}
        assert w.historial_empalme[-1]['resultado_archivado']==original_unitario
    finally:w.close()

def test_07_sesion_mixta_sin_fuente_y_rechazo_no_destructivo(tmp_path):
    v=VentanaPrincipal();w=cargar(v);otro=None
    try:
        registrar_seis(v.widget_empalme)
        for r,m in zip(w.valores,('unitario','total','heredado')):r['modalidad_importe']=m
        for i,r in enumerate(v.widget_empalme.calculos):r['modalidad_importe']=tuple(MODALIDADES)[i%3]
        ausente=tmp_path/'DANE_AUSENTE.xlsx'
        for a in w.historial:a['archivo_excel']=str(ausente)
        v.controlador.ruta_archivo=ausente;v._activar_analisis(w.historial[0]['id']);w.aplicar_global.setChecked(True);w.seleccionar_referencia_historica('2024-05')
        antes=deepcopy(w.datos_sesion());emp=deepcopy(v.widget_empalme.datos_sesion());p=tmp_path/'mixta.json';guardar(v,p);otro=VentanaPrincipal()
        with patch.object(otro.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor')),patch.object(otro.controlador,'cargar_archivo',side_effect=AssertionError('DANE')),patch('app_icociv.interfaz.widgets.empalme_iccp_icociv.calcular_empalme_iccp_icociv',side_effect=AssertionError('Servicio')):
            abrir(otro,p);n=otro.widget_proyecciones;assert n.datos_sesion()==cargar_sesion(p)['proyecciones_icociv'] and otro.widget_empalme.datos_sesion()==emp
            assert n.valores==antes['valores'] and n._periodo_por_analisis==w._periodo_por_analisis
        for module,datos,clave in ((n,n.datos_sesion(),'valores'),(otro.widget_empalme,emp,'calculos')):
            for malo in ('parcial',None,False,float('inf'),[]):
                copia=deepcopy(datos);copia[clave][0]['modalidad_importe']=malo;anterior=deepcopy(module.datos_sesion())
                try:module.restaurar_datos_sesion(copia)
                except ValueError:pass
                else:raise AssertionError('Metadata inválida aceptada')
                assert module.datos_sesion()==anterior
        viejo=WidgetProyeccionesICOCIV()
        try:viejo.restaurar_datos_sesion(reales());assert all('modalidad_importe' not in r for r in viejo.datos_sesion()['valores'])
        finally:viejo.close()
    finally:
        if otro:otro.close()
        v.close()

def bloque(ws):
    titulo=next(row[0].row for row in ws if isinstance(row[0].value,str) and row[0].value.startswith('Desglose presupuestal complementario'))
    return titulo+2

def test_08_excel_estructuras_formulas_vivas_ids_y_no_inyeccion(tmp_path):
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales());emp=WidgetEmpalmeICCPICOCIV()
    try:
        for r,m in zip(w.valores,('unitario','total','heredado')):r['modalidad_importe']=m
        w.valores[0]['item']='=1+1';p=tmp_path/'ico.xlsx';w._generar_excel(str(p));b=load_workbook(p)
        try:
            # Fase 4G: objetivos distintos (2027-01/02/03), I0 único → 18 columnas; M es la fecha de I proyectado.
            s=b['Valores proyectados'];t=b['Trazabilidad'];inicio=bloque(t);assert s.max_column==18 and len(b.sheetnames)==5
            assert s.cell(2,13).value=='Fecha de I proyectado (AAAA-MM)' and s.cell(2,11).value=='I0 ICOCIV\n2026-12'
            for f,r in enumerate(w.valores,3):
                assert s.cell(f,14).value==f'=L{f}/K{f}' and s.cell(f,17).value==f'=(N{f}-1)*100'
                if r['modalidad_importe']=='heredado':assert all(s.cell(f,c).value is None for c in (9,10,15,16))
                else:
                    base,valor=(9,15) if r['modalidad_importe']=='unitario' else (10,16)
                    letra='I' if base==9 else 'J'
                    assert s.cell(f,base).value==f'=G{f}-H{f}' and s.cell(f,valor).value==f'={letra}{f}+{letra}{f}*(N{f}-1)'
                assert t.cell(inicio+f-3,1).value==r['calculo_id']
            assert t.cell(inicio,11).value=="='Valores proyectados'!P3"
            assert t.cell(inicio+1,10).value=="='Valores proyectados'!O4"
            assert all(t.cell(inicio+2,c).value is None for c in (8,9,10,11))
            assert s['D3'].data_type=='s' and s['D3'].value=='=1+1' and not b._external_links
            assert not any(c.comment for h in b for row in h for c in row)
        finally:b.close()
        registrar_seis(emp)
        for i,r in enumerate(emp.calculos):r['modalidad_importe']=tuple(MODALIDADES)[i%3]
        p=tmp_path/'emp.xlsx';emp._generar_excel_empalme(str(p));b=load_workbook(p)
        try:
            assert b['Resumen operativo Empalme'].max_column==18 and b.sheetnames==['Empalme ICCP-ICOCIV','Resultados registrados','Cálculos editables','Historial de Empalme','Resumen operativo Empalme']
            t=b['Resultados registrados'];inicio=bloque(t)
            assert t.cell(inicio+1,11).value=="='Resumen operativo Empalme'!O4"
            assert t.cell(inicio+2,10).value=="='Resumen operativo Empalme'!N5"
            assert not any(c.comment for h in b for row in h for c in row)
        finally:b.close()
    finally:w.close();emp.close()

def test_09_documentos_sin_nuevas_columnas_ni_repetir_estadistica():
    w=sintetico();emp=WidgetEmpalmeICCPICOCIV()
    try:
        w.valores[0]['modalidad_importe']='unitario';w.seleccionar_referencia_historica('2024-02');d=asdict(construir_proyecciones(w));s=next(s for s in d['secciones'] if s['clave']=='valores')
        assert [len(b['encabezados']) for b in s['bloques'] if 'encabezados' in b]==[7,7,3]
        texto=json.dumps(s,ensure_ascii=False);assert 'Valor ajustado unitario: 945' in texto and 'Valor ajustado total: 1,890' in texto and 'Escenario vigente del slider' in texto
        registrar_seis(emp);emp.calculos[3]['modalidad_importe']='unitario';d=asdict(construir_empalme(emp.calculos,{},historial=emp.historial_empalme,operativos=emp.calculos))
        s=next(s for s in d['secciones'] if s['clave']=='resumen_operativo_empalme');assert max(len(b['encabezados']) for b in s['bloques'] if 'encabezados' in b)==7
        texto=json.dumps(s,ensure_ascii=False);assert emp.calculos[3]['calculo_id'] in texto and 'Valores unitarios' in texto and 'Valor ajustado total:' in texto
    finally:w.close();emp.close()

def test_10_ui_temas_resoluciones_seleccion_copia_y_vistas():
    v=VentanaPrincipal();w=cargar(v);emp=v.widget_empalme
    try:
        registrar_seis(emp)
        w.valores[0]['modalidad_importe']='unitario';emp.calculos[0]['modalidad_importe']='total'
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for ancho,alto in ((1100,700),(1366,768),(1920,1080)):
                v.resize(ancho,alto);v.show();v._cambiar_seccion(1);APP.processEvents();APP.processEvents();w.tabla_valores.selectRow(0)
                assert w.findChild(QScrollArea).horizontalScrollBar().maximum()==0 and w.valores[0]['calculo_id'] in w.resumen_presupuestal.text()
                QTest.keyClick(w.tabla_valores,Qt.Key_C,Qt.ControlModifier);QTest.keyRelease(w.tabla_valores,Qt.Key_Control);assert len(APP.clipboard().text().split('\t'))==18  # 4G: + fecha de I proyectado
                w.tabla_valores.clearSelection();assert 'Seleccione una sola fila' in w.resumen_presupuestal.text()
                v._cambiar_seccion(2);APP.processEvents();emp.tabla_operativa.selectRow(0)
                assert emp.findChild(QScrollArea).horizontalScrollBar().maximum()==0 and emp.calculos[0]['calculo_id'] in emp.contexto_registro.text()
                for i in (1,0):emp.vistas_tablas.setCurrentIndex(i);assert 'Valores totales' in emp.contexto_registro.text()
    finally:v.close()
