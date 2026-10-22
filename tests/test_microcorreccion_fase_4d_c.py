"""Contratos de presentación y eventos Qt reales; no se vuelve a ejecutar el motor."""
import json
from copy import deepcopy
from dataclasses import asdict
from unittest.mock import patch
from PySide6.QtCore import Qt,QPoint,QPointF
from PySide6.QtGui import QWheelEvent,QFontDatabase,QFont
from PySide6.QtTest import QTest,QSignalSpy
from PySide6.QtWidgets import QScrollArea,QAbstractSlider,QStyle,QStyleOptionSlider,QBoxLayout
from openpyxl import load_workbook
from tests.test_slider_comparacion_historica_fase_4d import APP,cargar,reales,sintetico
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV,COLUMNAS_VALORES,COLUMNAS_PROYECCIONES
from app_icociv.reportes.sesion_modulos import construir_proyecciones

TITULOS=['N.º','Proyecto','Componente del proyecto','Ítem / insumo','Unidad','Cantidad','P','A','Base ajustable unitaria (COP)','Base ajustable total (COP)','I0 ICOCIV','I ICOCIV proyectado','Factor','Valor proyectado unitario (COP)','Valor proyectado total (COP)','Diferencia %','Observación técnica']
for fuente in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):
    assert QFontDatabase.addApplicationFont('C:/Windows/Fonts/'+fuente)>=0
APP.setFont(QFont('Segoe UI',10))

def rueda(slider,delta,punto):
    evento=QWheelEvent(QPointF(punto),QPointF(slider.mapToGlobal(punto)),QPoint(),QPoint(0,delta),Qt.NoButton,Qt.NoModifier,Qt.NoScrollPhase,False)
    APP.sendEvent(slider,evento);APP.processEvents()

def test_columnas_0_1_3_y_multiples_sin_cambio_economico():
    assert [t for _,t in COLUMNAS_VALORES]==TITULOS and len(COLUMNAS_PROYECCIONES)==21
    for cantidad in (0,1,3,9):
        w=WidgetProyeccionesICOCIV()
        try:
            d=reales();d['valores']=(d['valores']*3)[:cantidad]
            for i,r in enumerate(d['valores']):r['calculo_id']=f'contrato_{i}'
            for a in d['historial_analisis']:a['registros_ids']=[r['calculo_id'] for r in d['valores'] if r['analisis_id']==a['id']]
            w.restaurar_datos_sesion(d);antes=deepcopy(w.datos_sesion())
            # Fase 4G: los objetivos de estos análisis difieren (2027-01/02/03) y todos los I0 son 2026-12:
            # con 2+ filas se añade solo la columna de fecha de I proyectado (18); con 0 o 1 fila, 17.
            assert w.modelo_valores.columnCount()==(18 if cantidad>1 else 17) and w.modelo_valores.rowCount()==cantidad
            assert [t for t in w.modelo_valores._dataframe.columns if t!='Fecha de I proyectado (AAAA-MM)']==TITULOS
            col=w.modelo_valores.claves.index
            for i,r in enumerate(w._valores_visibles):
                e=w.escenario_registro(r)[0]
                for c,k in ((col('i0_icociv'),'i0_icociv'),(col('i_icociv'),'i_icociv'),(col('diferencia'),'diferencia_porcentual')):assert w.modelo_valores.index(i,c).data(Qt.UserRole)==e[k]
            assert all(w.modelo_valores.index(i,col(k)).data(Qt.UserRole) is None for i in range(cantidad) for k in ('base_unitaria','base_total','valor_unitario','valor_total'))
            assert w.datos_sesion()==antes
        finally:w.close()

def test_fechas_mixtas_filtro_tooltip_y_tsv():
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales())
    try:
        w.establecer_analisis_activo(w.historial[0]['id']);w.seleccionar_referencia_historica('2024-02');m=w.modelo_valores
        # Fase 4G: I0 (2024-02 y 2026-12) e I (tres objetivos) difieren → 19 columnas, títulos técnicos sin fecha.
        assert m.columnCount()==19 and m.headerData(10,Qt.Horizontal)=='I0 ICOCIV' and m.headerData(12,Qt.Horizontal)=='I ICOCIV proyectado'
        assert [m.headerData(c,Qt.Horizontal) for c in (11,13)]==['Fecha I0 (AAAA-MM)','Fecha de I proyectado (AAAA-MM)']
        for i,r in enumerate(w._valores_visibles):
            assert [m.index(i,c).data() for c in (11,13)]==list(w.fechas_registro(r))
            for c in (10,12):
                tip=m.index(i,c).data(Qt.ToolTipRole)
                assert all(p in tip for p in w.fechas_registro(r)) and 'Referencia efectiva:' in tip and 'Objetivo:' in tip
        w.solo_analisis_activo.setChecked(True)
        assert m.rowCount()==1 and m.columnCount()==17 and m.headerData(10,Qt.Horizontal)=='I0 ICOCIV\n2024-02'
        assert m.headerData(11,Qt.Horizontal)=='I ICOCIV proyectado\n'+w.fechas_registro(w._valores_visibles[0])[1]
        w.tabla_valores.selectRow(0);QTest.keyClick(w.tabla_valores,Qt.Key_C,Qt.ControlModifier)
        celdas=APP.clipboard().text().split('\t');assert len(celdas)==17
        assert float(celdas[10])==m.index(0,10).data(Qt.UserRole)
        w.solo_analisis_activo.setChecked(False);assert m.rowCount()==3 and m.columnCount()==19 and m.headerData(10,Qt.Horizontal)=='I0 ICOCIV'
    finally:w.close()

def test_rueda_48_eventos_con_scroll_sin_senales_ni_recalculo():
    v=VentanaPrincipal();w=cargar(v);v._cambiar_seccion(1);v.resize(1366,768);v.show();APP.processEvents()
    try:
        v._activar_analisis(w.historial[0]['id']);s=w.slider_comparacion;scroll=w.findChild(QScrollArea);barra=scroll.verticalScrollBar()
        for global_,modificado in ((False,False),(False,True),(True,False),(True,True)):
            w.aplicar_global.setChecked(global_);w.seleccionar_referencia_historica('2024-02' if modificado else '2026-12')
            scroll.ensureWidgetVisible(s);APP.processEvents();spy=QSignalSpy(w.comparacion_cambiada)
            antes=deepcopy(w.datos_sesion());economia=[w.escenario_registro(r) for r in w.valores];resultado=v.lbl_comparacion_resultados.text();pos=s.value()
            with patch.object(v.controlador,'ejecutar_analisis',side_effect=AssertionError('Motor')),patch.object(v.controlador,'cargar_archivo',side_effect=AssertionError('DANE')),patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio')),patch.object(v,'guardar_sesion_actual',side_effect=AssertionError('Escritura')):
                for foco in (False,True):
                    s.setFocus() if foco else s.clearFocus()
                    for x in (1,s.width()//2,s.width()-2):
                        for delta in (-120,120):
                            barra.setValue(barra.maximum()//2);previo=barra.value();rueda(s,delta,QPoint(x,s.height()//2))
                            assert barra.value()>previo if delta<0 else barra.value()<previo
                            assert s.value()==pos and w.datos_sesion()==antes and [w.escenario_registro(r) for r in w.valores]==economia
                            assert v.lbl_comparacion_resultados.text()==resultado and spy.count()==0
    finally:v.close()

def test_interaccion_explicita_teclado_clic_arrastre_y_manual():
    w=sintetico();w.show();APP.processEvents()
    try:
        s=w.slider_comparacion;s.setValue(0);spy=QSignalSpy(w.comparacion_cambiada)
        QTest.keyClick(s,Qt.Key_Right);assert spy.count()==1 and w.selector_referencia.currentData()=='2024-02'
        w.seleccionar_referencia_historica('2026-12');antes=spy.count()
        opcion=QStyleOptionSlider();s.initStyleOption(opcion);mango=s.style().subControlRect(QStyle.CC_Slider,opcion,QStyle.SC_SliderHandle,s)
        QTest.mouseClick(s,Qt.LeftButton,pos=QPoint(2,s.height()//2));assert spy.count()==antes+1
        opcion=QStyleOptionSlider();s.initStyleOption(opcion);mango=s.style().subControlRect(QStyle.CC_Slider,opcion,QStyle.SC_SliderHandle,s)
        QTest.mousePress(s,Qt.LeftButton,pos=mango.center());QTest.mouseMove(s,QPoint(2,s.height()//2));QTest.mouseRelease(s,Qt.LeftButton,pos=QPoint(2,s.height()//2))
        assert w.selector_referencia.currentData()=='2020-01'
        antes=spy.count();w.selector_referencia.setEditText('2024-02');QTest.keyClick(w.selector_referencia.lineEdit(),Qt.Key_Return)
        assert spy.count()==antes+1 and w.selector_referencia.currentData()=='2024-02'
    finally:w.close()

def test_selector_compacto_nombres_largos_ids_y_geometria():
    for f in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):assert QFontDatabase.addApplicationFont('C:/Windows/Fonts/'+f)>=0
    APP.setFont(QFont('Segoe UI',10))
    v=VentanaPrincipal();w=cargar(v)
    try:
        nombre='Análisis con nombre completo y trazabilidad '+('muy largo '*8);w.historial[0]['nombre']=nombre;w._refrescar_historial();v._cambiar_seccion(1)
        for tema in ('claro','oscuro'):
            v._aplicar_tema(tema,persistir=False)
            for ancho,alto in ((1100,700),(1366,768),(1920,1080)):
                v.resize(ancho,alto);v.show();APP.processEvents()
                c=w.selector_analisis_valores;c.setCurrentIndex(c.findData(w.historial[0]['id']));APP.processEvents();APP.processEvents()
                # Entrega 1.1: «Proyección guardada» vive en el historial superior a ancho completo.
                assert c.width()>=300 and c.parentWidget() is w.selector_ambito.parentWidget()
                # Fase 4G-B: el desplegable de proyección usa etiqueta corta (objetivo, H, revisión); el nombre largo sigue en el superior.
                assert c.currentText()==c.toolTip() and w.selector_analisis.currentText()==nombre and c.currentData()==w.selector_analisis.currentData()==v.selector_resultados_icociv.currentData()
                assert w.findChild(QScrollArea).horizontalScrollBar().maximum()==0,(tema,ancho,w.width(),w.findChild(QScrollArea).widget().width())
                assert w.fila_filtro.direction()==QBoxLayout.LeftToRight
                c.showPopup();APP.processEvents();assert c.itemText(c.currentIndex())==c.currentText() and c.view().minimumWidth()<=c.screen().availableGeometry().width()
                assert c.view().horizontalScrollBarPolicy()==Qt.ScrollBarAsNeeded
                assert c.view().horizontalScrollBar().maximum()>0 or c.view().fontMetrics().horizontalAdvance(c.currentText())<=c.view().viewport().width(),(c.view().fontMetrics().horizontalAdvance(c.currentText()),c.view().viewport().width(),c.view().horizontalScrollBar().maximum())
                c.view().horizontalScrollBar().setValue(c.view().horizontalScrollBar().maximum());assert c.itemText(c.currentIndex())==c.currentText();c.hidePopup()
        for i in (2,3,1):
            w.selector_analisis_valores.setCurrentIndex(i);assert w.selector_analisis.currentData()==v.selector_resultados_icociv.currentData()==w.selector_analisis_valores.currentData()
    finally:v.close()

def test_excel_15_reales_formulas_y_fechas_en_traza(tmp_path):
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales())
    try:
        for r in w.valores:r['modalidad_importe']='unitario'
        w.establecer_analisis_activo(w.historial[0]['id']);w.seleccionar_referencia_historica('2024-02');p=tmp_path/'quince.xlsx';w._generar_excel(str(p));b=load_workbook(p)
        try:
            s=b['Valores proyectados'];t=b['Trazabilidad'];assert b.sheetnames==['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
            # Fase 4G: I0 e I heterogéneos en el Excel → 19 columnas; L y N son las fechas auxiliares.
            assert s.max_column==19 and s.calculate_dimension().endswith(':S5') and s.auto_filter.ref=='A2:S5'
            assert not any(c>19 for _,c in s._cells) and not any(k in s.column_dimensions for k in ('T','U'))
            assert [s.cell(2,c).value for c in (11,12,13,14)]==['I0 ICOCIV','Fecha I0 (AAAA-MM)','I ICOCIV proyectado','Fecha de I proyectado (AAAA-MM)']
            for i,r in enumerate(w.valores,3):
                assert [s.cell(i,c).value for c in (9,15,16,18)]==[f'=G{i}-H{i}',f'=M{i}/K{i}',f'=I{i}+I{i}*(O{i}-1)',f'=(O{i}-1)*100']
                assert all(s.cell(i,c).data_type=='n' for c in (6,7,8,11,13))
                assert [s.cell(i,c).value for c in (12,14)]==list(w.fechas_registro(r))
                assert (t.cell(i,9).value,t.cell(i,11).value)==w.fechas_registro(r) and t.cell(i,2).value==r['calculo_id']
            assert not any(c.comment for h in b for row in h for c in row)
        finally:b.close()
    finally:w.close()

def test_documentos_fechas_por_id_sin_columnas_adicionales():
    w=sintetico()
    try:
        w.seleccionar_referencia_historica('2024-02');informe=asdict(construir_proyecciones(w));valores=next(s for s in informe['secciones'] if s['clave']=='valores')
        tablas=[b for b in valores['bloques'] if 'encabezados' in b];assert [len(t['encabezados']) for t in tablas]==[7,7,3]
        assert not any('Referencia histórica' in t['encabezados'] or 'Objetivo proyectado' in t['encabezados'] for t in tablas)
        texto=json.dumps(valores,ensure_ascii=False);assert 'Referencia efectiva: 2024-02' in texto and 'Objetivo:' in texto and w.valores[0]['calculo_id'] in texto and 'Auditoría' in texto
    finally:w.close()
