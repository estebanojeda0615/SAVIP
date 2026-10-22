"""Formato durante edición real Qt; regresión de contratos contra el respaldo 2B."""

import os
import sys
import tempfile
import types
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from uuid import UUID

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PySide6.QtCore import Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtTest import QTest, QSignalSpy
from PySide6.QtWidgets import QStyle, QStyleOptionSpinBox
from openpyxl import load_workbook
from app_icociv.interfaz.widgets.numerico_icociv import NumeroICOCIV
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.tema import hoja_estilos
from app_icociv.persistencia.gestor_sesiones import guardar_sesion, cargar_sesion
from tests.test_proyecciones_redisenio_visual import _app, ventana_de_prueba

for nombre in ("segoeui.ttf", "segoeuib.ttf", "seguisb.ttf"):
    QFontDatabase.addApplicationFont(str(Path("C:/Windows/Fonts") / nombre))


def campo(decimales=4, maximo=1_000_000_000):
    w = NumeroICOCIV(0, maximo, decimales)
    w.resize(350, 48)
    w.show()
    w.activateWindow()
    w.setFocus()
    _app.processEvents()
    w.selectAll()
    return w


def pegar(w, texto):
    _app.clipboard().setText(texto)
    QTest.keyClick(w.lineEdit(), Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)


def test_escritura_caracter_a_caracter_antes_de_perder_foco():
    w = campo()
    try:
        esperados = ['1', '12', '125', '1,250', '12,500', '125,000', '1,250,000', '1,250,000.', '1,250,000.7', '1,250,000.75']
        cambios = QSignalSpy(w.valueChanged)
        for letra, esperado in zip('1250000.75', esperados):
            QTest.keyClicks(w.lineEdit(), letra)
            assert w.hasFocus() and w.text() == esperado, (letra, w.text())
            assert w.lineEdit().cursorPosition() == len(esperado)
        assert cambios.count() == 9  # El punto y el agrupado no crean cambios de valor.
        for entrada, esperado in (('1000', '1,000'), ('25000', '25,000'), ('1250000', '1,250,000'), ('1250000.5', '1,250,000.5'), ('1250000.0001', '1,250,000.0001'), ('12.', '12.'), ('12.00', '12.00')):
            w.selectAll()
            for letra in entrada:
                QTest.keyClicks(w.lineEdit(), letra)
            assert w.text() == esperado and w.hasFocus()
        w.clearFocus()
        _app.processEvents()
        assert w.text() == '12' and w.value() == 12
        w.setValue(1250.5)
        assert w.text() == '1,250.5'
    finally:
        w.close()


def test_cursor_edicion_intermedia_seleccion_y_borrado():
    w = campo()
    try:
        w.setValue(1234567)
        w.lineEdit().setCursorPosition(3)
        QTest.keyClicks(w.lineEdit(), '9')
        assert w.text() == '12,934,567' and w.value() == 12934567
        assert w.lineEdit().cursorPosition() == 4
        w.setValue(1234)
        w.lineEdit().setCursorPosition(2)
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_Backspace)
        assert w.text() == '234' and w.value() == 234 and w.lineEdit().cursorPosition() == 0
        w.setValue(1234)
        w.lineEdit().setCursorPosition(1)
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_Delete)
        assert w.text() == '134' and w.value() == 134 and w.lineEdit().cursorPosition() == 1
        w.setValue(1234567)
        w.lineEdit().setSelection(2, 3)
        QTest.keyClicks(w.lineEdit(), '9')
        assert w.text() == '19,567' and w.value() == 19567
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_Home)
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_Right, Qt.KeyboardModifier.ShiftModifier)
        assert w.lineEdit().selectedText() == '1'
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
        assert _app.clipboard().text() == '1'
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_X, Qt.KeyboardModifier.ControlModifier)
        assert w.text() == '9,567' and w.value() == 9567
    finally:
        w.close()


def test_pegado_real_formatos_y_ambiguedades_sin_mutacion():
    w = campo(2, 1_000_000_000_000_000)
    rechazos = QSignalSpy(w.entradaRechazada)
    try:
        for entrada in ('1234567.89', '1,234,567.89', '1.234.567,89'):
            w.selectAll()
            pegar(w, entrada)
            assert w.text() == '1,234,567.89' and w.value() == 1234567.89
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
        assert _app.clipboard().text() == '1,234,567.89'
        w.setValue(0)
        w.selectAll()
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)
        assert w.value() == 1234567.89
        # Fase 4I: «1 250» (espacio como separador de miles) pasó a ser un formato aceptado.
        # Fase 4I-A V4: «1,234» = 1234 (coma de miles); «1.234» y «10.001» son decimales que
        # exceden los 2 decimales de este campo y se rechazan sin mutar.
        w.selectAll(); pegar(w, '1,234'); assert w.value() == 1234
        for entrada in ('1.234', '10.001', '1.25,50', '-10', "1e6", '1000000000000001', '999999999999999.99'):
            w.selectAll()
            antes = (w.text(), w.value(), w.lineEdit().selectedText())
            n = rechazos.count()
            pegar(w, entrada)
            assert (w.text(), w.value(), w.lineEdit().selectedText()) == antes
            assert rechazos.count() == n + 1 and rechazos.at(n)[0]
        for entrada, esperado in (('1 250', '1,250'), ('$ 1.234.567', '1,234,567'), ('1.234.567', '1,234,567'), ('1,234,567', '1,234,567'), ('1000,50', '1,000.50'), ('1000.00', '1,000.00')):
            w.selectAll()
            pegar(w, entrada)
            assert w.text() == esperado
    finally:
        w.close()


def test_limites_precision_original_y_negativos():
    for decimales, maximo in ((4, 1_000_000_000), (2, 1_000_000_000_000_000)):
        w = campo(decimales, maximo)
        try:
            assert w.minimum() == 0 and w.maximum() == maximo and w.decimals() == decimales
            w.setValue(maximo)
            assert w.value() == maximo and "." not in w.text()
            w.setValue(0)
            w.selectAll()
            QTest.keyClicks(w.lineEdit(), '0.' + '0' * (decimales - 1) + '1')
            assert w.value() == 10 ** -decimales
            antes = w.text()
            QTest.keyClicks(w.lineEdit(), '1')
            assert w.text() == antes
            w.selectAll()
            pegar(w, '-1')
            assert w.text() == antes  # Los tres rangos originales son no negativos.
        finally:
            w.close()


def test_deshacer_rehacer_agrupado_y_carga_programatica():
    w = campo()
    try:
        QTest.keyClicks(w.lineEdit(), '1250000.00')
        for esperado in ('1,250,000.0', '1,250,000.', '1,250,000', '125,000'):
            QTest.keyClick(w.lineEdit(), Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
            assert w.text() == esperado
        for esperado in ('1,250,000', '1,250,000.', '1,250,000.0', '1,250,000.00'):
            QTest.keyClick(w.lineEdit(), Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)
            assert w.text() == esperado
        assert w.value() == 1250000
        w.setValue(99)
        QTest.keyClick(w.lineEdit(), Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        assert w.value() == 99 and w.text() == '99'
    finally:
        w.close()


def test_diez_cifras_dinamicas_en_los_tres_controles():
    w = WidgetProyeccionesICOCIV()
    w.show()
    w.activateWindow()
    try:
        esperados = ["1", "10", "100", "1,000", "10,000", "100,000", "1,000,000", "10,000,000", "100,000,000", "1,000,000,000"]
        for control in (w.cantidad, w.precio_base, w.anticipo):
            control.setFocus()
            control.selectAll()
            for letra, esperado in zip("1000000000", esperados):
                QTest.keyClicks(control.lineEdit(), letra)
                assert control.text() == esperado and control.lineEdit().cursorPosition() == len(esperado)
            assert control.value() == 1_000_000_000
            control.clearFocus()
            assert control.text() == "1,000,000,000"
        assert w.valores == [] and w.proyecciones == []
    finally:
        w.close()


def test_geometria_tres_campos_tamanos_temas_foco_y_sin_foco():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        for tema in ("claro", "oscuro"):
            v.setStyleSheet(hoja_estilos(tema))
            for ancho, alto in ((1100, 700), (1366, 768), (1920, 1080)):
                v.resize(ancho, alto)
                _app.processEvents()
                for control in (w.cantidad, w.precio_base, w.anticipo):
                    control.setValue(control.maximum())
                    for foco in (True, False):
                        control.setFocus() if foco else control.clearFocus()
                        _app.processEvents()
                        assert control.buttonSymbols() == control.ButtonSymbols.NoButtons
                        opcion = QStyleOptionSpinBox()
                        control.initStyleOption(opcion)
                        for sub in (QStyle.SubControl.SC_SpinBoxUp, QStyle.SubControl.SC_SpinBoxDown):
                            assert control.style().subControlRect(QStyle.ComplexControl.CC_SpinBox, opcion, sub, control).isEmpty()
                        editor = control.lineEdit()
                        margen = editor.textMargins()
                        disponible = editor.contentsRect().width() - margen.left() - margen.right() - 4
                        assert editor.fontMetrics().horizontalAdvance(editor.text()) <= disponible
    finally:
        v.close()


_CLAVES_4IC = {"factor_icociv", "r_total", "valor_proyectado", "diferencia_absoluta", "diferencia_porcentual",
               "trazabilidad_formula", "factor_icociv_indices", "criterio_factor"}


def sin_regla_4ic(valores):
    return [{k: v for k, v in r.items() if k not in _CLAVES_4IC} for r in valores]


def _sesion_sin_regla(datos):
    datos = dict(datos)
    datos["valores"] = sin_regla_4ic(datos.get("valores", []))
    return datos


def test_regresion_calculo_json_excel_antes_despues():
    from decimal import Decimal
    from app_icociv.servicios.actualizacion_icociv import factor_aplicado
    fuente = ROOT / "docs/fase_2c_b/respaldo_inicio/app_icociv__interfaz__widgets__proyecciones_icociv.py.bak"
    base = types.ModuleType("baseline_numericos")
    exec(compile(fuente.read_text(encoding="utf-8-sig"), str(fuente), "exec"), base.__dict__)
    fuente_numerico = ROOT / "docs/fase_2c_b/respaldo_inicio/app_icociv__interfaz__widgets__numerico_icociv.py.bak"
    numerico_anterior = types.ModuleType("numerico_2ca")
    exec(compile(fuente_numerico.read_text(encoding="utf-8-sig"), str(fuente_numerico), "exec"), numerico_anterior.__dict__)
    base.NumeroICOCIV = numerico_anterior.NumeroICOCIV
    viejo, nuevo = base.WidgetProyeccionesICOCIV(), WidgetProyeccionesICOCIV()
    fijo = datetime(2026, 10, 9, 12, 0)
    nuevo.modalidad_importe.setCurrentIndex(nuevo.modalidad_importe.findData('unitario'))
    from tests.compatibilidad_modalidad_4e_b import sin_modalidad
    try:
        for widget in (viejo, nuevo):
            widget.contexto = {"proyeccion_generada": True, "indice_base": 100., "indice_proyectado": 110., "serie_final": "Prueba", "ruta_texto": "Ruta de prueba", "periodo_base": "2026-12", "periodo_proyectado": "2027-06", "horizonte_solicitado": 6, "modelo": "Prueba"}
            widget._actualizar_estado_calculadora()
            widget.unidad.setCurrentIndex(1)
            widget.nombre_proyecto.setText("Proyecto de prueba")
            widget.nombre_componente.setText("Componente independiente")
        with patch("app_icociv.servicios.actualizacion_icociv.datetime") as reloj, \
             patch.object(base, "uuid4", return_value=UUID(int=1)), \
             patch("app_icociv.interfaz.widgets.proyecciones_icociv.uuid4", return_value=UUID(int=1)):
            reloj.now.return_value = fijo
            for cantidad, precio, anticipo in ((2., 1000., 100.), (1250.0001, 1234567.89, 25000.5), (1., 1e15, 0.)):
                for widget in (viejo, nuevo):
                    widget.show()
                    widget.activateWindow()
                    for control, valor in ((widget.cantidad, cantidad), (widget.precio_base, precio), (widget.anticipo, anticipo)):
                        control.setFocus()
                        control.selectAll()
                        texto = f"{valor:.{control.decimals()}f}".rstrip("0").rstrip(".")
                        QTest.keyClicks(control.lineEdit(), texto.replace(".", ",") if widget is viejo else texto)
                    widget.boton_calcular.click()
                assert all(r['modalidad_importe']=='unitario' for r in nuevo.valores)
                # Fase 4I-C: entradas, índices e IDs idénticos al respaldo 2B; los importes siguen la
                # regla nueva (factor desde la variación redondeada a 2 decimales), comprobada aparte.
                assert sin_regla_4ic(sin_modalidad(nuevo.valores)) == sin_regla_4ic(viejo.valores)
                for r in nuevo.valores:
                    p, f = factor_aplicado(r["i_icociv"], r["i0_icociv"])
                    assert (r["factor_icociv"], r["diferencia_porcentual"]) == (float(f), float(p))
                    assert r["valor_proyectado"] == float(Decimal(repr(r["base_ajustable"])) * f)
            assert nuevo.valores[0]["valor_proyectado"] == 990.0  # 2B: 990.0000000000001 con el factor crudo
        assert _sesion_sin_regla(sin_modalidad(nuevo.datos_sesion())) == _sesion_sin_regla(viejo.datos_sesion())
        with tempfile.TemporaryDirectory(prefix="savip_numericos_") as carpeta:
            carpeta = Path(carpeta)
            for nombre, widget in (("antes", viejo), ("despues", nuevo)):
                datos = widget.datos_sesion()
                datos['fecha_guardado'] = fijo.isoformat(timespec='seconds')
                guardar_sesion(carpeta / (nombre + ".json"), datos)
                widget._generar_excel(str(carpeta / (nombre + ".xlsx")))
            assert _sesion_sin_regla(cargar_sesion(carpeta / "antes.json")) == _sesion_sin_regla(sin_modalidad(cargar_sesion(carpeta / "despues.json")))
            a, b = load_workbook(carpeta / "antes.xlsx"), load_workbook(carpeta / "despues.xlsx")
            # Contrato 4D-B: cinco hojas; los originales permanecen en JSON y Trazabilidad.
            assert b.sheetnames == ['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
            for fila,r in enumerate(nuevo.valores,3):
                for col,key in ((6,'cantidad'),(7,'precio_base'),(8,'anticipo_amortizado'),(11,'i0_icociv'),(12,'i_icociv')):
                    assert b['Valores proyectados'].cell(fila,col).value == float(format(r[key],'.16g')) if r[key] is not None else b['Valores proyectados'].cell(fila,col).value is None
                assert b['Trazabilidad'].cell(fila,16).value == float(format(r['valor_proyectado'],'.16g'))
                assert b['Trazabilidad'].cell(fila,15).value == float(format(r['r_total'],'.16g'))
                # La sesión legada conserva sus importes sin atribuirles unidad o total.
                assert r['modalidad_importe']=='unitario'
                assert all(b['Valores proyectados'].cell(fila,col).data_type=='f' for col in (9,10,14,15))
                for col,formula in {13:f'=1+ROUND((L{fila}/K{fila}-1)*100,2)/100',16:f'=ROUND((M{fila}-1)*100,2)'}.items():
                    assert b['Valores proyectados'].cell(fila,col).value == formula and b['Valores proyectados'].cell(fila,col).data_type == 'f'
            a.close()
            b.close()
    finally:
        viejo.close()
        nuevo.close()


if __name__ == "__main__":
    for nombre, funcion in sorted(globals().copy().items()):
        if nombre.startswith("test_") and callable(funcion):
            funcion()
            print("OK " + nombre, flush=True)
