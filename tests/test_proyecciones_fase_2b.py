"""Fase 2B: Qt real, JSON y Excel temporales; integración con series sintéticas."""

from __future__ import annotations

import json
import math
import os
import sys
import tempfile
import time
import types
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication  # Fase 4I-G: la apertura de sesiones termina en el bucle de eventos
from PySide6.QtCore import Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtTest import QTest
from tests.test_proyecciones_redisenio_visual import _app, ventana_de_prueba
from app_icociv.interfaz.tema import hoja_estilos
from app_icociv.interfaz.widgets.proyecciones_icociv import COLUMNAS_PROYECCIONES, COLUMNAS_VALORES, TIPOS_COMPONENTE
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.servicios.actualizacion_icociv import calcular_actualizacion_icociv
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.utilidades.utilidades import COLUMNAS_FIJAS_TABLAS

for nombre in ("segoeui.ttf", "segoeuib.ttf", "seguisb.ttf", "seguiemj.ttf"):
    fuente = Path("C:/Windows/Fonts") / nombre
    if fuente.exists():
        QFontDatabase.addApplicationFont(str(fuente))


def llenar_y_calcular(v, tipo="Capítulo", componente="Pavimento rígido", serie="Serie sintética A"):
    w = v.widget_proyecciones
    w.contexto = {"proyeccion_generada": True, "indice_base": 100.0, "indice_proyectado": 110.0,
                  "ruta_texto": serie, "serie_final": "Ítem sintético", "periodo_base": "2026-12",
                  "periodo_proyectado": "2027-06", "horizonte_solicitado": 6, "modelo": "Prueba sin pronóstico"}
    w._actualizar_estado_calculadora()
    w.tipo_componente.setCurrentText(tipo)
    w.nombre_componente.setText(componente)
    w.unidad.setCurrentIndex(1)
    w.cantidad.setValue(2)
    w.precio_base.setValue(1000)
    w.anticipo.setValue(100)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('unitario')); w.boton_calcular.click()
    return w.valores[-1]


def llenar_proyecto(w):
    w.nombre_proyecto.setText("Proyecto de prueba — sector norte")
    w.objeto_proyecto.setPlainText("Objeto con texto largo\ny segunda línea.")
    w.responsable_tecnico.setText("Responsable técnico de prueba")
    w.observacion_general.setPlainText("Observación común: ñ y símbolos < & >.")


def test_proyecto_componente_registros_y_calculo_sin_regresion():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        llenar_proyecto(w)
        seleccion = v._seleccion_actual().copy()
        for tipo, etiqueta in TIPOS_COMPONENTE.items():
            w.tipo_componente.setCurrentText(tipo)
            w.nombre_componente.setText("Texto libre diferente de ICOCIV")
            assert w.lbl_nombre_componente.text() == etiqueta
            assert v._seleccion_actual() == seleccion
        primero = deepcopy(llenar_y_calcular(v))
        segundo = llenar_y_calcular(v, "Tipología", "Drenajes", "Serie sintética B")
        assert w.valores[0] == primero
        assert segundo["nombre_componente"] == "Drenajes" and primero["nombre_componente"] == "Pavimento rígido"
        assert segundo["ruta_icociv"] != primero["ruta_icociv"]
        referencia = calcular_actualizacion_icociv({"precio_base": 1000, "anticipo_amortizado": 100,
                                                    "i0_icociv": 100, "i_icociv": 110, "cantidad": 2})
        assert primero["valor_proyectado"] == referencia["valor_proyectado"]
        assert math.isclose(primero["valor_proyectado"], 990, rel_tol=0, abs_tol=1e-12)
        assert primero["cantidad"] == 2
        assert list(w.modelo_valores._dataframe.columns) == [t for _, t in COLUMNAS_VALORES]
        assert w.modelo_valores.columnCount() == 17
        w.tabla_valores.selectRow(0)
        formulario = w._datos_calculadora()
        w._detalle_valor()
        # Fase 4G-B: «Ver detalle» es consulta pura; no escribe en la calculadora.
        assert w._datos_calculadora() == formulario
        assert w.valores[0] == primero
        assert "ruta icociv" in w.resultado.toPlainText()
        assert w.nombre_proyecto.text() == primero["proyecto"]
    finally:
        v.close()


def test_guardar_cerrar_reabrir_y_sesion_antigua():
    with tempfile.TemporaryDirectory(prefix="savip_2b_") as carpeta:
        carpeta = Path(carpeta)
        fuente = carpeta / "fuente_sintetica.xlsx"
        Workbook().save(fuente)
        sesion = carpeta / "sesion.json"
        v = ventana_de_prueba()
        w = v.widget_proyecciones
        llenar_proyecto(w)
        llenar_y_calcular(v)
        llenar_y_calcular(v, "Actividad", "Replanteo", "Serie B")
        w.valores[0]["campo_legacy_adicional"] = {"valor": 0.12345678901234567, "texto": "conservar"}
        w.proyecciones = [{"numero": 1, "item": "Registro técnico original", "rmse": "0.123456789", "extra": [1, 2]}]
        w._refrescar_tabla_proyecciones()
        esperado = w.datos_sesion()
        v.controlador.ruta_archivo = fuente
        v.controlador.serie_actual = pd.DataFrame({"Periodo": ["2026_12"], "Indice": [100.0]})
        v.controlador.proyeccion_actual = {"periodo_proj": "2027-06", "y_proj": 110.0,
                                          "stats": {}, "backtesting": {"metricas": {"rmse": 0.12345678901234567}}}
        with patch.object(v, "_solicitar_usuario", return_value="Prueba"), \
             patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
             patch("app_icociv.interfaz.ventana_principal.QFileDialog.getSaveFileName", return_value=(str(sesion), "")):
            v.guardar_sesion_actual(como=True)
        guardado = cargar_sesion(sesion)
        assert guardado["proyecciones_icociv"] == esperado
        v.close()
        v.deleteLater()
        _app.processEvents()
        nuevo = ventana_de_prueba()
        nuevo.controlador.periodos = ["2026_11", "2026_12"]  # DANE sintético ya cargado (paso 1), corte 2026-12
        try:
            def cargar_sintetico(_ruta):
                raise AssertionError("La sesión no debe abrir su Excel original")
            with patch.object(nuevo, "_cargar_archivo", side_effect=cargar_sintetico), \
                 patch.object(nuevo, "ejecutar_proyeccion") as proyectar, \
                 patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
                 patch("app_icociv.interfaz.ventana_principal.QFileDialog.getOpenFileName", return_value=(str(sesion), "")):
                nuevo.abrir_sesion()
                while nuevo._operacion_sesion: QCoreApplication.processEvents()
                assert proyectar.call_count == 0  # Abrir una sesión no recalcula nada.
            assert nuevo.widget_proyecciones.datos_sesion() == esperado
            assert nuevo.widget_proyecciones.valores[0]["campo_legacy_adicional"] == esperado["valores"][0]["campo_legacy_adicional"]
            assert not nuevo._restaurando_registros_icociv  # La prueba aislada no inicia un trabajador.
            antiguos = deepcopy(guardado)
            del antiguos["proyecciones_icociv"]
            ruta_antigua = carpeta / "v1_original.json"
            ruta_antigua.write_text(json.dumps(antiguos, ensure_ascii=False), encoding="utf-8")
            bytes_originales = ruta_antigua.read_bytes()
            with patch.object(nuevo, "_cargar_archivo", side_effect=cargar_sintetico), \
                 patch.object(nuevo, "ejecutar_proyeccion") as proyectar, \
                 patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
                 patch("app_icociv.interfaz.ventana_principal.QFileDialog.getOpenFileName", return_value=(str(ruta_antigua), "")), \
                 patch.object(nuevo, "_decidir_cierre", return_value="descartar"):
                nuevo.abrir_sesion()
                while nuevo._operacion_sesion: QCoreApplication.processEvents()
                assert proyectar.call_count == 0
            assert nuevo.widget_proyecciones.datos_proyecto() == {"nombre_proyecto": "", "objeto": "", "responsable_tecnico": "", "observacion_general": ""}
            assert not nuevo._restaurando_registros_icociv
            assert ruta_antigua.read_bytes() == bytes_originales
            assert cargar_sesion(ruta_antigua)["resultado_proyeccion"] == guardado["resultado_proyeccion"]
        finally:
            nuevo.close()


def test_excel_operativa_y_regresion_hoja_tecnica():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        llenar_proyecto(w)
        llenar_y_calcular(v)
        llenar_y_calcular(v, "Otro", "=SUM(1,2)")
        for registro in w.valores:
            registro['modalidad_importe'] = 'unitario'
        w.valores[0]["adicional"] = {"precision": 0.12345678901234567}
        with tempfile.TemporaryDirectory(prefix="savip_excel_") as carpeta:
            carpeta = Path(carpeta)
            actual = carpeta / "actual.xlsx"
            with patch("app_icociv.interfaz.widgets.proyecciones_icociv.QFileDialog.getSaveFileName", return_value=(str(actual), "")), \
                 patch("app_icociv.interfaz.widgets.proyecciones_icociv.QMessageBox.information"):
                w._exportar_excel()
            libro = load_workbook(actual, data_only=False)
            encabezados = [c.value for c in libro["Valores proyectados"][2]]
            assert encabezados[:10] == [t for _,t in COLUMNAS_VALORES][:10] and encabezados[12:] == [t for _,t in COLUMNAS_VALORES][12:]
            # Fase 4G: la fecha única va después del título técnico ("I0 ICOCIV\nAAAA-MM").
            assert encabezados[10].startswith('I0 ICOCIV\n') and encabezados[11].startswith('I ICOCIV proyectado\n')
            assert libro["Valores proyectados"].cell(3, 14).data_type == 'f'
            assert libro["Trazabilidad"].cell(3, 16).value == w.valores[0]["valor_proyectado"]
            assert libro["Valores proyectados"].cell(4, 3).data_type == "s"
            assert any(c.value==w.nombre_proyecto.text() for row in libro["Resumen"] for c in row)
            assert w.datos_sesion()["valores"][0] == w.valores[0] and w.valores[0]["adicional"]["precision"] == 0.12345678901234567
            assert libro["Proyecciones ICOCIV"].max_column == 21
            assert libro.sheetnames == ['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
            for fila,r in enumerate(w.valores,3):
                assert libro['Trazabilidad'].cell(fila,8).value == r['i0_icociv']
                assert libro['Valores proyectados'].cell(fila,14).value == f'=I{fila}+I{fila}*(M{fila}-1)'
            libro.close()
    finally:
        v.close()


def test_tablas_vacias_filas_y_distribucion_real():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        assert w.tabla_valores.isHidden() and not w.sin_valores.isHidden()
        assert w.tabla_proyecciones.isHidden() and not w.sin_proyecciones.isHidden()
        assert v.scroll_ruta_jerarquica.maximumHeight() == 40
        llenar_proyecto(w)
        llenar_y_calcular(v)
        assert not w.tabla_valores.isHidden() and w.sin_valores.isHidden()
        altura_una = w.tabla_valores.height()
        llenar_y_calcular(v, "Actividad", "Texto extenso " * 20)
        assert w.tabla_valores.height() > altura_una
        v.combo_grupo.setCurrentIndex(1)
        for ancho, alto in ((1100, 700), (1366, 768), (1920, 1080)):
            v.resize(ancho, alto)
            for tema in ("claro", "oscuro"):
                v._aplicar_tema(tema, persistir=False)
                _app.processEvents()
                assert w.modelo_valores.columnCount() == 17
                assert all(not w.tabla_valores.isColumnHidden(i) for i in range(17))
                for boton in v.botones_periodo:
                    assert boton.width() >= 38
        w.valores.clear()
        w._refrescar_tabla_valores()
        assert w.tabla_valores.isHidden() and not w.sin_valores.isHidden()
    finally:
        v.close()


def test_sesion_malformada_no_reemplaza_metadatos():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        llenar_proyecto(w)
        antes = w.datos_sesion()
        for invalido in ({"valores": "no es una lista"}, {"calculadora": {"precio_base": "texto"}}, {"proyecto": {"nombre_proyecto": 123}}, {"calculadora": {"precio_base": w.precio_base.maximum() * 2}}):
            try:
                w.validar_datos_sesion(invalido)
                w.restaurar_datos_sesion(invalido)
            except ValueError:
                pass
            else:
                raise AssertionError("Debió rechazar el bloque inválido sin cambiar la sesión.")
            assert w.datos_sesion() == antes
    finally:
        v.close()


def test_integracion_sesion_con_excel_y_trabajadores_reales():
    def esperar(v, condicion):
        fin = time.monotonic() + 45
        while time.monotonic() < fin:
            _app.processEvents()
            trabajador = v.trabajador
            if condicion() and (trabajador is None or not trabajador.isRunning()):
                return
            QTest.qWait(20)
        raise AssertionError("No finalizó el flujo real: " + v.etiqueta_estado.text())

    with tempfile.TemporaryDirectory(prefix="savip_integracion_2b_") as carpeta:
        carpeta = Path(carpeta)
        excel = carpeta / "anexo_sintetico.xlsx"
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Anexo 16"
        periodos = [f"{p.year}_{p.month}" for p in pd.period_range("2024-01", periods=36, freq="M")]
        for tag, columnas in COLUMNAS_FIJAS_TABLAS.items():
            hoja.append([tag + ". DEMOSTRACION - datos sintéticos"])
            hoja.append(columnas + periodos)
            valores = []
            for columna in columnas:
                if "Ponderacion" in columna:
                    valores.append(100.0)
                elif "cod" in columna.lower():
                    valores.append(1)
                else:
                    valores.append(columna + " sintético")
            hoja.append(valores + [100 + i * 0.2 + math.sin(i / 3) * 0.1 for i in range(36)])
            for _ in range(3):
                hoja.append([None])
        libro.save(excel)
        v = VentanaPrincipal()
        v._cargar_archivo(str(excel))
        esperar(v, lambda: v.combo_grupo.count() > 1 and v.combo_grupo.isEnabled())
        v.combo_grupo.setCurrentIndex(1)
        v.ejecutar_proyeccion()
        esperar(v, lambda: v.resultado_ui_actual is not None)
        w = v.widget_proyecciones
        assert w.contexto["proyeccion_generada"]
        llenar_proyecto(w)
        for tipo, nombre in (("Capítulo", "Pavimento rígido"), ("Actividad", "Replanteo")):
            w.tipo_componente.setCurrentText(tipo)
            w.nombre_componente.setText(nombre)
            w.unidad.setCurrentIndex(1)
            w.precio_base.setValue(1000)
            w.anticipo.setValue(100)
            w.cantidad.setValue(2)
            w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('unitario')); w.boton_calcular.click()
        assert len(w.valores) == 2
        sesion = carpeta / "real.json"
        with patch.object(v, "_solicitar_usuario", return_value="Integración"), \
             patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
             patch("app_icociv.interfaz.ventana_principal.QFileDialog.getSaveFileName", return_value=(str(sesion), "")):
            v.guardar_sesion_actual(como=True)
        guardado = cargar_sesion(sesion)
        v.close()
        nuevo = VentanaPrincipal()
        nuevo._archivo_cargado(nuevo.controlador.cargar_archivo(guardado["archivo_excel"]))  # paso 1: DANE
        try:
            with patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
                 patch("app_icociv.interfaz.ventana_principal.QFileDialog.getOpenFileName", return_value=(str(sesion), "")):
                nuevo.abrir_sesion()
                while nuevo._operacion_sesion: QCoreApplication.processEvents()
            esperar(nuevo, lambda: nuevo.resultado_ui_actual is not None and not nuevo._restaurando_registros_icociv)
            actual = nuevo.widget_proyecciones
            assert actual.datos_proyecto() == guardado["proyecciones_icociv"]["proyecto"]
            assert actual.valores == guardado["proyecciones_icociv"]["valores"]
            assert actual.proyecciones == guardado["proyecciones_icociv"]["proyecciones"]
            assert len(actual.valores) == 2 and len(actual.proyecciones) == 1
            copia = carpeta / "reabierta.json"
            with patch.object(nuevo, "_solicitar_usuario", return_value="Integración"), \
                 patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
                 patch("app_icociv.interfaz.ventana_principal.QFileDialog.getSaveFileName", return_value=(str(copia), "")):
                nuevo.guardar_sesion_actual(como=True)
            assert cargar_sesion(copia)["proyecciones_icociv"]["valores"] == guardado["proyecciones_icociv"]["valores"]
            antigua = deepcopy(guardado)
            del antigua["proyecciones_icociv"]
            ruta_antigua = carpeta / "v1_sin_metadatos.json"
            ruta_antigua.write_text(json.dumps(antigua, ensure_ascii=False), encoding="utf-8")
            bytes_antiguos = ruta_antigua.read_bytes()
            legacy = VentanaPrincipal()
            legacy._archivo_cargado(legacy.controlador.cargar_archivo(guardado["archivo_excel"]))
            try:
                with patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", carpeta), \
                     patch("app_icociv.interfaz.ventana_principal.QFileDialog.getOpenFileName", return_value=(str(ruta_antigua), "")):
                    legacy.abrir_sesion()
                    while legacy._operacion_sesion: QCoreApplication.processEvents()
                # V1.0 no guardaba análisis: se recuperan selección y parámetros; proyectar es decisión del usuario.
                assert legacy.controlador.proyeccion_actual is None
                legacy.ejecutar_proyeccion()
                esperar(legacy, lambda: legacy.resultado_ui_actual is not None)
                assert legacy.widget_proyecciones.nombre_proyecto.text() == ""
                assert legacy.widget_proyecciones.valores == []  # V1.0 no guardaba estas filas.
                assert ruta_antigua.read_bytes() == bytes_antiguos
                assert legacy.controlador.proyeccion_actual["y_proj"] == guardado["resultado_proyeccion"]["y_proj"]
            finally:
                legacy.close()
        finally:
            nuevo.close()


if __name__ == "__main__":
    for nombre, funcion in sorted(globals().copy().items()):
        if nombre.startswith("test_") and callable(funcion):
            funcion()
            print(f"OK {nombre}", flush=True)
