"""Historial con Qt real y motor intacto sobre anexo sintético temporal."""
from __future__ import annotations

import json
import math
import os
import sys
import tempfile
import time
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication  # Fase 4I-G: la apertura de sesiones termina en el bucle de eventos
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QScrollArea, QPushButton
from matplotlib.figure import Figure
from tests.test_proyecciones_fase_2b import _app, llenar_proyecto
from tests.test_proyecciones_redisenio_visual import ventana_de_prueba
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.tema.colores import paleta, contraste
from app_icociv.persistencia.historial_icociv import serializar_resultado, recuperar_resultado
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.utilidades.utilidades import COLUMNAS_FIJAS_TABLAS
from tests.fixtures.salidas import carpeta_salida


def esperar(v, condicion):
    fin = time.monotonic() + 90
    while time.monotonic() < fin:
        _app.processEvents()
        if condicion() and (v.trabajador is None or not v.trabajador.isRunning()):
            return
        QTest.qWait(20)
    raise AssertionError("No finalizó el flujo real: " + v.etiqueta_estado.text())


def anexo_sintetico(ruta):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Anexo 16"
    periodos = [f"{p.year}_{p.month}" for p in pd.period_range("2024-01", periods=36, freq="M")]
    for tag, columnas in COLUMNAS_FIJAS_TABLAS.items():
        hoja.append([tag + ". PRUEBA - datos sintéticos"])
        hoja.append(columnas + periodos)
        for grupo in (1, 2):
            valores = [100.0 if "Ponderacion" in c else grupo if "cod" in c.lower()
                       else f"{c} sintético {grupo}" for c in columnas]
            hoja.append(valores + [100 + grupo + i * .2 + math.sin(i / 3) * .1 for i in range(36)])
        for _ in range(3):
            hoja.append([None])
    libro.save(ruta)


def guardar(v, ruta):
    with patch.object(v, "_solicitar_usuario", return_value="Prueba Fase 3"), \
         patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", ruta.parent), \
         patch("app_icociv.interfaz.ventana_principal.QFileDialog.getSaveFileName", return_value=(str(ruta), "")):
        v.guardar_sesion_actual(como=True)  # destino explícito (no el JSON abierto)


def abrir(v, ruta, dane=None):
    # Flujo vigente: primero se carga el Excel DANE (por defecto, el mismo anexo de la sesión) y luego la sesión.
    dane = dane or cargar_sesion(ruta)["archivo_excel"]
    v._archivo_cargado(v.controlador.cargar_archivo(str(dane)))
    with patch("app_icociv.interfaz.ventana_principal.CARPETA_SESIONES", ruta.parent), \
         patch("app_icociv.interfaz.ventana_principal.QFileDialog.getOpenFileName", return_value=(str(ruta), "")):
        v.abrir_sesion()
        while v._operacion_sesion: QCoreApplication.processEvents()
    esperar(v, lambda: v.resultado_ui_actual is not None and v.sesion_pendiente is None)


def calcular(w, componente):
    w.tipo_componente.setCurrentText("Actividad")
    w.nombre_componente.setText(componente)
    w.unidad.setCurrentIndex(1)
    w.cantidad.setValue(2)
    w.precio_base.setValue(1000)
    w.anticipo.setValue(100)
    w.modalidad_importe.setCurrentIndex(w.modalidad_importe.findData('unitario')); w.boton_calcular.click()


def capturar(v, nombre, widget=None):
    out = carpeta_salida("fase_3_capturas")  # Fase 1A (AS-007): antes en docs/, no publicado
    out.mkdir(parents=True, exist_ok=True)
    v.resize(1366, 768)
    v._cambiar_seccion(1)
    v.show()
    _app.processEvents()
    scroll = v.widget_proyecciones.findChild(QScrollArea)
    if widget:
        scroll.ensureWidgetVisible(widget, 20, 100)
    else:
        scroll.verticalScrollBar().setValue(0)
    v.notificaciones.cerrar_todas()
    QTest.qWait(400)
    assert not v.velo_carga.isVisible()
    assert v.grab().save(str(out / (nombre + ".png")))


def test_serializacion_exacta_sin_codigo_ejecutable():
    df = pd.DataFrame({"entero": pd.Series([2**53 + 1, 7], dtype="int64"),
                       "real": [float("nan"), .12345678901234567], "texto": ["ñ", "<>&"]})
    original = {"df": df, "array": np.array([1.1, float("inf")]), "pares": {1: (2, 3)},
                "predict_func": lambda _: 0, "neg_inf": float("-inf")}
    serial = serializar_resultado(original)
    vuelta = recuperar_resultado(json.loads(json.dumps(serial, allow_nan=False)))
    pd.testing.assert_frame_equal(vuelta["df"], df)
    np.testing.assert_array_equal(vuelta["array"], original["array"])
    assert vuelta["pares"] == original["pares"] and vuelta["predict_func"] is None
    assert vuelta["neg_inf"] == float("-inf")
    try:
        serializar_resultado({"otra_funcion": lambda _: 0})
    except ValueError:
        pass
    else:
        raise AssertionError("No debe serializar código ejecutable.")


def test_sidebar_titulo_y_historial_vacio():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        assert not w.historial and w.selector_analisis.count() == 1
        assert "historial_analisis" not in w.datos_sesion()
        assert any(b.text() == "Datos generales del proyecto" for b in w.findChildren(QPushButton))
        for tema in ("claro", "oscuro"):
            v._aplicar_tema(tema, persistir=False)
            p = paleta(tema)
            assert contraste(p["navegacion_texto"], p["barra_lateral"]) >= 4.5
            assert contraste(p["navegacion_secundario"], p["barra_lateral"]) >= 4.5
            v.cabecera_app.establecer_estado_archivo("PRUEBA UI · datos sintéticos", "cargado")
            capturar(v, "historial_vacio_" + tema)
    finally:
        v.close()


def test_recuperar_ruta_profunda_con_checkbox_dinamicos():
    v = ventana_de_prueba()
    try:
        w = v.widget_proyecciones
        v.controlador.ruta_archivo = ROOT / "fuente_sintetica_jerarquia.xlsx"
        v.combo_grupo.setCurrentIndex(1)
        v.combo_nivel2.setCurrentIndex(1)
        v.combo_nivel3.setCurrentIndex(1)
        v.combo_nivel4.setCurrentIndex(1)
        v.chk_t16_3.setChecked(True)
        v.combo_nivel5.setCurrentIndex(1)
        v.combo_nivel6.setCurrentIndex(1)
        seleccion = deepcopy(v._seleccion_actual())
        ruta = v.controlador.construir_ruta_jerarquica(seleccion)
        # La prueba cubre la restauración UI, no inventa una proyección técnica.
        referencia = {"fuente": "T_16_13", "ruta_jerarquica": ruta}
        analisis = {"id": "ruta-profunda", "nombre": "Prueba de restauración jerárquica", "creado": "2026-10-09",
                    "archivo_excel": str(v.controlador.ruta_archivo), "seleccion": seleccion,
                    "parametros": {"anio": 2027, "mes": 1, "horizonte": 1},
                    "resultado": serializar_resultado(referencia), "proyecto": w.datos_proyecto(),
                    "calculadora": w._datos_calculadora(), "registros_ids": []}
        w.historial = [analisis]
        w._refrescar_historial()
        v.combo_grupo.setCurrentIndex(2)
        assert v.combo_nivel6.isHidden() and not v.chk_t16_3.isChecked()
        with patch.object(v, "_mostrar_resultado_guardado") as mostrar, \
             patch.object(v.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor en restauración UI")) as motor:
            w.selector_analisis.setCurrentIndex(1)
            assert mostrar.call_count == 1 and motor.call_count == 0
        assert v._seleccion_actual() == seleccion
        assert v.chk_t16_3.isChecked() and not v.chk_t16_3.isHidden()
        assert not v.combo_nivel6.isHidden() and v.lbl_nivel6.text() == "Insumo"
        assert not v.chk_t16_1.isHidden() and not v.chk_t16_2.isHidden()
        assert w.analisis_activo_id == "ruta-profunda"
    finally:
        v.close()


def test_tres_analisis_motor_real_recuperacion_sesion_excel_y_regresion():
    with tempfile.TemporaryDirectory(prefix="savip_fase3_") as tmp:
        carpeta = Path(tmp)
        fuente = carpeta / "anexo_sintetico.xlsx"
        anexo_sintetico(fuente)
        v = VentanaPrincipal()
        nuevo = None
        try:
            v._aplicar_tema("claro", persistir=False)
            v._cargar_archivo(str(fuente))
            esperar(v, lambda: v.combo_grupo.count() == 3 and v.combo_grupo.isEnabled())
            w = v.widget_proyecciones
            llenar_proyecto(w)
            for numero, (grupo, horizonte) in enumerate(((1, 1), (1, 3), (1, 1)), 1):
                v.combo_grupo.setCurrentIndex(grupo)
                v.combo_horizonte.setCurrentIndex(v.combo_horizonte.findData(horizonte))
                # Desde la Fase 4I-J la misma serie con los mismos datos reutiliza su trayectoria; para reproducir una
                # sesión con varias ejecuciones de la misma serie (creada con versiones anteriores) se fuerza el motor.
                with patch.object(v, "_analisis_equivalente", return_value=(None, None)):
                    v.ejecutar_proyeccion()
                esperar(v, lambda: len(w.historial) == numero)
                assert w.contexto["proyeccion_generada"]
                calcular(w, f"Actividad {numero}")
                assert w.valores[-1]["analisis_id"] == w.historial[-1]["id"]
            assert len(w.proyecciones) == 3  # Incluso la misma fecha no sobrescribe una revisión previa.
            assert len({a["id"] for a in w.historial}) == len({a["nombre"] for a in w.historial}) == 3
            assert w.historial[0]["contexto"]["serie_final"] == w.historial[1]["contexto"]["serie_final"]
            assert w.historial[0]["parametros"]["horizonte"] != w.historial[1]["parametros"]["horizonte"]
            ids = [a["id"] for a in w.historial]
            resultados = [deepcopy(a["resultado"]) for a in w.historial]
            operativas = deepcopy(w.valores)
            capturar(v, "historial_tres_analisis_claro")
            w.selector_analisis.showPopup()
            QTest.qWait(100)
            assert w.selector_analisis.view().window().grab().save(str(
                ROOT / "docs/fase_3/capturas/selector_desplegado_tres_analisis.png"))
            w.selector_analisis.hidePopup()
            with patch.object(v.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor invocado al recuperar")) as motor, \
                 patch.object(v, "ejecutar_proyeccion", side_effect=AssertionError("Ejecutar invocado al recuperar")):
                w.solo_analisis_activo.setChecked(True)
                for i in (0, 1, 2, 0, 2, 1, 0):
                    w.selector_analisis.setCurrentIndex(w.selector_analisis.findData(ids[i]))
                    assert w.analisis_activo_id == ids[i]
                    assert w.contexto == w.historial[i]["contexto"]
                    assert serializar_resultado(v.resultado_ui_actual) == resultados[i]
                    assert v._seleccion_actual() == w.historial[i]["seleccion"]
                    assert v._horizonte_solicitado() == w.historial[i]["parametros"]["horizonte"]
                    assert w.modelo_valores.rowCount() == 1
                    assert w._valores_visibles[0]["analisis_id"] == ids[i]
                assert motor.call_count == 0
                assert w.valores == operativas and len(w.proyecciones) == 3
                v._dibujar_grafica(Figure(), v.controlador.serie_actual, v.controlador.proyeccion_actual)
                capturar(v, "analisis_1_activo_claro")
                w.selector_analisis.setCurrentIndex(w.selector_analisis.findData(ids[1]))
                v._aplicar_tema("oscuro", persistir=False)
                capturar(v, "analisis_2_activo_oscuro")
                calcular(w, "Registro adicional desde análisis recuperado")
                assert w.valores[-1]["analisis_id"] == ids[1]
                assert w.modelo_valores.rowCount() == 2 and len(w.valores) == 4
                w.solo_analisis_activo.setChecked(False)
                # Fase 4G: objetivos 2027-01 y 2027-03 con I0 único → +1 columna de fecha de I proyectado.
                assert w.modelo_valores.rowCount() == 4 and w.modelo_valores.columnCount() == 18
                capturar(v, "tabla_todos_analisis_oscuro", w.tabla_valores)
                # Borrar desde el filtro elimina la fila visible, nunca otra fila global.
                w.solo_analisis_activo.setChecked(True)
                w.tabla_valores.selectRow(1)
                w._detalle_valor()
                assert w.nombre_componente.text() == "Registro adicional desde análisis recuperado"
                w._eliminar_valor()
                assert len(w.valores) == 3 and w.valores == operativas
                calcular(w, "Registro final de revisión")
                esperado = w.datos_sesion()
                w.validar_datos_sesion(esperado)
                ruta = carpeta / "historial.json"
                guardar(v, ruta)
                datos_guardados = cargar_sesion(ruta)
                assert datos_guardados["proyecciones_icociv"]["historial_analisis"] == esperado["historial_analisis"]
            v.close()
            v.deleteLater()
            _app.processEvents()
            nuevo = VentanaPrincipal()
            with patch.object(nuevo.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor al reabrir")) as motor, \
                 patch.object(nuevo, "ejecutar_proyeccion", side_effect=AssertionError("Ejecutar al reabrir")):
                abrir(nuevo, ruta)
                actual = nuevo.widget_proyecciones
                guardado = datos_guardados["proyecciones_icociv"]
                assert actual.datos_sesion() == guardado
                assert serializar_resultado(nuevo.resultado_ui_actual) == resultados[1]
                assert len(actual.proyecciones) == 3 and len(actual.historial) == 3
                assert actual.modelo_valores.rowCount() == 2
                nuevo._aplicar_tema("claro", persistir=False)
                capturar(nuevo, "sesion_reabierta_claro")
                nuevo._aplicar_tema("oscuro", persistir=False)
                capturar(nuevo, "sesion_reabierta_oscuro", actual.tabla_valores)
                copia = carpeta / "copia.json"
                guardar(nuevo, copia)
                assert cargar_sesion(copia)["proyecciones_icociv"] == guardado
                for i in (2, 0, 1):
                    actual.selector_analisis.setCurrentIndex(actual.selector_analisis.findData(ids[i]))
                    assert serializar_resultado(nuevo.resultado_ui_actual) == resultados[i]
                assert motor.call_count == 0
                xlsx = carpeta / "historial.xlsx"
                actual._generar_excel(str(xlsx))
                libro = load_workbook(xlsx, data_only=False)
                assert libro["Valores proyectados"].max_column == 18  # 4G: fecha de I proyectado heterogénea
                assert libro["Valores proyectados"].max_row == len(actual.valores) + 2
                assert libro.sheetnames == ['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
                for i,r in enumerate(actual.valores,3):
                    assert libro['Trazabilidad'].cell(i,3).value == r['analisis_id']
                    assert libro['Trazabilidad'].cell(i,2).value == r['calculo_id']
                    assert libro['Valores proyectados'].cell(i,14).data_type == 'f'  # 4G: factor tras la columna de fecha
                    assert libro['Trazabilidad'].cell(i,16).value == r['valor_proyectado']
                assert actual.datos_sesion()['historial_analisis'] == actual.historial
                libro.close()
                validacion = {"analisis": len(actual.historial), "registros": len(actual.valores),
                    "misma_serie_horizontes": [1, 3, 1], "recuperacion_motor_llamadas": motor.call_count,
                    "sesion_exacta": True, "excel_json_completo_exacto": True,
                    "sin_duplicacion_tecnica": True, "datos": "Anexo sintético; motor real intacto",
                    "sesion_bytes": ruta.stat().st_size, "excel_bytes": xlsx.stat().st_size}
                (carpeta_salida("fase_3") / "INTEGRACION_REAL.json").write_text(json.dumps(validacion, ensure_ascii=False, indent=2), encoding="utf-8")
            # La ruta de Excel guardada en la sesión ya no existe: no se busca ni se abre; con un DANE
            # compatible cargado antes, la sesión se recupera desde sus instantáneas sin ejecutar el motor.
            sin_fuente = deepcopy(datos_guardados)
            sin_fuente["archivo_excel"] = str(carpeta / "fuente_no_disponible.xlsx")
            ruta_sin_fuente = carpeta / "sin_fuente.json"
            ruta_sin_fuente.write_text(json.dumps(sin_fuente, ensure_ascii=False, allow_nan=False), encoding="utf-8")
            offline = VentanaPrincipal()
            try:
                with patch.object(offline.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor offline")) as motor_offline, \
                     patch.object(offline, "_cargar_archivo", side_effect=AssertionError("Carga de archivo inexistente")):
                    abrir(offline, ruta_sin_fuente, dane=datos_guardados["archivo_excel"])
                    assert offline.widget_proyecciones.datos_sesion() == datos_guardados["proyecciones_icociv"]
                    assert offline.boton_guardar_sesion.isEnabled()
                    # Los selectores se restauran sobre el DANE cargado; la ruta inexistente de la sesión no se usa.
                    assert str(offline.controlador.ruta_archivo) == datos_guardados["archivo_excel"]
                    assert motor_offline.call_count == 0
                    guardar(offline, carpeta / "copia_offline.json")
                    capturar(offline, "historial_sin_excel_original")
            finally:
                offline.close()
            # Datos corruptos se rechazan antes de reemplazar la sesión viva.
            antes = actual.datos_sesion()
            invalidos = []
            for cambio in ("id_repetido", "activo_inexistente", "vinculo_roto", "resultado_incompleto", "parametro", "selector", "visible"):
                invalido = deepcopy(antes)
                if cambio == "id_repetido": invalido["historial_analisis"][1]["id"] = ids[0]
                elif cambio == "activo_inexistente": invalido["analisis_activo_id"] = "no existe"
                elif cambio == "vinculo_roto": invalido["valores"][0]["analisis_id"] = "no existe"
                elif cambio == "resultado_incompleto": invalido["historial_analisis"][0]["resultado"] = {}
                elif cambio == "parametro": invalido["historial_analisis"][0]["parametros"]["horizonte"] = 25
                elif cambio == "selector": invalido["historial_analisis"][0]["seleccion"]["idx_g"] = "texto"
                else: invalido["resultado_visible"] = {}
                invalidos.append(invalido)
            for invalido in invalidos:
                try:
                    actual.restaurar_datos_sesion(invalido)
                except ValueError:
                    pass
                else:
                    raise AssertionError("Debe rechazar el historial inválido.")
                assert actual.datos_sesion() == antes
        finally:
            if nuevo: nuevo.close()
            if nuevo is None: v.close()


if __name__ == "__main__":
    for nombre, funcion in sorted(globals().copy().items()):
        if nombre.startswith("test_") and callable(funcion):
            funcion()
            print("OK " + nombre, flush=True)
