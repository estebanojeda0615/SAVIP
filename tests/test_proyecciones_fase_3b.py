"""Fase 3B: formatos UI, componentes y consulta de resultados sin motor."""
from __future__ import annotations
import json
import math
import sys
import tempfile
import types
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from uuid import UUID

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tests.test_proyecciones_fase_3 import _app, anexo_sintetico, esperar, calcular, guardar, abrir
from tests.test_proyecciones_redisenio_visual import ventana_de_prueba
from PySide6.QtCore import QPoint, Qt
from PySide6.QtGui import QTextDocument
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QScrollArea
from openpyxl import load_workbook
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV
from app_icociv.interfaz.formato_numerico import numero_ui, detalle_ui
from app_icociv.interfaz.presentacion_resultados import construir_html_resultados
from app_icociv.persistencia.historial_icociv import serializar_resultado
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
from app_icociv.interfaz.tema import tokens
from app_icociv.interfaz.widgets.modelo_tabla import ModeloTablaPandas
import pandas as pd

OUT = ROOT / "docs/fase_3b"


def capturar(v, nombre, resultados=False, widget=None):
    carpeta = OUT / "capturas"
    carpeta.mkdir(parents=True, exist_ok=True)
    v.resize(1366, 768)
    v._cambiar_seccion(v.INDICE_RESULTADOS if resultados else 1)
    v.show()
    _app.processEvents()
    if widget:
        v.widget_proyecciones.findChild(QScrollArea).ensureWidgetVisible(widget, 20, 100)
    elif not resultados:
        v.widget_proyecciones.findChild(QScrollArea).verticalScrollBar().setValue(0)
    v.notificaciones.cerrar_todas()
    QTest.qWait(400)
    assert not v.velo_carga.isVisible()
    assert v.grab().save(str(carpeta / (nombre + ".png")))


def test_formato_ui_y_modelo_empalme_no_cambia():
    ejemplos = [(100000., "100,000"), (147.81, "147.81"), (156.29674485838413, "156.297"),
                (5.7416581140546255, "5.742"), (0., "0"), (-100000.25, "-100,000.25"),
                (.0000001, "< 0.001"), (-.0000001, "> -0.001")]
    for numero, esperado in ejemplos:
        assert numero_ui(numero) == esperado
    assert numero_ui(float("nan")) == "No disponible"
    assert numero_ui(float("inf")) == "No disponible"
    df = pd.DataFrame({"valor": [156.29674485838413]})
    original = ModeloTablaPandas(df)
    ui = ModeloTablaPandas(df, formateador=detalle_ui)
    assert original.index(0, 0).data() == "156.2967"  # Empalme conserva su presentación original.
    assert ui.index(0, 0).data() == "156.297"
    pd.testing.assert_frame_equal(ui._dataframe, df)
    (OUT / "EJEMPLOS_NUMERICOS.json").write_text(json.dumps(ejemplos, ensure_ascii=False, indent=2), encoding="utf-8")


def test_banda_lateral_continua_temas_tamanos_contraccion_y_carga():
    v = ventana_de_prueba()
    try:
        for tema in ("claro", "oscuro"):
            v._aplicar_tema(tema, persistir=False)
            for ancho, alto in ((1100, 700), (1366, 768), (1920, 1080)):
                v.resize(ancho, alto)
                _app.processEvents()
                assert v.navegacion.mapTo(v.centralWidget(), QPoint(0, 0)).y() == 0
                assert v.cabecera_app.mapTo(v.centralWidget(), QPoint(0, 0)).y() == 0
                assert v.navegacion.mapTo(v.centralWidget(), QPoint(0, 0)).x() == 0
                assert v.navegacion.height() == v.centralWidget().height()
                assert v.cabecera_app.mapTo(v.centralWidget(), QPoint(0, 0)).x() == v.navegacion.width()
            capturar(v, "lateral_continuo_" + tema)
            v.navegacion.establecer_contraccion(True)
            QTest.qWait(300)
            assert v.navegacion.width() == tokens.NAVEGACION_ANCHO_CONTRAIDO
            assert v.cabecera_app.mapTo(v.centralWidget(), QPoint(0, 0)).x() == v.navegacion.width()
            capturar(v, "lateral_contraido_" + tema)
            v.navegacion.establecer_contraccion(False)
            QTest.qWait(300)
        v._establecer_ocupado(True)
        assert not v.navegacion.isEnabled() and not v.selector_resultados_icociv.isEnabled()
        v._establecer_ocupado(False)
        assert v.navegacion.isEnabled() and v.selector_resultados_icociv.isEnabled()
    finally:
        v.close()


def test_inferencia_componente_y_nombres_compatibles():
    w = WidgetProyeccionesICOCIV()
    try:
        for nivel, visible, guardado in (("Grupo de obra", "Grupo", "Grupo de obra"),
                ("Subclase CPC", "Subclase", "Subclase de obra"), ("Tipología de obra", "Tipología", "Tipología"),
                ("Capítulo constructivo", "Capítulo", "Capítulo"), ("Grupo de costos", "Grupo", "Grupo de obra"),
                ("Ítem", "Ítem", "Ítem"), ("Insumo", "Insumo", "Insumo"), ("Otro nivel", "Otro", "Otro")):
            w.contexto = {"ruta": [{"nivel": "Tabla ICOCIV", "valor": "A16"}, {"nivel": nivel, "valor": "Prueba"}]}
            w.nombre_componente.setText("Nombre libre conservado")
            w._preseleccionar_componente()
            assert w.tipo_componente.currentText() == visible and w.tipo_componente.currentData() == guardado
            assert w.nombre_componente.text() == "Nombre libre conservado"
            w.tipo_componente.setCurrentText("Actividad")
            assert w.tipo_componente.currentData() == "Actividad"
        for tipo, visible in (("Grupo de obra", "Grupo"), ("Subclase de obra", "Subclase")):
            w.restaurar_calculadora({"tipo_componente": tipo, "nombre_componente": "Anterior"})
            assert w.tipo_componente.currentText() == visible and w._datos_calculadora()["tipo_componente"] == tipo
    finally:
        w.close()


def test_calculo_sesion_excel_identicos_al_respaldo_3b():
    fuente = OUT / "respaldo_inicio/app_icociv__interfaz__widgets__proyecciones_icociv.py.bak"
    base = types.ModuleType("baseline_fase3b")
    exec(compile(fuente.read_text(encoding="utf-8-sig"), str(fuente), "exec"), base.__dict__)
    antiguo = base.WidgetProyeccionesICOCIV()
    nuevo = WidgetProyeccionesICOCIV()
    nuevo.modalidad_importe.setCurrentIndex(nuevo.modalidad_importe.findData('unitario'))
    from tests.compatibilidad_modalidad_4e_b import sin_modalidad
    try:
        for w in (antiguo, nuevo):
            w.contexto = {"proyeccion_generada": True, "indice_base": 147.81, "indice_proyectado": 156.29674485838413,
                          "ruta_texto": "Ruta de prueba", "serie_final": "Serie de prueba", "horizonte_solicitado": 3,
                          "modelo": "Prueba de presentación", "periodo_base": "2026-12", "periodo_proyectado": "2027-03"}
            w._actualizar_estado_calculadora()
            w.unidad.setCurrentIndex(1)
            w.precio_base.setValue(100000)
            w.anticipo.setValue(100)
            w.cantidad.setValue(1234.5678)
        with patch.object(base, "uuid4", return_value=UUID(int=1)), \
             patch("app_icociv.interfaz.widgets.proyecciones_icociv.uuid4", return_value=UUID(int=1)), \
             patch("app_icociv.servicios.actualizacion_icociv.datetime") as reloj:
            reloj.now.return_value = datetime(2026, 10, 9, 12, 0)
            antiguo.boton_calcular.click()
            nuevo.boton_calcular.click()
        assert all(r['modalidad_importe']=='unitario' for r in nuevo.valores)
        assert sin_modalidad(nuevo.valores) == antiguo.valores
        assert sin_modalidad(nuevo.datos_sesion()) == antiguo.datos_sesion()
        texto = nuevo.resultado.toPlainText()
        assert "156.297" in texto and "100,000" in texto and "1234.5678" not in texto
        assert "156.29674485838413" not in texto and "100000.00000000" not in texto
        assert nuevo.valores[0]["i_icociv"] == 156.29674485838413
        with tempfile.TemporaryDirectory(prefix="savip_3b_excel_") as tmp:
            tmp = Path(tmp)
            for nombre, w in (("anterior", antiguo), ("actual", nuevo)):
                w._generar_excel(str(tmp / (nombre + ".xlsx")))
            a, b = [load_workbook(tmp / (n + ".xlsx"), data_only=False) for n in ("anterior", "actual")]
            assert b.sheetnames == ['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']
            for fila,r in enumerate(nuevo.valores,3):
                assert b['Trazabilidad'].cell(fila,16).value == float(format(r['valor_proyectado'],'.16g'))
                for col,key in ((7,'precio_base'),(8,'anticipo_amortizado'),(11,'i0_icociv'),(12,'i_icociv')):
                    assert b['Valores proyectados'].cell(fila,col).value == float(format(r[key],'.16g'))
                assert r['modalidad_importe']=='unitario'
                assert all(b['Valores proyectados'].cell(fila,col).data_type=='f' for col in (9,10,14,15))
                assert b['Valores proyectados'].cell(fila,13).value == f'=L{fila}/K{fila}'
                assert b['Valores proyectados'].cell(fila,16).value == f'=(M{fila}-1)*100'
            a.close(); b.close()
    finally:
        antiguo.close(); nuevo.close()


def test_tres_proyecciones_selectores_precision_historial_y_exportacion_reales():
    with tempfile.TemporaryDirectory(prefix="savip_3b_real_") as tmp:
        tmp = Path(tmp)
        fuente = tmp / "anexo_sintetico_3b.xlsx"
        anexo_sintetico(fuente)
        v = VentanaPrincipal()
        nuevo = None
        try:
            v._cargar_archivo(str(fuente))
            esperar(v, lambda: v.combo_grupo.count() == 3 and v.combo_grupo.isEnabled())
            w = v.widget_proyecciones
            w.nombre_proyecto.setText("Prueba Fase 3B · anexo sintético")
            for secuencia, (grupo, horizonte) in enumerate(((1, 1), (1, 3), (2, 1)), 1):
                v.combo_grupo.setCurrentIndex(grupo)
                v.combo_horizonte.setCurrentIndex(v.combo_horizonte.findData(horizonte))
                if secuencia == 3:
                    v.combo_nivel2.setCurrentIndex(1)
                v.ejecutar_proyeccion()
                esperar(v, lambda: len(w.historial) == secuencia)
                assert w.tipo_componente.currentText() == ("Subclase" if secuencia == 3 else "Grupo")
                assert w.historial[-1]["calculadora"]["tipo_componente"] == ("Subclase de obra" if secuencia == 3 else "Grupo de obra")
                if secuencia == 1:
                    v._aplicar_tema("claro", persistir=False)
                    capturar(v, "componente_grupo_inferido", widget=w.tipo_componente)
                if secuencia == 3:
                    capturar(v, "componente_subclase_inferido", widget=w.tipo_componente)
                calcular(w, f"Actividad manual {secuencia}")
            ids = [a["id"] for a in w.historial]
            formulario = (w.tipo_componente.currentText(), w.nombre_componente.text())
            expected = deepcopy(w.datos_sesion())
            estados = []
            with patch.object(v.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor al consultar Resultados")) as motor, \
                 patch.object(v, "ejecutar_proyeccion", side_effect=AssertionError("Proyección al consultar")):
                for tema in ("claro", "oscuro"):
                    v._aplicar_tema(tema, persistir=False)
                    for i in (0, 1, 2, 0):
                        v.selector_resultados_icociv.setCurrentIndex(v.selector_resultados_icociv.findData(ids[i]))
                        assert w.analisis_activo_id == ids[i]
                        assert w.selector_analisis.currentData() == ids[i]
                        assert v.selector_resultados_icociv.currentData() == ids[i]
                        assert serializar_resultado(v.resultado_ui_actual) == w.historial[i]["resultado"]
                        esperado_html = QTextDocument()
                        esperado_html.setHtml(construir_html_resultados(v.resultado_ui_actual, tema))
                        assert v.texto_resultados.toPlainText() == esperado_html.toPlainText()
                        assert v.valores_kpi["indice"].text() == numero_ui(w.historial[i]["contexto"]["indice_proyectado"])
                        # Fase 4G-B: activar un análisis NO recupera ni sobrescribe la calculadora.
                        assert (w.tipo_componente.currentText(), w.nombre_componente.text()) == formulario
                        capturar(v, f"resultados_analisis_{i+1}_{tema}", resultados=True)
                        assert v.selector_resultados_icociv.isVisible()
                        estados.append({"tema": tema, "analisis": i + 1, "id": ids[i],
                                        "indice_real": v.controlador.proyeccion_actual["y_proj"],
                                        "indice_ui": v.valores_kpi["indice"].text(), "tipo_manual_conservado": True})
                # Selección inversa desde Proyecciones actualiza también el selector Resultados.
                w.selector_analisis.setCurrentIndex(w.selector_analisis.findData(ids[1]))
                assert v.selector_resultados_icociv.currentData() == ids[1]
                assert w.tipo_componente.currentText() == "Actividad"
                assert w.valores == expected["valores"] and len(w.proyecciones) == 3
                capturar(v, "calculadora_resultado_tres_decimales", widget=w.resultado)
                capturar(v, "tabla_operativa_tres_decimales", widget=w.tabla_valores)
                assert motor.call_count == 0
                w.selector_analisis.setCurrentIndex(w.selector_analisis.findData(expected["analisis_activo_id"]))
                assert w.datos_sesion() == expected
                ruta = tmp / "sesion.json"
                guardar(v, ruta)
                guardado = cargar_sesion(ruta)
                xlsx = tmp / "historial.xlsx"
                w._generar_excel(str(xlsx))
                libro = load_workbook(xlsx, data_only=False)
                # Fase 4G: el factor se ubica por su encabezado (puede haber columnas de fecha antes).
                col_factor = next(c.column for c in libro["Valores proyectados"][2] if c.value == 'Factor')
                for i, fila in enumerate(w.valores, 3):
                    assert libro["Valores proyectados"].cell(i, col_factor).data_type == 'f'
                    assert libro["Trazabilidad"].cell(i, 16).value == float(format(fila["valor_proyectado"],'.16g'))
                    assert libro['Trazabilidad'].cell(i,3).value == fila['analisis_id']
                libro.close()
            v.close()
            nuevo = VentanaPrincipal()
            with patch.object(nuevo.controlador, "ejecutar_analisis", side_effect=AssertionError("Motor al reabrir")) as motor:
                abrir(nuevo, ruta)
                assert nuevo.widget_proyecciones.datos_sesion() == guardado["proyecciones_icociv"]
                assert nuevo.selector_resultados_icociv.currentData() == ids[2]
                assert nuevo.selector_resultados_icociv.count() == 4
                capturar(nuevo, "resultados_sesion_reabierta", resultados=True)
                assert motor.call_count == 0
            (OUT / "INTEGRACION_REAL.json").write_text(json.dumps({"analisis": 3, "registros": 3,
                "motor_consulta_llamadas": 0, "sesion_exacta": True, "excel_precision_interna": True,
                "estados": estados, "datos": "Anexo sintético, cargador y motor reales intactos"}, ensure_ascii=False, indent=2), encoding="utf-8")
        finally:
            if nuevo: nuevo.close()
            v.close()


if __name__ == "__main__":
    for nombre, funcion in sorted(globals().copy().items()):
        if nombre.startswith("test_") and callable(funcion):
            funcion()
            print("OK " + nombre, flush=True)
