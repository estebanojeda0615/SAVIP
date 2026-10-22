import json,math
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest,QSignalSpy
from PySide6.QtWidgets import QAbstractSlider,QStyle,QStyleOptionSlider
from openpyxl import load_workbook
from tests.test_slider_comparacion_historica_fase_4d import APP,reales,sintetico,cargar,sin_motor,originales
from tests.test_historial_empalme_fase_4b import guardar,abrir
from app_icociv.interfaz.widgets.proyecciones_icociv import WidgetProyeccionesICOCIV,COLUMNAS_VALORES
from app_icociv.interfaz.ventana_principal import VentanaPrincipal
from app_icociv.persistencia.historial_icociv import recuperar_resultado,serializar_resultado
from app_icociv.persistencia.gestor_sesiones import cargar_sesion
ROOT=Path(__file__).resolve().parents[1]
HOJAS=['Resumen','Valores proyectados','Proyecciones ICOCIV','Trazabilidad','Metodología y fórmulas']

def comprobar(w,r,referencia):
    # Fase 4I-C: el escenario usa el factor APLICADO (variación redondeada a 2 decimales, mitad hacia arriba).
    from decimal import Decimal
    from app_icociv.servicios.actualizacion_icociv import factor_aplicado
    e,_=w.escenario_registro(r);base=r['precio_base']-r['anticipo_amortizado'];p,F=factor_aplicado(r['i_icociv'],referencia)
    assert e['i0_icociv']==referencia and e['i_icociv']==r['i_icociv'] and e['cantidad']==r['cantidad']
    assert e['base_ajustable']==base and e['factor_icociv']==float(F) and e['diferencia_porcentual']==float(p)
    assert e['r_total']==float(Decimal(repr(base))*(F-1)) and e['valor_proyectado']==float(Decimal(repr(base))*F)
    return e

def test_inicializacion_movimiento_real_y_retorno_sin_drift():
    w=sintetico()
    try:
        r=deepcopy(w.valores[0]);antes=originales(w);spy=QSignalSpy(w.comparacion_cambiada)
        assert w.escenario_registro(r)[0]==r and w._periodo_por_analisis[w.analisis_activo_id]=='2026-12'
        w.slider_comparacion.setValue(0);assert w.escenario_registro(r)[0]==r and spy.count()==0
        with sin_motor(),patch('app_icociv.interfaz.widgets.proyecciones_icociv.calcular_actualizacion_icociv',side_effect=AssertionError('Servicio en consulta')):
            w.slider_comparacion.triggerAction(QAbstractSlider.SliderSingleStepAdd)
            assert w.selector_referencia.currentData()=='2024-02' and spy.count()==1
            e=comprobar(w,r,104.7619047619);assert math.isclose(e['valor_proyectado'],945,abs_tol=1e-9)
            assert w.modelo_valores.index(0,10).data(Qt.UserRole)==e['i0_icociv']
            assert 'Valor vigente:' in w.modelo_valores.index(0,10).data(Qt.ToolTipRole)
            assert w.modelo_valores.index(0,13).data(Qt.UserRole) is None  # Sin clasificar: principal en resumen y detalle.
            assert '945' in w.resultado.toPlainText() and '990' in w.resultado.toPlainText()
            QTest.keyClick(w.slider_comparacion,Qt.Key_Right);assert w.escenario_registro(r)[0]['diferencia_porcentual']<0
            for _ in range(5):
                assert w.seleccionar_referencia_historica('2024-02')
                assert w.seleccionar_referencia_historica('2026-12') and w.escenario_registro(r)[0]==r
            assert originales(w)==antes
    finally:w.close()

def test_selectores_sincronizados_filtro_senales_y_scroll():
    v=VentanaPrincipal();w=cargar(v)
    try:
        ids=[a['id'] for a in w.historial];spy=QSignalSpy(w.analisis_solicitado)
        # Fase 4G-B: el desplegable inferior es «Proyección» del ámbito: «Todas…» + los mismos análisis en el mismo orden.
        assert [w.selector_analisis_valores.itemData(i) for i in range(w.selector_analisis_valores.count())]==[w.selector_analisis.itemData(i) for i in range(w.selector_analisis.count())]
        assert w.selector_analisis_valores.itemText(0)=='Todas las proyecciones de esta selección'
        with sin_motor() as motor,patch.object(v.controlador,'cargar_archivo',side_effect=AssertionError('DANE')):
            for combo,identificador in ((w.selector_analisis_valores,ids[0]),(w.selector_analisis,ids[1]),(w.selector_analisis_valores,ids[2])):
                cuenta=spy.count();combo.setCurrentIndex(combo.findData(identificador));APP.processEvents()
                assert spy.count()==cuenta+1 and w.analisis_activo_id==identificador
                assert w.selector_analisis.currentData()==w.selector_analisis_valores.currentData()==v.selector_resultados_icociv.currentData()==identificador
            w.solo_analisis_activo.setChecked(True);assert len(w._valores_visibles)==1
            actual=w.analisis_activo_id;w.selector_analisis_valores.setCurrentIndex(0)
            # Fase 4G-B: «Todas» solo amplía la consulta; no desactiva ni cambia el análisis para nuevos cálculos.
            assert w.analisis_activo_id==actual and w.solo_analisis_activo.isChecked() and len(w._valores_visibles)==len(w.valores)
            w.selector_analisis.setCurrentIndex(0)  # «Sin análisis activo» real (selector superior)
            assert w.analisis_activo_id is None and v.selector_resultados_icociv.currentData() is None
            w.solo_analisis_activo.setChecked(False);w.tabla_valores.selectRow(2)
            calculo=w._valores_visibles[2]['calculo_id'];scroll=w.tabla_valores.horizontalScrollBar();scroll.setValue(50);pos=scroll.value()
            w.seleccionar_referencia_historica('2024-02');assert scroll.value()==pos
            assert w._valores_visibles[w.tabla_valores.selectedIndexes()[0].row()]['calculo_id']==calculo
            assert motor.call_count==0
        w.historial[0]['nombre']='Nombre cambiado';w._refrescar_historial();assert w.selector_analisis.itemText(1)=='Nombre cambiado' and w.selector_analisis_valores.itemData(1)==w.historial[0]['id']  # 4G-B: etiqueta corta abajo
        v._establecer_ocupado(True);assert not w.selector_analisis_valores.isEnabled()
        v._establecer_ocupado(False);assert w.selector_analisis_valores.isEnabled()
    finally:v.close()

def test_individual_global_atomico_y_desactivacion():
    v=VentanaPrincipal();w=cargar(v)
    try:
        ids=[a['id'] for a in w.historial];v._activar_analisis(ids[0]);original=originales(w)
        previo=[w.escenario_registro(r)[0] for r in w.valores]
        assert w.seleccionar_referencia_historica('2024-02')
        assert [w.escenario_registro(r)[0] for r in w.valores][1:]==previo[1:]
        w.solo_analisis_activo.setChecked(True);antes=[w.escenario_registro(r)[0] for r in w.valores]
        w.aplicar_global.setChecked(True);assert [w.escenario_registro(r)[0] for r in w.valores]==antes
        assert w.seleccionar_referencia_historica('2024-05') and len(w._valores_visibles)==1
        for r in w.valores:comprobar(w,r,w._historicos_comparacion[r['analisis_id']]['2024-05'])
        fechas=deepcopy(w._periodo_por_analisis);economia=[w.escenario_registro(r)[0] for r in w.valores]
        assert not w.seleccionar_referencia_historica('1900-01') and w._periodo_por_analisis==fechas
        assert [w.escenario_registro(r)[0] for r in w.valores]==economia
        w.aplicar_global.setChecked(False);assert [w.escenario_registro(r)[0] for r in w.valores]==economia
        assert w.seleccionar_referencia_historica('2024-08') and [w.escenario_registro(r)[0] for r in w.valores][1:]==economia[1:]
        w.solo_analisis_activo.setChecked(False);assert originales(w)==original
    finally:v.close()

def test_global_fuente_ambigua_sin_interseccion_y_error_economico():
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales())
    try:
        antes=deepcopy(w._periodo_por_analisis)
        w._historicos_comparacion[w.historial[-1]['id']]={};w.aplicar_global.setChecked(True)
        assert not w.aplicar_global.isChecked() and w._periodo_por_analisis==antes and 'globalmente' in w.lbl_estado_comparacion.text()
        w._preparar_comparacion_historica();w.aplicar_global.setChecked(True)
        ultimo=w.valores[-1];p=ultimo['precio_base'];ultimo['precio_base']=float('inf')
        assert not w.seleccionar_referencia_historica('2024-02') and w._periodo_por_analisis==antes
        ultimo['precio_base']=p
        ultimo_id=w.historial[-1]['id'];h=w._historicos_comparacion[ultimo_id];indice=h['2024-02'];h['2024-02']=1e-320
        assert not w.seleccionar_referencia_historica('2024-02') and w._periodo_por_analisis==antes
        h['2024-02']=indice
        for i,identificador in enumerate(w._historicos_comparacion):w._historicos_comparacion[identificador]={f'202{i}-01':100.}
        w.aplicar_global.setChecked(False);w.aplicar_global.setChecked(True);assert not w.aplicar_global.isChecked()
    finally:w.close()

def test_migracion_v1_sin_repricing_y_aplicacion_explicita():
    datos=reales();identificador=datos['historial_analisis'][0]['id']
    datos['comparacion_historica']={'version':1,'periodo_por_analisis':{identificador:'2024-02'}}
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(datos);w.establecer_analisis_activo(identificador)
    try:
        r=w.valores[0];assert w.escenario_registro(r)[0]==r and not w._overrides[identificador]
        assert w._periodo_por_analisis[identificador]=='2026-12' and 'aún no aplicada' in w.lbl_estado_comparacion.text()
        w.boton_aplicar_consulta.click();assert w._overrides[identificador] and w._periodo_por_analisis[identificador]=='2024-02'
        assert w.escenario_registro(r)[0]!=r and not w._consultas_v1
        datos=w.datos_sesion();otro=WidgetProyeccionesICOCIV()
        try:otro.restaurar_datos_sesion(datos);assert otro.escenario_registro(otro.valores[0])==w.escenario_registro(r)
        finally:otro.close()
    finally:w.close()

def test_i0_original_discrepante_y_legados_no_reinterpretados():
    w=sintetico()
    try:
        r=w.valores[0];r['i0_icociv']=95.;antes=deepcopy(r);w._actualizar_controles_comparacion()
        assert w.escenario_registro(r)[0]==antes and 'distinto' in w.escenario_registro(r)[1]
        w.show();APP.processEvents();slider=w.slider_comparacion
        opcion=QStyleOptionSlider();slider.initStyleOption(opcion)
        mango=slider.style().subControlRect(QStyle.CC_Slider,opcion,QStyle.SC_SliderHandle,slider)
        QTest.mouseClick(slider,Qt.LeftButton,pos=mango.center());comprobar(w,r,100.)
        assert w.valores[0]==antes
        for clave,valor in (('analisis_id',None),('ruta_icociv','otra'),('i_icociv',111.)):
            legado=deepcopy(r);legado[clave]=valor
            assert w.escenario_registro(legado)[0]==legado and 'vínculo no demostrable' in w.escenario_registro(legado)[1]
    finally:w.close()

def test_fechas_encabezados_uniformes_mixtos_y_copia_numerica():
    w=sintetico()
    try:
        w.seleccionar_referencia_historica('2024-02');m=w.modelo_valores
        # Fase 4G: fecha única después del título técnico.
        assert m.headerData(10,Qt.Horizontal)=='I0 ICOCIV\n2024-02' and m.headerData(11,Qt.Horizontal).startswith('I ICOCIV proyectado\n')
        w.tabla_valores.selectRow(0);QTest.keyClick(w.tabla_valores,Qt.Key_C,Qt.ControlModifier)
        copia=APP.clipboard().text().split('\t');assert float(copia[10])==104.7619047619 and len(copia)==17 and '2024-02' in m.index(0,10).data(Qt.ToolTipRole)
        assert len(COLUMNAS_VALORES)==17
    finally:w.close()
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales())
    try:
        # Fase 4G: I0 único (fecha en el encabezado) e I heterogéneo (columna auxiliar junto a I).
        assert w.modelo_valores.headerData(10,Qt.Horizontal)=='I0 ICOCIV\n2026-12' and w.modelo_valores.headerData(12,Qt.Horizontal)=='Fecha de I proyectado (AAAA-MM)'
        assert all(w.fechas_registro(r)[1] in w.modelo_valores.index(i,11).data(Qt.ToolTipRole) for i,r in enumerate(w.valores))
    finally:w.close()

def test_excel_cinco_hojas_formulas_y_trazabilidad(tmp_path):
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales());w.establecer_analisis_activo(w.historial[0]['id'])
    try:
        for r in w.valores:r['modalidad_importe']='unitario'  # Declaración explícita para comprobar fórmulas fuente.
        antes=originales(w);w.seleccionar_referencia_historica('2024-02');w.valores[0]['item']='=1+1'
        p=tmp_path/'escenario.xlsx';w._generar_excel(str(p));b=load_workbook(p);s=b['Valores proyectados'];traza=b['Trazabilidad']
        assert b.sheetnames==HOJAS and all(h.sheet_state=='visible' for h in b) and not b.defined_names and not b._external_links
        # Fase 4G: I0 (2024-02/2026-12) e I (tres objetivos) heterogéneos → 19 columnas; L y N son fechas.
        assert s.max_column==19 and s['D3'].data_type=='s' and s['D3'].value=='=1+1'
        for fila,r in enumerate(w.valores,3):
            e=w.escenario_registro(r)[0]
            assert s.cell(fila,11).value==float(format(e['i0_icociv'],'.16g'))
            # Fase 4I-C: escenario recalculado → factor aplicado con ROUND; registro histórico → I/I0 original.
            from app_icociv.servicios.actualizacion_icociv import es_factor_legado
            factor,diferencia=((f'=M{fila}/K{fila}',f'=(O{fila}-1)*100') if es_factor_legado(e)
                               else (f'=1+ROUND((M{fila}/K{fila}-1)*100,2)/100',f'=ROUND((O{fila}-1)*100,2)'))
            assert [s.cell(fila,c).value for c in (9,15,16,18)]==[f'=G{fila}-H{fila}',factor,f'=I{fila}+I{fila}*(O{fila}-1)',diferencia]
            assert all(s.cell(fila,c).data_type=='f' for c in (9,15,16,18))
            assert tuple(traza.cell(fila,c).value for c in (9,11))==w.fechas_registro(r)
            assert traza.cell(fila,8).value==float(format(r['i0_icociv'],'.16g')) and traza.cell(fila,16).value==float(format(r['valor_proyectado'],'.16g'))
        assert not any(c.comment for h in b for row in h for c in row)
        assert not any('[' in c.value or 'IF(' in c.value for h in b for row in h for c in row if c.data_type=='f')
        assert not any('!' in c.value for row in s for c in row if c.data_type=='f')  # 4E autoriza enlaces internos sólo en el desglose.
        b.close();w.valores[0]['item']=antes['valores'][0]['item'];assert originales(w)==antes
    finally:w.close()

def test_sesion_v2_mixta_reapertura_sin_fuente_y_corrupcion(tmp_path):
    v=VentanaPrincipal();w=cargar(v);otro=None
    try:
        emp=json.loads((ROOT/'docs/fase_4c/FIXTURE_DATOS.json').read_text(encoding='utf-8'))['empalme_iccp_icociv'];v.widget_empalme.restaurar_datos_sesion(emp)
        ausente=tmp_path/'DANE_AUSENTE.xlsx';assert not ausente.exists()
        for a in w.historial:a['archivo_excel']=str(ausente)
        v.controlador.ruta_archivo=ausente
        v._activar_analisis(w.historial[0]['id']);w.aplicar_global.setChecked(True);w.seleccionar_referencia_historica('2024-03')
        escenario=[w.escenario_registro(r) for r in w.valores];antes=deepcopy(w.datos_sesion());emp_antes=deepcopy(v.widget_empalme.datos_sesion())
        p=tmp_path/'sin_dane.json';guardar(v,p);guardado=cargar_sesion(p)['proyecciones_icociv']
        assert all(guardado[k]==valor for k,valor in antes.items());otro=VentanaPrincipal()
        with sin_motor() as motor,patch.object(otro.controlador,'cargar_archivo',side_effect=AssertionError('DANE')):
            abrir(otro,p);n=otro.widget_proyecciones
            assert n.datos_sesion()==guardado and [n.escenario_registro(r) for r in n.valores]==escenario and n.aplicar_global.isChecked()
            assert otro.widget_empalme.datos_sesion()==emp_antes and motor.call_count==0
        for alterar in ('periodo','override','id','global'):
            d=deepcopy(guardado);b=d['comparacion_historica'];identificador=w.historial[0]['id']
            if alterar=='periodo':b['selecciones'][identificador]['periodo']='1900-01'
            elif alterar=='override':b['selecciones'][identificador]['override']='sí'
            elif alterar=='id':b['selecciones']['ajeno']=b['selecciones'].pop(identificador)
            else:b['global']='sí'
            previo=n.datos_sesion()
            try:n.restaurar_datos_sesion(d)
            except ValueError:pass
            else:raise AssertionError('Corrupción aceptada')
            assert n.datos_sesion()==previo
    finally:
        if otro:otro.close()
        v.close()

def test_historial_nuevo_global_y_referencias_huerfanas():
    w=WidgetProyeccionesICOCIV();w.restaurar_datos_sesion(reales())
    try:
        w.aplicar_global.setChecked(True);w.establecer_analisis_activo(w.historial[0]['id']);w.seleccionar_referencia_historica('2024-02');previo=deepcopy(w._periodo_por_analisis)
        a=deepcopy(w.historial[0]);a['id']='nuevo';a['registros_ids']=[];w.historial.append(a);w._preparar_comparacion_historica();w._refrescar_historial()
        assert w._periodo_por_analisis['nuevo']=='2026-12' and not w._overrides['nuevo'] and all(w._periodo_por_analisis[k]==p for k,p in previo.items())
        w.historial.pop();w._preparar_comparacion_historica();w._refrescar_historial();assert 'nuevo' not in w._periodo_por_analisis
    finally:w.close()

def test_documentos_escenario_efectivo_sin_capitulo_duplicado():
    from app_icociv.reportes.sesion_modulos import construir_proyecciones
    from dataclasses import asdict
    w=sintetico()
    try:
        original=deepcopy(w.valores);w.seleccionar_referencia_historica('2024-02');d=asdict(construir_proyecciones(w));s=next(s for s in d['secciones'] if s['clave']=='valores')
        assert not any(s['clave']=='comparacion_historica' for s in d['secciones'])
        texto=json.dumps(s,ensure_ascii=False);assert '945' in texto and '2024-02' in texto and 'Auditoría' in texto and '990' in texto
        assert w.valores==original
    finally:w.close()
